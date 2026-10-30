"""
Sinh rubric chấm điểm theo CLO từ một nguồn duy nhất: <hoc-phan>/rubric/rubric.json

Cách dùng:
    uv run --with openpyxl python build-rubric.py lap-trinh-ung-dung-mobile

Đầu ra:
    <hoc-phan>/rubric/Rubric-<mã học phần>.xlsx   rubric + bảng điểm có công thức + mức đạt CLO
    <hoc-phan>/cham-diem.html                      công cụ chấm điểm trên web (từ templates/rubric-grading.html),
                                                   xuất điểm theo tiêu chí và CLO ra Excel

Kiểm tra: tổng điểm tiêu chí của mỗi CLO bằng điểm tối đa của CLO, tổng các CLO bằng thang điểm,
mỗi tiêu chí có đủ mô tả các mức, gợi ý chấm và minh chứng.
"""
import argparse
import json
import sys
from pathlib import Path

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT_DIR = Path(__file__).resolve().parent
TEMPLATE = ROOT_DIR / "templates" / "rubric-grading.html"
DATA_ROWS = 60  # số dòng sinh viên sẵn trong bảng điểm Excel


# ---------------------------------------------------------------- kiểm tra dữ liệu
def validate(rubric: dict) -> list[str]:
    errors = []
    level_codes = [lv["code"] for lv in rubric["levels"]]
    described = [c for c in level_codes if c != "0"]
    total = 0.0
    seen = set()
    for clo in rubric["clos"]:
        s = round(sum(c["max"] for c in clo["criteria"]), 4)
        total += clo["max"]
        if abs(s - clo["max"]) > 1e-9:
            errors.append(f"{clo['id']}: tổng tiêu chí {s} khác điểm tối đa {clo['max']}")
        for c in clo["criteria"]:
            if c["id"] in seen:
                errors.append(f"Trùng mã tiêu chí {c['id']}")
            seen.add(c["id"])
            if c["scope"] not in ("group", "individual"):
                errors.append(f"{c['id']}: scope phải là group hoặc individual")
            missing = [code for code in described if not c.get("levels", {}).get(code)]
            if missing:
                errors.append(f"{c['id']}: thiếu mô tả mức {', '.join(missing)}")
            for key in ("guide", "evidence"):
                if not c.get(key):
                    errors.append(f"{c['id']}: thiếu {key}")
    if abs(total - rubric["course"]["scale"]) > 1e-9:
        errors.append(f"Tổng điểm các CLO {total} khác thang điểm {rubric['course']['scale']}")
    return errors


# ---------------------------------------------------------------- Excel
def build_excel(rubric: dict, out: Path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as col
    from openpyxl.worksheet.datavalidation import DataValidation

    course = rubric["course"]
    levels = rubric["levels"]
    described = [lv for lv in levels if lv["code"] != "0"]
    clos = rubric["clos"]
    criteria = [(clo, c) for clo in clos for c in clo["criteria"]]

    BLUE, DARK, LIGHT = "013D88", "01295C", "EEF5FC"
    CLO_FILL = ["E7F1FF", "EAFAF1", "FFF5EB", "F3E5F5"]
    thin = Side(style="thin", color="B8C4D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=BLUE)
    scope_text = {"group": "Nhóm", "individual": "Cá nhân"}

    def header(ws, row, values, fill=head_fill, font=head_font):
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.fill, cell.font, cell.alignment, cell.border = fill, font, center, border

    wb = Workbook()

    # ----- Hướng dẫn
    ws = wb.active
    ws.title = "Hướng dẫn"
    ws.column_dimensions["A"].width = 120
    lines = [
        (f"RUBRIC CHẤM ĐIỂM CUỐI KỲ - {course['name'].upper()} ({course['code']})", Font(bold=True, size=14, color=BLUE)),
        (course["assessment"], Font(italic=True)),
        ("", None),
        ("Cấu trúc", Font(bold=True, color=DARK)),
        (f"- {len(clos)} CLO, {len(criteria)} tiêu chí; mỗi tiêu chí chỉ chấm một vấn đề. Tổng điểm: {course['scale']}.", None),
        ("- " + "; ".join(f"{clo['id']} ({clo['plo']}): {clo['max']} điểm, {len(clo['criteria'])} tiêu chí" for clo in clos) + ".", None),
        ("- Mức chấm: " + "; ".join(f"{lv['code']} {lv['name']} = {int(lv['ratio'] * 100)}% điểm tối đa" for lv in levels) + ".", None),
        ("- Có thể cho điểm lẻ giữa hai mức (bước 0,05) khi minh chứng nằm giữa hai mô tả.", None),
        ("- Phạm vi 'Nhóm': chấm trên sản phẩm chung, mặc định cùng điểm cho cả nhóm; hạ riêng cho sinh viên không giải thích được phần mã liên quan.", None),
        ("- Phạm vi 'Cá nhân': chấm riêng từng sinh viên dựa trên minh chứng của chính sinh viên đó.", None),
        ("", None),
        ("Ngưỡng đạt", Font(bold=True, color=DARK)),
        (f"- {course['note']}", None),
        ("", None),
        ("Cách dùng file", Font(bold=True, color=DARK)),
        ("1. Sheet 'Rubric': mô tả mức, gợi ý chấm, minh chứng cần có của từng tiêu chí.", None),
        (f"2. Sheet 'Bảng điểm': nhập MSSV, họ tên, nhóm và điểm từng tiêu chí (0 đến điểm tối đa, có kiểm tra dữ liệu). Điểm CLO, tổng, Đạt/Chưa đạt tự tính. Sẵn {DATA_ROWS} dòng.", None),
        ("3. Sheet 'Mức đạt CLO': tỷ lệ sinh viên đạt từng CLO và kết luận đạt ở cấp lớp; điểm trung bình theo tiêu chí.", None),
        ("4. Có thể chấm trên trang web cham-diem.html của học phần rồi xuất Excel; file xuất có cùng cấu trúc cột.", None),
    ]
    for r, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True)
        if font:
            cell.font = font

    # ----- Rubric
    ws = wb.create_sheet("Rubric")
    cols = ["Mã", "Tiêu chí", "Phạm vi", "Điểm tối đa"] + [f"{lv['code']} - {lv['name']} ({int(lv['ratio'] * 100)}%)" for lv in described] + ["Gợi ý chấm", "Minh chứng cần có"]
    widths = [7, 26, 10, 9] + [34] * len(described) + [40, 34]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[col(i)].width = w
    header(ws, 1, cols)
    r = 2
    for k, clo in enumerate(clos):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
        cell = ws.cell(row=r, column=1, value=f"{clo['id']} ({clo['plo']}) - {clo['name']} - {clo['max']} điểm (đạt khi ≥ {clo['max'] * course['studentThreshold']:g})")
        cell.font = Font(bold=True, color=DARK)
        cell.fill = PatternFill("solid", fgColor=CLO_FILL[k % 4])
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 30
        r += 1
        for c in clo["criteria"]:
            values = [c["id"], c["name"], scope_text[c["scope"]], c["max"]] + [c["levels"][lv["code"]] for lv in described] + [c["guide"], c["evidence"]]
            for i, v in enumerate(values, start=1):
                cell = ws.cell(row=r, column=i, value=v)
                cell.alignment, cell.border = (center if i in (1, 3, 4) else wrap), border
            ws.cell(row=r, column=2).font = Font(bold=True)
            r += 1
    ws.freeze_panes = "C2"

    # ----- Bảng điểm
    ws = wb.create_sheet("Bảng điểm")
    info_cols = ["STT", "MSSV", "Họ và tên", "Nhóm"]
    first_crit = len(info_cols) + 1
    crit_col = {}
    c_idx = first_crit
    for clo, c in criteria:
        crit_col[c["id"]] = c_idx
        c_idx += 1
    clo_col = {clo["id"]: c_idx + i for i, clo in enumerate(clos)}
    total_col = c_idx + len(clos)
    pass_col = {clo["id"]: total_col + 1 + i for i, clo in enumerate(clos)}
    note_col = total_col + 1 + len(clos)
    last_col = note_col

    # Hàng 1: nhóm CLO; hàng 2: tiêu chí; hàng 3: điểm tối đa
    for i, name in enumerate(info_cols, start=1):
        ws.merge_cells(start_row=1, start_column=i, end_row=3, end_column=i)
    header(ws, 1, info_cols)
    for k, clo in enumerate(clos):
        ids = [c["id"] for c in clo["criteria"]]
        a, b = crit_col[ids[0]], crit_col[ids[-1]]
        ws.merge_cells(start_row=1, start_column=a, end_row=1, end_column=b)
        cell = ws.cell(row=1, column=a, value=f"{clo['id']} ({clo['max']} điểm)")
        for cc in range(a, b + 1):
            ws.cell(row=1, column=cc).fill = PatternFill("solid", fgColor=CLO_FILL[k % 4])
            ws.cell(row=1, column=cc).border = border
        cell.font, cell.alignment = Font(bold=True, color=DARK), center
        for c in clo["criteria"]:
            x = crit_col[c["id"]]
            h = ws.cell(row=2, column=x, value=f"{c['id']} {c['name']}" + (" (CN)" if c["scope"] == "individual" else ""))
            h.alignment, h.border, h.font = center, border, Font(bold=True, size=9)
            h.fill = PatternFill("solid", fgColor=CLO_FILL[k % 4])
            m = ws.cell(row=3, column=x, value=c["max"])
            m.alignment, m.border, m.font = center, border, Font(italic=True, size=9)
            ws.column_dimensions[col(x)].width = 11
    for title, colmap, fill in (("Điểm CLO", clo_col, "D6E4F5"), ("Đạt CLO", pass_col, "FCE8C8")):
        a, b = min(colmap.values()), max(colmap.values())
        ws.merge_cells(start_row=1, start_column=a, end_row=1, end_column=b)
        cell = ws.cell(row=1, column=a, value=title)
        cell.font, cell.alignment = Font(bold=True, color=DARK), center
        for clo in clos:
            x = colmap[clo["id"]]
            ws.cell(row=1, column=x).fill = PatternFill("solid", fgColor=fill)
            ws.cell(row=1, column=x).border = border
            h = ws.cell(row=2, column=x, value=clo["id"])
            h.font, h.alignment, h.border, h.fill = Font(bold=True), center, border, PatternFill("solid", fgColor=fill)
            m = ws.cell(row=3, column=x, value=clo["max"] if colmap is clo_col else f"≥ {clo['max'] * course['studentThreshold']:g}")
            m.alignment, m.border, m.font = center, border, Font(italic=True, size=9)
            ws.column_dimensions[col(x)].width = 10
    for x, title, sub in ((total_col, f"Tổng ({course['scale']})", course["scale"]), (note_col, "Ghi chú", "")):
        ws.merge_cells(start_row=1, start_column=x, end_row=2, end_column=x)
        cell = ws.cell(row=1, column=x, value=title)
        cell.fill, cell.font, cell.alignment, cell.border = head_fill, head_font, center, border
        m = ws.cell(row=3, column=x, value=sub)
        m.alignment, m.border = center, border
    ws.column_dimensions[col(total_col)].width = 10
    ws.column_dimensions[col(note_col)].width = 30
    for i, w in enumerate([5, 12, 24, 8], start=1):
        ws.column_dimensions[col(i)].width = w
    ws.row_dimensions[2].height = 75

    first, last = 4, 3 + DATA_ROWS
    for r in range(first, last + 1):
        ws.cell(row=r, column=1, value=r - first + 1).alignment = center
        for x in range(1, last_col + 1):
            ws.cell(row=r, column=x).border = border
        for clo in clos:
            ids = [c["id"] for c in clo["criteria"]]
            rng = f"{col(crit_col[ids[0]])}{r}:{col(crit_col[ids[-1]])}{r}"
            # Chỉ tính điểm CLO khi đã chấm đủ mọi tiêu chí của CLO
            ws.cell(row=r, column=clo_col[clo["id"]], value=f'=IF(COUNT({rng})={len(ids)},ROUND(SUM({rng}),2),"")')
            cref = f"{col(clo_col[clo['id']])}{r}"
            ws.cell(row=r, column=pass_col[clo["id"]], value=f'=IF({cref}="","",IF({cref}>={clo["max"] * course["studentThreshold"]},"Đạt","Chưa đạt"))')
        crng = f"{col(min(clo_col.values()))}{r}:{col(max(clo_col.values()))}{r}"
        ws.cell(row=r, column=total_col, value=f'=IF(COUNT({crng})={len(clos)},ROUND(SUM({crng}),2),"")').font = Font(bold=True)
        for x in list(clo_col.values()) + list(pass_col.values()) + [total_col]:
            ws.cell(row=r, column=x).alignment = center

    for clo, c in criteria:
        letter = col(crit_col[c["id"]])
        dv = DataValidation(type="decimal", operator="between", formula1="0", formula2=str(c["max"]), allow_blank=True,
                            showErrorMessage=True, errorTitle="Điểm không hợp lệ",
                            error=f"Tiêu chí {c['id']}: nhập từ 0 đến {c['max']}")
        ws.add_data_validation(dv)
        dv.add(f"{letter}{first}:{letter}{last}")
    red = PatternFill("solid", fgColor="FDE2E1")
    green = PatternFill("solid", fgColor="E3F6E8")
    prng = f"{col(min(pass_col.values()))}{first}:{col(max(pass_col.values()))}{last}"
    ws.conditional_formatting.add(prng, CellIsRule(operator="equal", formula=['"Chưa đạt"'], fill=red, font=Font(color="B42318", bold=True)))
    ws.conditional_formatting.add(prng, CellIsRule(operator="equal", formula=['"Đạt"'], fill=green, font=Font(color="067647")))
    ws.freeze_panes = ws.cell(row=first, column=first_crit)

    # ----- Mức đạt CLO
    ws2 = wb.create_sheet("Mức đạt CLO")
    for i, w in enumerate([10, 50, 12, 14, 14, 14, 12, 14, 18], start=1):
        ws2.column_dimensions[col(i)].width = w
    ws2.cell(row=1, column=1, value=f"MỨC ĐẠT CLO - {course['name']} ({course['code']})").font = Font(bold=True, size=13, color=BLUE)
    ws2.cell(row=2, column=1, value=course["note"]).font = Font(italic=True)
    header(ws2, 4, ["CLO", "Nội dung", "Điểm tối đa", "Ngưỡng SV đạt", "Số SV có điểm", "Số SV đạt", "Tỷ lệ đạt", "Điểm TB", "Kết luận cấp lớp"])
    sheet = "'Bảng điểm'"
    for i, clo in enumerate(clos):
        r = 5 + i
        sc = f"{sheet}!{col(clo_col[clo['id']])}{first}:{col(clo_col[clo['id']])}{last}"
        pc = f"{sheet}!{col(pass_col[clo['id']])}{first}:{col(pass_col[clo['id']])}{last}"
        values = [clo["id"], clo["name"], clo["max"], clo["max"] * course["studentThreshold"], f"=COUNT({sc})", f'=COUNTIF({pc},"Đạt")',
                  f'=IF(E{r}=0,"",F{r}/E{r})', f'=IF(E{r}=0,"",ROUND(AVERAGE({sc}),2))',
                  f'=IF(E{r}=0,"",IF(G{r}>={course["classThreshold"]},"Đạt","Chưa đạt"))']
        for x, v in enumerate(values, start=1):
            cell = ws2.cell(row=r, column=x, value=v)
            cell.border, cell.alignment = border, (wrap if x == 2 else center)
        ws2.cell(row=r, column=7).number_format = "0.0%"
    concl = f"I5:I{4 + len(clos)}"
    ws2.conditional_formatting.add(concl, CellIsRule(operator="equal", formula=['"Chưa đạt"'], fill=red, font=Font(color="B42318", bold=True)))
    ws2.conditional_formatting.add(concl, CellIsRule(operator="equal", formula=['"Đạt"'], fill=green, font=Font(color="067647", bold=True)))

    r0 = 7 + len(clos)
    ws2.cell(row=r0 - 1, column=1, value="Điểm trung bình theo tiêu chí").font = Font(bold=True, color=DARK)
    header(ws2, r0, ["Mã", "Tiêu chí", "Điểm tối đa", "Phạm vi", "Số SV có điểm", "Điểm TB", "TB / tối đa"])
    for i, (clo, c) in enumerate(criteria):
        r = r0 + 1 + i
        rng = f"{sheet}!{col(crit_col[c['id']])}{first}:{col(crit_col[c['id']])}{last}"
        values = [c["id"], c["name"], c["max"], scope_text[c["scope"]], f"=COUNT({rng})",
                  f'=IF(E{r}=0,"",ROUND(AVERAGE({rng}),2))', f'=IF(E{r}=0,"",F{r}/C{r})']
        for x, v in enumerate(values, start=1):
            cell = ws2.cell(row=r, column=x, value=v)
            cell.border, cell.alignment = border, (wrap if x == 2 else center)
        ws2.cell(row=r, column=7).number_format = "0%"
    low = f"G{r0 + 1}:G{r0 + len(criteria)}"
    ws2.conditional_formatting.add(low, FormulaRule(formula=[f'AND(ISNUMBER(G{r0 + 1}),G{r0 + 1}<{course["studentThreshold"]})'], fill=red))

    wb.active = wb.sheetnames.index("Bảng điểm")
    wb.save(out)


# ---------------------------------------------------------------- HTML
def build_html(rubric: dict, out: Path):
    html = TEMPLATE.read_text(encoding="utf-8")
    data = json.dumps(rubric, ensure_ascii=False).replace("</", "<\\/")
    course = rubric["course"]
    html = (html.replace("__RUBRIC_JSON__", data)
                .replace("__COURSE_NAME__", course["name"])
                .replace("__COURSE_CODE__", course["code"]))
    out.write_text(html, encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description="Sinh rubric Excel và công cụ chấm điểm HTML từ rubric.json")
    parser.add_argument("course", help="Thư mục học phần")
    args = parser.parse_args()
    course_dir = ROOT_DIR / args.course.strip("/\\")
    src = course_dir / "rubric" / "rubric.json"
    if not src.is_file():
        print(f"[LỖI] Không có {src.relative_to(ROOT_DIR)}")
        return 1
    rubric = json.loads(src.read_text(encoding="utf-8"))
    errors = validate(rubric)
    if errors:
        for e in errors:
            print(f"[LỖI] {e}")
        return 1
    n = sum(len(clo["criteria"]) for clo in rubric["clos"])
    xlsx = course_dir / "rubric" / f"Rubric-{rubric['course']['code']}.xlsx"
    build_excel(rubric, xlsx)
    print(f"[+] Đã ghi {xlsx.relative_to(ROOT_DIR)} ({len(rubric['clos'])} CLO, {n} tiêu chí)")
    html = course_dir / "cham-diem.html"
    build_html(rubric, html)
    print(f"[+] Đã ghi {html.relative_to(ROOT_DIR)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
