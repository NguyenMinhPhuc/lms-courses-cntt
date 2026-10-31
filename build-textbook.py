"""
Ghép giáo trình nội bộ (Markdown) của một học phần thành một file Word (.docx) có trang bìa và mục lục.

Cách dùng:
    uv run --with python-docx python build-textbook.py lap-trinh-ung-dung-mobile
    uv run --with python-docx python build-textbook.py lap-trinh-ung-dung-mobile --srs

Giáo trình:  <hoc-phan>/giao-trinh/00-loi-noi-dau.md, chuong-01.md, ..., phu-luc-a-*.md  → giao-trinh/Giao-trinh-<hoc-phan>.docx
             (giao-trinh/bai-tap-huong-dan/bt-XX.md, nếu có, được chèn vào chương XX trước "Tài liệu tham khảo")
SRS (--srs): <hoc-phan>/srs/01-*.md, 02-*.md, ...  → srs/SRS-<srsProject>.docx
             (tên dự án lấy từ "srsProject" trong course.json)
Agile (--agile): <hoc-phan>/srs/agile/01-*.md, ...  → srs/agile-scrum.docx (phần nội dung khổ ngang cho bảng backlog)
                 và srs/agile/product-backlog.csv (Epic + backlog, nhập Jira/Trello/Notion)
Hồ sơ (--ho-so): <hoc-phan>/giao-trinh/ho-so/01-*.md, ...  → giao-trinh/ho-so/Ho-so-bien-soan-<hoc-phan>.docx
                 (đề cương chi tiết, tổng hợp góp ý, báo cáo quá trình biên soạn theo Thông tư 62/2026/TT-BGDĐT)
Watermark:   chữ mờ giữa mọi trang, mặc định "Công nghệ thông tin - LHU" ("watermark" trong course.json, "" để tắt)

Hỗ trợ phần Markdown dùng trong giáo trình: tiêu đề #..####, đoạn văn, **đậm**, *nghiêng*, `code`,
danh sách - và 1. (lồng 2 cấp), bảng |...|, khối ``` code ```, trích dẫn >,
hình ![Hình X.Y. Chú thích](đường/dẫn.png) trên một dòng riêng (đường dẫn tính từ thư mục chứa file .md).
Mục lục là trường TOC của Word: mở file trong Word rồi chọn Update Field (hoặc F9) để cập nhật số trang.
"""
import re
import sys
import json
import argparse
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Fix Windows console encoding for Vietnamese unicode characters
if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

ROOT_DIR = Path(__file__).resolve().parent
BODY_FONT = "Times New Roman"
HEADING_FONT = "Arial"
CODE_FONT = "Consolas"
LHU_BLUE = RGBColor(0x01, 0x3D, 0x88)
CODE_BG = "F1F3F5"
TABLE_HEADER_BG = "E7F1FF"
DEFAULT_WATERMARK = "Công nghệ thông tin - LHU"  # ghi đè bằng "watermark" trong course.json; "" để tắt


# ---------------------------------------------------------------- helpers
def set_cell_shading(cell, color: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def set_paragraph_shading(paragraph, color: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    paragraph._p.get_or_add_pPr().append(shd)


def set_paragraph_left_border(paragraph, color: str = "013D88"):
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for key, value in (("val", "single"), ("sz", "18"), ("space", "8"), ("color", color)):
        left.set(qn(f"w:{key}"), value)
    borders.append(left)
    paragraph._p.get_or_add_pPr().append(borders)


def add_field(paragraph, instruction: str, placeholder: str = ""):
    """Chèn một trường Word (TOC, PAGE...) vào paragraph."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, sep, text, end):
        run._r.append(element)


def add_inline(paragraph, text: str, size: float | None = None, color: RGBColor | None = None):
    """Ghi text có **đậm**, *nghiêng*, `code` vào paragraph."""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = CODE_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
            run.font.size = Pt((size or 13) - 2)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size and not (part.startswith("`") and part.endswith("`")):
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


# ---------------------------------------------------------------- document setup
def setup_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(13)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    for level, size, before in ((1, 20, 0), (2, 16, 14), (3, 14, 10), (4, 13, 8)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = HEADING_FONT
        rfonts = style.element.rPr.rFonts
        rfonts.set(qn("w:eastAsia"), HEADING_FONT)
        # Style tiêu đề mặc định trỏ tới font theme (Calibri Light); gỡ để Arial có hiệu lực
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = LHU_BLUE
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.5)


def add_cover(doc: Document, config: dict, course_dir: Path, kind: str = "GIÁO TRÌNH", title: str | None = None,
              subtitle_lines: list[str] | None = None):
    def centered(text: str, size: float, bold=False, color=None, space_after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = color
        return p

    centered("TRƯỜNG ĐẠI HỌC LẠC HỒNG", 14, bold=True)
    faculty = (config.get("faculty") or {}).get("name") or "Khoa Công nghệ Thông tin"
    centered(faculty.upper(), 14, bold=True, space_after=120)
    centered(kind, 22, bold=True, color=LHU_BLUE)
    centered((title or config.get("name", course_dir.name)).upper(), 26, bold=True, color=LHU_BLUE)
    centered(f"Mã học phần: {config.get('code', '')}", 14, space_after=24)
    lines = subtitle_lines if subtitle_lines is not None else config.get("textbookSubtitle", [])
    for line in lines:
        centered(line, 13)
    centered("", 13, space_after=120)
    author = config.get("textbookAuthor") or config.get("instructor")
    if author:
        centered(f"Biên soạn: {author}", 14, bold=True)
    centered("Bản thảo - Lưu hành nội bộ", 13)
    centered(f"Đồng Nai, {config.get('year', date.today().year)}", 13)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc: Document):
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("MỤC LỤC")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = LHU_BLUE
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u',
              "Nhấn chuột phải vào đây và chọn Update Field để tạo mục lục.")


def setup_header_footer(section, title: str):
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(title)
    run.font.size = Pt(10)
    run.italic = True
    run.font.color.rgb = RGBColor(0x6C, 0x75, 0x7D)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE", "1")


WATERMARK_SHAPETYPE = (
    '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" '
    'path="m@7,l@8,m@5,21600l@6,21600e"><v:formulas>'
    + "".join(f'<v:f eqn="{e}"/>' for e in [
        "sum #0 0 10800", "prod #0 2 1", "sum 21600 0 @1", "sum 0 0 @2", "sum 21600 0 @3",
        "if @0 @3 0", "if @0 21600 @1", "if @0 0 @2", "if @0 @4 21600", "mid @5 @6",
        "mid @8 @5", "mid @7 @8", "mid @6 @7", "sum @6 0 @5"])
    + '</v:formulas><v:path textpathok="t" o:connecttype="custom" '
    'o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>'
    '<v:textpath on="t" fitshape="t"/><v:handles><v:h position="#0,bottomRight" xrange="6629,14971"/>'
    '</v:handles><o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>'
)


def add_watermark(section, text: str):
    """Watermark chữ nghiêng mờ giữa trang (WordArt VML trong header, kiểu Design > Watermark của Word)."""
    from docx.oxml import parse_xml
    from xml.sax.saxutils import quoteattr
    paragraph = section.header.paragraphs[0]
    run = parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
        '<w:pict>' + WATERMARK_SHAPETYPE +
        '<v:shape id="LhuWatermark" o:spid="_x0000_s2049" type="#_x0000_t136" '
        'style="position:absolute;margin-left:0;margin-top:0;width:470pt;height:70pt;rotation:315;'
        'z-index:-251654144;mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
        'mso-position-vertical:center;mso-position-vertical-relative:margin" '
        'o:allowincell="f" fillcolor="#c8d3e6" stroked="f">'
        '<v:fill opacity=".35"/>'
        f'<v:textpath style="font-family:&quot;Arial&quot;;font-size:1pt" string={quoteattr(text)}/>'
        '</v:shape></w:pict></w:r>'
    )
    paragraph._p.append(run)


# ---------------------------------------------------------------- markdown → docx
def add_table(doc: Document, rows: list[list[str]]):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Chia độ rộng cột theo lượng chữ: cột mô tả dài được rộng hơn cột mã, cột ưu tiên
    section = doc.sections[-1]
    usable = section.page_width - section.left_margin - section.right_margin
    minimums, weights = [], []
    for c in range(len(header)):
        texts = [re.sub(r"[`*]", "", r[c]) for r in [header] + body if c < len(r)]
        lengths = [len(t) for t in texts]
        longest_word = max((len(w) for t in texts for w in t.split()), default=1)
        # mỗi cột tối thiểu đủ chứa từ dài nhất (khoảng 0,21 cm mỗi ký tự ở cỡ 11pt, cộng lề ô)
        minimums.append(Cm(min(longest_word * 0.21 + 0.5, 6)))
        weights.append(min(max(sum(lengths) / len(lengths), max(lengths) * 0.35), 60))
    spare = max(usable - sum(minimums), 0)
    widths = [int(m + spare * w / sum(weights)) for m, w in zip(minimums, weights)]
    if sum(widths) > usable:  # quá nhiều cột: thu nhỏ đều
        widths = [int(w * usable / sum(widths)) for w in widths]

    for r, cells in enumerate([header] + body):
        for c in range(len(header)):
            cell = table.cell(r, c)
            cell.width = widths[c]
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            add_inline(cell.paragraphs[0], cells[c] if c < len(cells) else "", size=11)
            if r == 0:
                set_cell_shading(cell, TABLE_HEADER_BG)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(doc: Document, lines: list[str]):
    for i, line in enumerate(lines or [""]):
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(4 if i == 0 else 0)
        fmt.space_after = Pt(4 if i == len(lines) - 1 else 0)
        fmt.line_spacing = 1.0
        fmt.left_indent = Cm(0.3)
        fmt.keep_with_next = i < len(lines) - 1 and len(lines) <= 25
        set_paragraph_shading(p, CODE_BG)
        run = p.add_run(line.replace("\t", "    ") or " ")
        run.font.name = CODE_FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
        run.font.size = Pt(9.5)


def add_list_item(doc: Document, text: str, level: int, number: str | None):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    indent = Cm(0.8 + 0.8 * level)
    fmt.left_indent = indent
    fmt.first_line_indent = -Cm(0.6)
    fmt.space_after = Pt(3)
    bullet = f"{number}\t" if number else ("•\t" if level == 0 else "–\t")
    p.add_run(bullet)
    fmt.tab_stops.add_tab_stop(indent)
    add_inline(p, text)


def add_figure(doc: Document, caption: str, path: Path):
    """Hình căn giữa, rộng tối đa 15 cm, cao tối đa 9 cm, chú thích in nghiêng bên dưới."""
    if path.is_file():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        pic = p.add_run().add_picture(str(path), width=Cm(15))
        if pic.height > Cm(9):  # hình khổ dọc: giữ tỷ lệ, không chiếm cả trang
            pic.width, pic.height = int(pic.width * Cm(9) / pic.height), Cm(9)
    else:
        print(f"  [!] Không tìm thấy hình: {path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(caption)
    run.italic = True
    run.font.size = Pt(11)


def render_markdown(doc: Document, text: str, first_chapter: bool, base: Path | None = None):
    lines = text.replace("\r\n", "\n").split("\n")
    i, para = 0, []

    def flush():
        if para:
            add_inline(doc.add_paragraph(), " ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code_block(doc, block)
        elif m := re.match(r"^!\[(.*)\]\((.+)\)\s*$", stripped):
            flush()
            add_figure(doc, m.group(1), (base or ROOT_DIR) / m.group(2))
        elif m := re.match(r"^(#{1,4})\s+(.*)$", line):
            flush()
            level = len(m.group(1))
            heading = doc.add_heading(m.group(2).strip(), level=level)
            if level == 1 and not first_chapter:
                heading.paragraph_format.page_break_before = True
        elif stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        elif m := re.match(r"^(\s*)[-*]\s+(.*)$", line):
            flush()
            add_list_item(doc, m.group(2), len(m.group(1)) // 2, None)
        elif m := re.match(r"^(\s*)(\d+)\.\s+(.*)$", line):
            flush()
            add_list_item(doc, m.group(3), len(m.group(1)) // 2, f"{m.group(2)}.")
        elif stripped.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            set_paragraph_left_border(p)
            add_inline(p, " ".join(quote))
            continue
        elif not stripped:
            flush()
        else:
            para.append(line)
        i += 1
    flush()


def chapter_text(f: Path) -> str:
    """Nội dung chương, kèm bài tập hướng dẫn bai-tap-huong-dan/bt-XX.md (nếu có) chèn trước Tài liệu tham khảo."""
    text = f.read_text(encoding="utf-8")
    m = re.match(r"chuong-(\d+)\.md$", f.name)
    extra = f.parent / "bai-tap-huong-dan" / f"bt-{m.group(1)}.md" if m else None
    if extra and extra.is_file():
        block = extra.read_text(encoding="utf-8").strip() + "\n\n"
        marker = "\n## Tài liệu tham khảo"
        text = text.replace(marker, "\n" + block + marker.lstrip("\n"), 1) if marker in text else text + "\n\n" + block
    return text


def render_document(files: list[Path], out: Path, header: str, cover, watermark: str | None = None,
                    landscape: bool = False) -> Path:
    doc = Document()
    setup_styles(doc)
    cover(doc)
    add_toc(doc)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:  # bảng nhiều cột (ví dụ product backlog) cần trang ngang
        body.orientation = WD_ORIENT.LANDSCAPE
        body.page_width, body.page_height = body.page_height, body.page_width
    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False
    setup_header_footer(body, header)

    for index, f in enumerate(files):
        render_markdown(doc, chapter_text(f), first_chapter=index == 0, base=f.parent)

    if watermark:
        for section in doc.sections:
            add_watermark(section, watermark)

    doc.save(out)
    return out


def load_config(course_dir: Path) -> dict:
    config_file = course_dir / "course.json"
    return json.loads(config_file.read_text(encoding="utf-8")) if config_file.is_file() else {}


def build(course_dir: Path) -> Path:
    src = course_dir / "giao-trinh"
    files = sorted(src.glob("00-*.md")) + sorted(src.glob("chuong-*.md")) + sorted(src.glob("phu-luc-*.md"))
    if not files:
        raise FileNotFoundError(f"Không có file Markdown trong {src.relative_to(ROOT_DIR)}")
    config = load_config(course_dir)
    # "textbookKind" đổi tên loại tài liệu, ví dụ "Tài liệu học tập" thay cho "Giáo trình"
    kind = config.get("textbookKind", "Giáo trình")
    return render_document(files, src / f"Giao-trinh-{course_dir.name}.docx",
                           f"{kind} {config.get('name', course_dir.name)}",
                           lambda doc: add_cover(doc, config, course_dir, kind=kind.upper()),
                           watermark=config.get("watermark", DEFAULT_WATERMARK))


def read_md_table(path: Path, first_header: str) -> list[dict]:
    """Đọc bảng Markdown đầu tiên có cột đầu là first_header, trả về danh sách dòng dạng dict."""
    lines = path.read_text(encoding="utf-8").splitlines()
    start = next((i for i, l in enumerate(lines) if re.match(rf"^\|\s*{re.escape(first_header)}\s*\|", l)), None)
    if start is None:
        raise FileNotFoundError(f"Không thấy bảng có cột '{first_header}' trong {path.name}")
    split = lambda l: [c.strip() for c in l.strip().strip("|").split("|")]
    header, rows = split(lines[start]), []
    for line in lines[start + 2:]:
        if not line.strip().startswith("|"):
            break
        cells = split(line)
        if len(cells) != len(header):
            raise ValueError(f"{path.name}: dòng '{cells[0]}' có {len(cells)} cột, cần {len(header)}")
        rows.append(dict(zip(header, cells)))
    return rows


JIRA_PRIORITY = {"Must-have": "Highest", "Should-have": "High", "Could-have": "Medium", "Won't-have": "Lowest"}


def export_backlog_csv(src: Path, out: Path) -> int:
    """Xuất Epic + Product Backlog ra CSV nhập được vào Jira (Issue Id/Parent Id), Trello, Notion.

    Một dòng mỗi Epic (Issue Type = Epic) và mỗi story; story liên kết Epic qua Parent Id.
    Priority theo thang Jira, MoSCoW giữ ở cột riêng và Labels. UTF-8 có BOM để Excel đọc đúng tiếng Việt.
    """
    import csv

    epics = read_md_table(next(src.glob("[0-9][0-9]-epic*.md")), "ID")
    stories = read_md_table(next(src.glob("[0-9][0-9]-product-backlog*.md")), "ID")
    columns = ["Issue Id", "Issue Type", "Summary", "Epic", "Parent Id", "Description", "Acceptance Criteria",
               "Priority", "MoSCoW", "Story Points", "Sprint", "Labels"]
    plain = lambda s: re.sub(r"\*\*|`", "", s)

    epic_names = {e["ID"]: e["Epic"] for e in epics}
    rows = []
    for e in epics:
        sp = e["Tổng SP"].strip("()")
        rows.append({
            "Issue Id": e["ID"], "Issue Type": "Epic", "Summary": f"[{e['ID']}] {e['Epic']}", "Epic": e["Epic"],
            "Parent Id": "", "Description": f"{e['Mô tả']}\nModule SRS: {e['Module SRS']}\nSprint: {e['Sprint']}\nTổng story point: {sp}",
            "Acceptance Criteria": "", "Priority": JIRA_PRIORITY.get(e["Ưu tiên"], ""), "MoSCoW": e["Ưu tiên"],
            "Story Points": "", "Sprint": "", "Labels": e["Ưu tiên"].lower().replace("'", ""),
        })
    for s in stories:
        story = plain(s["User Story"])
        want = re.search(r"I want (.+?), so that", story)
        title = (want.group(1) if want else story)[:200]
        criteria = "\n".join(f"- {part.strip()}" for part in plain(s["Acceptance Criteria"]).split(" · ") if part.strip())
        epic_id = s["Epic"].split()[0]
        if epic_id not in epic_names:
            raise ValueError(f"{s['ID']}: Epic {epic_id} không có trong danh sách Epic")
        rows.append({
            "Issue Id": s["ID"], "Issue Type": "Task" if s["ID"].startswith("TS-") else "Story",
            "Summary": f"[{s['ID']}] {title[:1].upper()}{title[1:]}", "Epic": epic_names[epic_id], "Parent Id": epic_id,
            "Description": f"{story}\n\nAcceptance Criteria:\n{criteria}", "Acceptance Criteria": criteria,
            "Priority": JIRA_PRIORITY.get(s["Priority"], ""), "MoSCoW": s["Priority"], "Story Points": s["Story Points"],
            "Sprint": s["Target Sprint"], "Labels": s["Priority"].lower().replace("'", ""),
        })

    with out.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def build_agile(course_dir: Path) -> Path:
    src = course_dir / "srs" / "agile"
    files = sorted(src.glob("[0-9][0-9]-*.md"))
    if not files:
        raise FileNotFoundError(f"Không có file Markdown trong {src.relative_to(ROOT_DIR)}")
    csv_out = src / "product-backlog.csv"
    count = export_backlog_csv(src, csv_out)
    print(f"[+] Đã ghi {csv_out.relative_to(ROOT_DIR)} ({count} dòng: Epic + story)")
    config = load_config(course_dir)
    project = config.get("srsProject", course_dir.name)
    subtitle = ["Epic · User Story · Acceptance Criteria · Story Point · MoSCoW · Sprint Planning",
                f"Chuyển đổi từ tài liệu SRS dự án {project}"]
    return render_document(files, course_dir / "srs" / "agile-scrum.docx",
                           f"Agile Project Plan {project} - {config.get('name', course_dir.name)}",
                           lambda doc: add_cover(doc, config, course_dir, kind="KẾ HOẠCH DỰ ÁN AGILE & PRODUCT BACKLOG",
                                                 title=f"Dự án {project}", subtitle_lines=subtitle),
                           watermark=config.get("watermark", DEFAULT_WATERMARK), landscape=True)


def build_srs(course_dir: Path) -> Path:
    src = course_dir / "srs"
    files = sorted(src.glob("[0-9][0-9]-*.md"))
    if not files:
        raise FileNotFoundError(f"Không có file Markdown trong {src.relative_to(ROOT_DIR)}")
    config = load_config(course_dir)
    project = config.get("srsProject", course_dir.name)
    subtitle = [f"Dự án mẫu của học phần {config.get('name', course_dir.name)}",
                "Phụ lục của giáo trình"]
    out = src / f"SRS-{project.replace(' ', '-')}.docx"
    return render_document(files, out, f"SRS {project} - {config.get('name', course_dir.name)}",
                           lambda doc: add_cover(doc, config, course_dir, kind="ĐẶC TẢ YÊU CẦU PHẦN MỀM (SRS)",
                                                 title=f"Dự án {project}", subtitle_lines=subtitle),
                           watermark=config.get("watermark", DEFAULT_WATERMARK))


def build_ho_so(course_dir: Path) -> Path:
    """Hồ sơ biên soạn giáo trình (đề cương chi tiết, tổng hợp góp ý, báo cáo quá trình biên soạn...)
    theo Thông tư 62/2026/TT-BGDĐT. Xem skill giao-trinh-builder."""
    src = course_dir / "giao-trinh" / "ho-so"
    files = sorted(src.glob("[0-9][0-9]-*.md"))
    if not files:
        raise FileNotFoundError(f"Không có file Markdown trong {src.relative_to(ROOT_DIR)}")
    config = load_config(course_dir)
    name = config.get("name", course_dir.name)
    kind = config.get("textbookKind", "Giáo trình")
    subtitle = [f"{kind} {name}", "Theo Thông tư 62/2026/TT-BGDĐT"]
    return render_document(files, src / f"Ho-so-bien-soan-{course_dir.name}.docx", f"Hồ sơ biên soạn - {name}",
                           lambda doc: add_cover(doc, config, course_dir, kind="HỒ SƠ BIÊN SOẠN GIÁO TRÌNH",
                                                 title=name, subtitle_lines=subtitle),
                           watermark=config.get("watermark", DEFAULT_WATERMARK))


def main():
    parser = argparse.ArgumentParser(description="Ghép giáo trình Markdown thành file Word")
    parser.add_argument("course", help="Thư mục học phần")
    parser.add_argument("--srs", action="store_true", help="Ghép tài liệu SRS trong <hoc-phan>/srs/ thay cho giáo trình")
    parser.add_argument("--agile", action="store_true", help="Ghép kế hoạch Agile/Scrum trong <hoc-phan>/srs/agile/")
    parser.add_argument("--ho-so", action="store_true", help="Ghép hồ sơ biên soạn trong <hoc-phan>/giao-trinh/ho-so/")
    args = parser.parse_args()
    course_dir = ROOT_DIR / args.course.strip("/\\")
    try:
        out = (build_agile(course_dir) if args.agile else build_srs(course_dir) if args.srs
               else build_ho_so(course_dir) if args.ho_so else build(course_dir))
    except FileNotFoundError as e:
        print(f"[LỖI] {e}")
        return 1
    print(f"[+] Đã ghi {out.relative_to(ROOT_DIR)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
