"""
Sinh bộ hồ sơ xây dựng chương trình đào tạo (BM01–BM09, file PLO, đề cương BM05) từ một bộ dữ liệu duy nhất.

Thư mục một CTĐT:
    chuong-trinh-dao-tao/<nganh>-<nam>/
    ├── du-lieu/                  dữ liệu nhập (sửa ở đây)
    │   ├── thong-tin.json        ngành, niên khóa, ngày và số văn bản, Hội đồng, căn cứ, người góp ý
    │   ├── ke-hoach.csv          học kỳ, mã, tên, tín chỉ, LT/TH/BT, khối, nhóm tự chọn, học phần trước
    │   ├── chuan-dau-ra.json     PEO, PLO, PI, ma trận PEO–PLO, vị trí việc làm
    │   ├── doi-sanh.csv          (tùy chọn) học phần của chương trình đối sánh ↔ mã LHU tương đương
    │   ├── ke-hoach-khoa-truoc.csv (tùy chọn) để lập bảng thay đổi (BM06) và đánh giá khối lượng (BM08)
    │   └── de-cuong/<mã>.md      đề cương từng học phần (CLO, kế hoạch, ma trận LLO, đánh giá, tài liệu)
    ├── ho-so/                    file Word/Excel sinh tự động (không sửa tay)
    └── ra-soat/                  báo cáo rà soát

Cách dùng:
    uv run --with python-docx --with openpyxl python build-ctdt.py init chuong-trinh-dao-tao/cntt-2027
    uv run --with python-docx --with openpyxl python build-ctdt.py init chuong-trinh-dao-tao/cntt-2026 --tu-ho-so "templates/xay-dung-ctdt-Dai hoc chinh quy"
    uv run --with python-docx --with openpyxl python build-ctdt.py build chuong-trinh-dao-tao/cntt-2026
    uv run --with python-docx --with openpyxl python build-ctdt.py build chuong-trinh-dao-tao/cntt-2026 --chi BM03,BM05

`init` tạo bộ dữ liệu mẫu (hoặc nhập từ một bộ hồ sơ Word/Excel đã có bằng --tu-ho-so).
`build` sinh hồ sơ vào ho-so/, lấy đầu trang, logo, khổ giấy từ file mẫu trong thư mục mẫu
(mặc định templates/xay-dung-ctdt-Dai hoc chinh quy, đổi bằng "thu_muc_mau" trong thong-tin.json),
rồi chạy validate-ctdt.py trên ho-so/ và ghi ra-soat/ra-soat-tu-dong.md.
Chỗ chưa có dữ liệu được để {{...}}; ý kiến góp ý, phát biểu trong họp, chữ ký không bao giờ được sinh tự động.
"""
import re
import sys
import csv
import json
import argparse
import unicodedata
import importlib.util
from fnmatch import fnmatch
from datetime import date
from pathlib import Path
from collections import defaultdict

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm
from docx.table import Table
from docx.text.paragraph import Paragraph
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

ROOT_DIR = Path(__file__).resolve().parent
DEFAULT_MAU = ROOT_DIR / "templates" / "xay-dung-ctdt-Dai hoc chinh quy"
FONT = "Times New Roman"
GIO_MOI_TIN_CHI = 50
PH = "{{%s}}"

CSV_COLS = ["hoc_ky", "ma", "ten", "ten_en", "tc", "lt", "th", "bt", "khoi", "nhom_tu_chon", "hoc_truoc", "ghi_chu"]
# Khối kiến thức dùng trong cột "khoi" của ke-hoach.csv, theo nhóm của bảng "Tỷ lệ các khối kiến thức" (BM09.2)
KHOI = {
    "Lý luận chính trị": "Khối kiến thức giáo dục đại cương",
    "Ngoại ngữ": "Khối kiến thức giáo dục đại cương",
    "Đại cương khác": "Khối kiến thức giáo dục đại cương",
    "Cơ sở ngành": "Khối kiến thức giáo dục chuyên nghiệp",
    "Chuyên ngành": "Khối kiến thức giáo dục chuyên nghiệp",
    "Chuyên nghiệp khác": "Khối kiến thức giáo dục chuyên nghiệp",
    "Thực tập": "Tốt nghiệp",
    "Tốt nghiệp": "Tốt nghiệp",
    "Giáo dục thể chất": None,          # không tính vào tỷ lệ khối kiến thức
    "Giáo dục quốc phòng": None,
}
KHOI_TEN_DE_CUONG = {"Lý luận chính trị": "Đại cương", "Ngoại ngữ": "Đại cương", "Đại cương khác": "Đại cương",
                     "Chuyên nghiệp khác": "Chuyên ngành", "Thực tập": "Thực tập, tốt nghiệp", "Tốt nghiệp": "Thực tập, tốt nghiệp"}
# Thứ tự lập hồ sơ: ngày của văn bản sau không được sớm hơn văn bản trước
THU_TU_NGAY = ["BM08", "BM04", "BM01", "BM02", "BM03", "BM05", "BM06", "BM09", "BM07"]

# Mẫu khung: (mẫu tên file trong thư mục mẫu, regex tiêu đề). Giữ phần thân trước tiêu đề (đầu trang, logo), xóa phần còn lại.
# Regex None: xóa toàn bộ thân (logo nằm ở header của section).
FORMS = {
    "BM01.1": ("BM1.1*/*.docx", r"^Ý KIẾN CÁC BÊN"),
    "BM01.2": ("BM1.2*/*.docx", r"^PHIẾU GÓP Ý"),
    "BM02": ("BM02*.docx", r"^CHUẨN ĐẦU RA$"),
    "BM04": ("BM04*.docx", r"^BẢNG SO SÁNH"),
    "BM05": ("BM05*/HK3-111166*.docx", r"^ĐỀ CƯƠNG HỌC PHẦN"),
    "BM05a": ("02 mau*/*BM05a*.docx", None),
    "BM05b": ("02 mau*/*BM05b*.docx", None),
    "BM06": ("BM06*.docx", r"^BIÊN BẢN"),
    "BM07": ("BM07*.docx", r"^TỜ TRÌNH"),
    "BM08": ("BM08*.docx", r"^BẢNG ĐÁNH GIÁ"),
    "BM09": ("BM09*.docx", r"^MỤC LỤC$"),
    "DH": ("*điều hành*.docx", r"^CHƯƠNG TRÌNH ĐÀO TẠO TÍN CHỈ"),
    "TT": ("*Tóm tắt học phần*.docx", r"^TÓM TẮT HỌC PHẦN"),
}


def nfc(s):
    return unicodedata.normalize("NFC", str(s or ""))


def safe_name(s):
    return re.sub(r'[<>:"/\\|?*]+', "", nfc(s)).strip()


def fnum(v):
    if v is None or str(v).strip() == "":
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def g(v):
    """In số gọn: 3.0 -> '3', None -> ''."""
    if v is None:
        return ""
    return f"{v:g}" if isinstance(v, float) else str(v)


def parse_date(s):
    m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})$", str(s or "").strip())
    return date(int(m[1]), int(m[2]), int(m[3])) if m else None


def ngay_dai(iso, noi="Đồng Nai", nam=None):
    """'2026-04-07' -> 'Đồng Nai, ngày 07 tháng 4 năm 2026' (ngày < 10 và tháng 1, 2 có số 0 phía trước)."""
    d = parse_date(iso)
    if not d:
        return f"{noi}, ngày … tháng … năm {nam or '…'}"
    month = f"{d.month:02d}" if d.month < 3 else str(d.month)
    return f"{noi}, ngày {d.day:02d} tháng {month} năm {d.year}"


def ngay_ngan(iso):
    d = parse_date(iso)
    return f"{d.day}/{d.month}/{d.year}" if d else PH % "ngày"


def load_validator():
    spec = importlib.util.spec_from_file_location("validate_ctdt", ROOT_DIR / "validate-ctdt.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ================================================================ dữ liệu

def read_plan(path):
    rows = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for i, r in enumerate(csv.DictReader(f)):
            r = {k: (nfc(v).strip() if v is not None else "") for k, v in r.items()}
            if not r.get("ma"):
                continue
            c = {"hoc_ky": int(fnum(r["hoc_ky"]) or 0), "ma": r["ma"], "ten": r["ten"], "ten_en": r.get("ten_en", ""),
                 "tc": fnum(r["tc"]), "lt": fnum(r["lt"]) or 0.0, "th": fnum(r["th"]) or 0.0, "bt": fnum(r["bt"]) or 0.0,
                 "khoi": r.get("khoi", ""), "nhom": r.get("nhom_tu_chon", ""), "ghi_chu": r.get("ghi_chu", ""),
                 "hoc_truoc": [], "order": i}
            for m in re.finditer(r"(\d{6})\s*(\*{1,3})?", r.get("hoc_truoc", "")):
                c["hoc_truoc"].append((m[1], m[2] or "**"))
            c["tiet"] = (c["tc"] or 0) * GIO_MOI_TIN_CHI
            rows.append(c)
    rows.sort(key=lambda c: (c["hoc_ky"], c["order"]))
    return rows


def write_plan(path, courses):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(CSV_COLS)
        for c in courses:
            w.writerow([c["hoc_ky"], c["ma"], c["ten"], c.get("ten_en", ""), g(c["tc"]), g(c["lt"]), g(c["th"]), g(c["bt"]),
                        c.get("khoi", ""), c.get("nhom", ""), "; ".join(f"{a}{k}" for a, k in c.get("hoc_truoc", [])),
                        c.get("ghi_chu", "")])


def parse_md(path):
    """Đề cương Markdown: frontmatter 'khóa: giá trị' giữa hai dòng ---, rồi các mục '## Tên mục'."""
    text = nfc(path.read_text(encoding="utf-8"))
    meta, sections = {}, []
    if text.startswith("---"):
        _, fm, text = text.split("---", 2)
        for line in fm.splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip()] = v.strip()
    cur = None
    for line in text.splitlines():
        if line.startswith("## "):
            cur = [line[3:].strip(), []]
            sections.append(cur)
        elif cur is not None:
            cur[1].append(line)
    return {"meta": meta, "sections": sections, "path": path}


def md_tables(lines):
    """Các bảng Markdown trong một mục -> list[list[list[str]]] (bỏ hàng phân cách |---|)."""
    tables, cur = [], None
    for line in lines + [""]:
        s = line.strip()
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            if cur is None:
                cur = []
                tables.append(cur)
            cur.append(cells)
        else:
            cur = None
    return tables


def section(syl, key):
    for name, lines in syl["sections"]:
        if nfc(name).casefold().startswith(key.casefold()):
            return lines
    return None


def clo_pi(syl):
    """{số CLO: (nội dung, tập PI)} từ bảng của mục Chuẩn đầu ra."""
    out = {}
    lines = section(syl, "Chuẩn đầu ra") or []
    for t in md_tables(lines):
        for r in t:
            m = re.fullmatch(r"CLO\s*(\d+)", r[0]) if r else None
            if m and int(m[1]) not in out:
                out[int(m[1])] = (r[1] if len(r) > 1 else "", set(re.findall(r"PI\s*(\d+\.\d+)", " ".join(r))))
    return out


class Data:
    def __init__(self, folder):
        self.dir = folder
        d = folder / "du-lieu"
        self.info = json.loads((d / "thong-tin.json").read_text(encoding="utf-8"))
        self.cdr = json.loads((d / "chuan-dau-ra.json").read_text(encoding="utf-8"))
        self.courses = read_plan(d / "ke-hoach.csv")
        self.by = {c["ma"]: c for c in self.courses}
        prev = d / "ke-hoach-khoa-truoc.csv"
        self.prev = read_plan(prev) if prev.exists() else None
        ds = d / "doi-sanh.csv"
        self.doisanh = list(csv.DictReader(open(ds, encoding="utf-8-sig"))) if ds.exists() else []
        self.syl = {}
        for p in sorted((d / "de-cuong").glob("*.md")):
            s = parse_md(p)
            code = s["meta"].get("ma") or re.sub(r"\D", "", p.stem)[:6]
            self.syl[code] = s
        self.mau = Path(self.info.get("thu_muc_mau") or DEFAULT_MAU)
        if not self.mau.is_absolute():
            self.mau = ROOT_DIR / self.mau
        self.warnings = []

    # ---- tiện ích
    def i(self, key, default=None):
        v = self.info.get(key)
        return v if v not in (None, "") else (default if default is not None else PH % key)

    def fpart(self, key, default=""):
        """Giá trị dùng trong tên file: bỏ dấu {{ }} của chỗ chưa điền."""
        return safe_name(re.sub(r"[{}]", "", str(self.info.get(key) or default)))

    def ngay(self, form):
        return self.info.get("ngay", {}).get(form)

    def so(self, form):
        return self.info.get("so_van_ban", {}).get(form) or PH % f"số {form}"

    def nganh(self):
        return self.i("nganh")

    def semesters(self):
        """[(hk, [mục])], mục = ("hp", c) | ("nhom", k, [c...]). Thứ tự giữ như ke-hoach.csv."""
        out = defaultdict(list)
        seen = {}
        for c in self.courses:
            if c["nhom"]:
                key = c["nhom"].split(":")[0]
                if key not in seen:
                    k = int(c["nhom"].split(":")[1]) if ":" in c["nhom"] else 1
                    seen[key] = ["nhom", k, []]
                    out[c["hoc_ky"]].append(seen[key])
                seen[key][2].append(c)
            else:
                out[c["hoc_ky"]].append(["hp", c])
        return sorted(out.items())

    def sem_total(self, items):
        tot = {"tc": 0.0, "lt": 0.0, "th": 0.0, "bt": 0.0, "tiet": 0.0}
        for it in items:
            for c in ([it[1]] if it[0] == "hp" else it[2][:it[1]]):
                for k in tot:
                    tot[k] += c[k] or 0
        return tot

    def total(self):
        tot = {"tc": 0.0, "lt": 0.0, "th": 0.0, "bt": 0.0, "tiet": 0.0}
        for _, items in self.semesters():
            for k, v in self.sem_total(items).items():
                tot[k] += v
        return tot

    def prereq_text(self, c, with_name=False):
        if not c["hoc_truoc"]:
            return ""
        parts = []
        for code, kind in c["hoc_truoc"]:
            if with_name:
                p = self.by.get(code)
                label = {"*": "tiên quyết", "**": "học trước", "***": "song hành"}[kind]
                parts.append(f"{code} – {p['ten'] if p else PH % 'tên'} ({label})")
            else:
                parts.append(f"{code}{kind}")
        return "; ".join(parts) if with_name else " / ".join(parts)

    def map_clo_pi(self):
        """{mã: {CLO: tập PI}} lấy từ đề cương (nguồn của ma trận CLO–PI)."""
        return {code: {k: v[1] for k, v in clo_pi(s).items()} for code, s in self.syl.items()}

    def check_dates(self):
        prev = None
        for f in THU_TU_NGAY:
            d = parse_date(self.ngay(f))
            if d and prev and d < prev[1]:
                self.warnings.append(f"Ngày {f} ({d:%d/%m/%Y}) sớm hơn {prev[0]} ({prev[1]:%d/%m/%Y}): sai thứ tự lập hồ sơ")
            if d:
                prev = (f, d)
        hop = parse_date(self.ngay("BM06"))
        for cc in self.info.get("can_cu", []):
            d = parse_date(cc.get("ngay"))
            if hop and d and d > hop:
                self.warnings.append(f"Căn cứ ban hành sau ngày họp BM06: {cc.get('van_ban', '')[:80]} ({d:%d/%m/%Y})")


# ================================================================ docx

def body_elements(doc):
    return [e for e in doc.element.body.iterchildren() if not e.tag.endswith('}sectPr')]


def el_text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def find_template(data, form):
    pattern, _ = FORMS[form]
    pat = nfc(pattern).casefold()
    for p in sorted(data.mau.rglob("*.docx")):
        rel = nfc(p.relative_to(data.mau).as_posix()).casefold()
        if not p.name.startswith("~$") and fnmatch(rel, pat):
            return p
    return None


def save_docx(doc, path):
    """Lưu docx, bỏ font nhúng thừa từ file mẫu (Tahoma… ~900 KB/file) vì hồ sơ dùng Times New Roman."""
    import io
    import zipfile
    buf = io.BytesIO()
    doc.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            if item.filename.startswith("word/fonts/"):
                continue
            data = src.read(item.filename)
            if item.filename == "word/fontTable.xml":
                data = re.sub(rb"<w:embed(Regular|Bold|Italic|BoldItalic)[^>]*/>", b"", data)
            elif item.filename == "word/_rels/fontTable.xml.rels":
                data = re.sub(rb"<Relationship[^>]*/font\"[^>]*/>", b"", data)
                data = re.sub(rb'<Relationship[^>]*Type="[^"]*/font"[^>]*/>', b"", data)
            elif item.filename == "word/settings.xml":
                data = re.sub(rb"<w:(embedTrueTypeFonts|saveSubsetFonts)[^>]*/>", b"", data)
            elif item.filename == "[Content_Types].xml":
                data = re.sub(rb'<Override PartName="/word/fonts/[^"]*"[^>]*/>', b"", data)
            z.writestr(item, data)
    Path(path).write_bytes(out.getvalue())


def set_run_font(run, size=13, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(a), FONT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def replace_par_text(p, new):
    runs = p.runs
    if not runs:
        p.add_run(new)
        return
    runs[0].text = new
    for r in runs[1:]:
        r.text = ""


def iter_paragraphs(els, doc):
    for el in els:
        if el.tag.endswith('}p'):
            yield Paragraph(el, doc)
        elif el.tag.endswith('}tbl'):
            for p in el.iter(qn("w:p")):
                yield Paragraph(p, doc)


def base_doc(data, form, replacements=()):
    """Mở file mẫu, giữ phần đầu trang (trước tiêu đề), xóa phần còn lại; thay ngày, số văn bản ở phần giữ lại."""
    tpl = find_template(data, form)
    if not tpl:
        data.warnings.append(f"{form}: không tìm thấy file mẫu trong {data.mau}, dùng khung trống")
        doc = Document()
        s = doc.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3), Cm(2), Cm(2), Cm(2)
        return doc, False
    doc = Document(tpl)
    els = body_elements(doc)
    regex = FORMS[form][1]
    cut = 0
    if regex:
        cut = next((i for i, e in enumerate(els) if re.search(regex, nfc(el_text(e)).strip())), len(els))
    for e in els[cut:]:
        doc.element.body.remove(e)
    for p in iter_paragraphs(els[:cut], doc):
        t = nfc(p.text).strip()
        for pat, new in replacements:
            if re.search(pat, t):
                replace_par_text(p, new)
                break
    return doc, cut > 0


def admin_header(doc, data, form, kept):
    """Đầu văn bản hành chính khi mẫu không có (hoặc không giữ được) khối đầu trang."""
    if kept:
        return
    t = doc.add_table(rows=2, cols=2)
    vals = [["TRƯỜNG ĐẠI HỌC LẠC HỒNG\n" + data.i("don_vi", "Khoa Công nghệ thông tin").upper(),
             "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập – Tự do – Hạnh phúc"],
            [f"Số: {data.so(form)}" if form in ("BM06", "BM07") else "", ngay_dai(data.ngay(form), data.i("noi_ky", "Đồng Nai"), data.i("nam"))]]
    for i, row in enumerate(vals):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(v)
            set_run_font(r, 12, bold=(i == 0), italic=(i == 1 and j == 1))
    doc.add_paragraph()


def header_repl(data, form):
    noi = data.i("noi_ky", "Đồng Nai")
    reps = [(r"^\S.*,\s*ngày\b.*\bnăm\b", ngay_dai(data.ngay(form), noi, data.i("nam")))]
    if form in ("BM06", "BM07"):
        reps.append((r"^Số\s*:", f"Số: {data.so(form)}"))
    reps.append((r"^KHOA\s", data.i("don_vi", "Khoa Công nghệ thông tin").upper()))
    return reps


def para(doc, text="", size=13, bold=False, italic=False, align=None, space_after=4, indent=None):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    inline(p, text, size, bold, italic)
    return p


def inline(p, text, size=13, bold=False, italic=False):
    for part in re.split(r"(\*\*.+?\*\*)", str(text)):
        if not part:
            continue
        b = bold or (part.startswith("**") and part.endswith("**"))
        part = part[2:-2] if part.startswith("**") and part.endswith("**") else part
        lines = part.split("<br>")
        for k, line in enumerate(lines):
            r = p.add_run(line)
            set_run_font(r, size, bold=b, italic=italic)
            if k < len(lines) - 1:
                r.add_break()


def heading(doc, text, level=1):
    size = {1: 14, 2: 13, 3: 13}[level]
    p = doc.add_paragraph()
    try:
        p.style = doc.styles[f"Heading {level}"]
        # bỏ đánh số tự động của style (số mục ghi tay trong text)
        ppr = p.style.element.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:numPr")) is not None:
            ppr.remove(ppr.find(qn("w:numPr")))
    except KeyError:
        pass
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size, bold=True, italic=(level == 3))
    r.font.color.rgb = None
    return p


def set_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "000000")):
            e.set(qn(f"w:{k}"), v)
        borders.append(e)
    tblPr.append(borders)


def shade(cell, color="DDEBF7"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def table(doc, rows, header=1, size=11, widths=None, bold_rows=(), center_cols=(), borders=True):
    """rows: list[list]; ô None = gộp với ô bên trái, '^' = gộp với ô bên trên."""
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if borders:
        set_borders(t)
    merges = []
    for i, r in enumerate(rows):
        r = list(r) + [""] * (ncol - len(r))
        for j, v in enumerate(r):
            cell = t.cell(i, j)
            if v is None:
                merges.append(("L", i, j))
                continue
            if v == "^":
                merges.append(("U", i, j))
                continue
            cell.text = ""
            p = cell.paragraphs[0]
            if i < header or j in center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, g(v) if not isinstance(v, str) else v, size, bold=(i < header or i in bold_rows))
            if i < header:
                shade(cell)
    for kind, i, j in merges:
        a = t.cell(i, j)
        b = t.cell(i, j - 1) if kind == "L" else t.cell(i - 1, j)
        b.merge(a)
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths[:len(row.cells)]):
                row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def render_lines(doc, lines, size=13):
    """Markdown đơn giản: bảng, '- ', '1. ', '### ', đoạn văn; bỏ chú thích <!-- -->."""
    buf = []

    def flush_table():
        if buf:
            rows = md_tables(buf)
            for t in rows:
                table(doc, t, header=1, size=10 if len(t[0]) > 5 else 11)
            buf.clear()

    for line in lines:
        s = line.rstrip()
        if s.strip().startswith("|"):
            buf.append(s)
            continue
        flush_table()
        st = s.strip()
        if not st or st.startswith("<!--"):
            continue
        if st.startswith("### "):
            para(doc, st[4:], size, bold=True)
        elif st.startswith(("- ", "* ")):
            para(doc, "– " + st[2:], size, indent=0.6, space_after=2)
        elif re.match(r"^\d+\.\s", st):
            para(doc, st, size, indent=0.6, space_after=2)
        else:
            para(doc, st, size, align="justify")
    flush_table()


def signatures(doc, cols, size=13):
    """cols: list[(chức danh, tên)]. Bảng không viền."""
    rows = [[c[0] for c in cols], ["(Ký và ghi rõ họ tên)"] * len(cols), [""] * len(cols), [""] * len(cols), [c[1] for c in cols]]
    t = table(doc, rows, header=0, size=size, borders=False)
    for j in range(len(cols)):
        for i in range(len(rows)):
            p = t.cell(i, j).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = i in (0, 4)
                r.italic = i == 1


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    for kind, val in (("begin", None), ("instr", 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), ("text", "Nhấn F9 để cập nhật mục lục."), ("end", None)):
        if kind == "instr":
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = val
        elif kind == "text":
            e = OxmlElement("w:t")
            e.text = val
        else:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                e.set(qn("w:dirty"), "true")
        run._r.append(e)


def plan_table_rows(data, prereq=False, detail=True):
    """Bảng CTĐT theo học kỳ: dùng cho chương trình điều hành (prereq=True), BM06, BM07."""
    if prereq:
        rows = [["Mã HP", "Tên học phần", "Số tín chỉ", "Học phần trước", "Ghi chú"]]
    else:
        rows = [["Mã MH", "Môn học", "Số tín chỉ", None, None, None, "Số tiết tổng"],
                ["^", "^", "Tổng TC", "Lý thuyết", "Thực hành", "Bài tập", "^"]]
    bold = []
    width = len(rows[-1])
    for hk, items in data.semesters():
        for it in items:
            group = [it[1]] if it[0] == "hp" else it[2]
            if it[0] == "nhom":
                rows.append([f"Môn tự chọn (Chọn {it[1]} trong {len(it[2])} môn)"] + [None] * (width - 1))
                bold.append(len(rows) - 1)
            for c in group:
                note = c["ghi_chu"] or ("Môn cơ bản" if c["khoi"] in ("Lý luận chính trị", "Ngoại ngữ", "Đại cương khác", "Giáo dục thể chất") else "")
                if prereq:
                    rows.append([c["ma"], c["ten"], g(c["tc"]), data.prereq_text(c), note])
                else:
                    rows.append([c["ma"], c["ten"], g(c["tc"]), g(c["lt"]), g(c["th"]), g(c["bt"]), g(c["tiet"])])
        tot = data.sem_total(items)
        if prereq:
            rows.append([f"TỔNG CỘNG HỌC KỲ {hk}", None, g(tot["tc"]), "", ""])
        else:
            rows.append([f"TỔNG CỘNG HỌC KỲ {hk}", None, g(tot["tc"]), g(tot["lt"]), g(tot["th"]), g(tot["bt"]), g(tot["tiet"])])
        bold.append(len(rows) - 1)
    tot = data.total()
    if prereq:
        rows.append(["TỔNG CỘNG TOÀN KHÓA", None, g(tot["tc"]), "", ""])
    else:
        rows.append(["TỔNG CỘNG TOÀN KHÓA", None, g(tot["tc"]), g(tot["lt"]), g(tot["th"]), g(tot["bt"]), g(tot["tiet"])])
    bold.append(len(rows) - 1)
    return rows, bold


# ================================================================ từng biểu mẫu

def build_dh(data, out):
    doc, kept = base_doc(data, "DH")
    if not kept:
        para(doc, "TRƯỜNG ĐẠI HỌC LẠC HỒNG", 13, bold=True)
        para(doc, data.i("don_vi", "Khoa Công nghệ thông tin").upper(), 13, bold=True)
    para(doc, f"CHƯƠNG TRÌNH ĐÀO TẠO TÍN CHỈ NIÊN KHÓA {data.i('nien_khoa')}", 14, bold=True, align="center")
    para(doc, f"NGÀNH: {data.nganh().upper()}", 13, bold=True, align="center")
    rows, bold = plan_table_rows(data, prereq=True)
    table(doc, rows, header=1, size=11, widths=[2, 7.5, 1.8, 3.5, 2.5], bold_rows=bold, center_cols=(0, 2))
    para(doc, "Lưu ý: Ký hiệu loại môn học trong chương trình đào tạo", 12, bold=True)
    for k, v in (("*", "Môn học tiên quyết"), ("**", "Môn học trước"), ("***", "Môn học song hành")):
        para(doc, f"{k}\t{v}", 12, space_after=0)
    save_docx(doc, out / "Chuong-trinh-dao-tao-dieu-hanh.docx")


def build_bm03(data, out):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"CTDT_{data.fpart('khoa_tuyen')}"
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    f = lambda **k: Font(name=FONT, size=k.pop("size", 12), **k)
    ws.append(["TRƯỜNG ĐẠI HỌC LẠC HỒNG", None, None, None, None, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"])
    ws.append([data.i("don_vi", "Khoa Công nghệ thông tin").upper(), None, None, None, None, "Độc lập - Tự do - Hạnh phúc"])
    ws.append([None, None, None, None, None, ngay_dai(data.ngay("BM03"), data.i("noi_ky", "Đồng Nai"), data.i("nam"))])
    ws.append(["CHƯƠNG TRÌNH ĐÀO TẠO THEO HỌC CHẾ TÍN CHỈ"])
    ws.append([f"NGÀNH: {data.nganh().upper()} – NIÊN KHÓA {data.i('nien_khoa')}"])
    for r in range(1, 6):
        for c in ws[r]:
            c.font = f(bold=r in (1, 2, 4, 5), size=14 if r == 4 else 12, italic=r == 3)
    ws.merge_cells("A4:K4")
    ws.merge_cells("A5:K5")
    ws["A4"].alignment = ws["A5"].alignment = Alignment(horizontal="center")
    h1 = ["Mã MH", "Môn học", "Số tín chỉ", None, None, None, "Số tiết tổng", "Số tiết LT lên lớp", "Số tiết TH lên lớp", "Số tiết BT / Tự học", "Ghi chú"]
    h2 = [None, None, "Tổng TC", "Lý thuyết", "Thực hành", "Bài tập"]
    ws.append(h1)
    ws.append(h2)
    hr = ws.max_row - 1
    ws.merge_cells(start_row=hr, start_column=3, end_row=hr, end_column=6)
    for col in (1, 2, 7, 8, 9, 10, 11):
        ws.merge_cells(start_row=hr, start_column=col, end_row=hr + 1, end_column=col)
    q = data.info.get("quy_doi_tiet") or {}
    fill_total = PatternFill("solid", fgColor="DDEBF7")
    total_rows = []
    for hk, items in data.semesters():
        for it in items:
            if it[0] == "nhom":
                ws.append([f"Môn tự chọn (Chọn {it[1]} trong {len(it[2])} môn)"])
                ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=11)
                ws.cell(ws.max_row, 1).font = f(italic=True, bold=True)
            for c in ([it[1]] if it[0] == "hp" else it[2]):
                lt_t = c["lt"] * q["lt"] if "lt" in q else None
                th_t = c["th"] * q["th"] if "th" in q else None
                bt_t = c["tiet"] - (lt_t or 0) - (th_t or 0) if q else None
                ws.append([int(c["ma"]) if c["ma"].isdigit() else c["ma"], c["ten"], c["tc"], c["lt"], c["th"], c["bt"], c["tiet"],
                           lt_t, th_t, bt_t, c["ghi_chu"]])
        tot = data.sem_total(items)
        ws.append([f"TỔNG CỘNG HỌC KỲ {hk}", None, tot["tc"], tot["lt"], tot["th"], tot["bt"], tot["tiet"]])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=2)
        total_rows.append(ws.max_row)
    tot = data.total()
    ws.append(["TỔNG CỘNG TOÀN KHÓA", None, tot["tc"], tot["lt"], tot["th"], tot["bt"], tot["tiet"]])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=2)
    total_rows.append(ws.max_row)
    last = ws.max_row
    for row in ws.iter_rows(min_row=hr, max_row=last, max_col=11):
        for c in row:
            c.border = box
            c.font = f(bold=c.row in total_rows or c.row < hr + 2)
            c.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left" if c.column == 2 else "center")
            if c.row in total_rows or c.row < hr + 2:
                c.fill = fill_total
    for col, w in zip("ABCDEFGHIJK", (10, 44, 8, 9, 9, 8, 9, 10, 10, 11, 24)):
        ws.column_dimensions[col].width = w
    ws.append([])
    ws.append([("Số tiết lên lớp quy đổi: 1 TC LT = %s tiết, 1 TC TH = %s tiết." % (q.get("lt"), q.get("th"))) if q
               else "Số tiết lên lớp: chưa khai báo quy tắc quy đổi (\"quy_doi_tiet\" trong thong-tin.json)."])
    ws.append(["File sinh tự động bằng build-ctdt.py từ du-lieu/ke-hoach.csv. Sửa dữ liệu rồi build lại, không sửa tay file này."])
    ws.append([])
    ws.append(["Nơi nhận:", None, None, None, None, None, None, "TRƯỞNG KHOA"])
    ws.append(["- Ban Giám hiệu;"])
    ws.append(["- Phòng Đào tạo;"])
    ws.append([f"- Lưu: VT, K.{data.i('ky_hieu_don_vi', 'CNTT')}.", None, None, None, None, None, None, data.i("truong_khoa")])
    for row in ws.iter_rows(min_row=last + 1, max_row=ws.max_row):
        for c in row:
            c.font = f(italic=c.row <= last + 3, bold=c.value in ("TRƯỞNG KHOA", "Nơi nhận:"))
    wb.save(out / f"BM03-Ke-hoach-dao-tao-{data.fpart('ky_hieu_don_vi', 'CNTT')}-{data.fpart('khoa_tuyen')}.xlsx")


def build_plo(data, out):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PLO"
    bold = Font(name=FONT, size=12, bold=True)
    ws.append(["#", "Nội dung", "Năng lực"])
    for p in data.cdr.get("peo", []):
        ws.append([p["ma"], p["noi_dung"], p.get("nang_luc", "")])
    ws.append([])
    ws.append(["#", "Nội dung", "Miền"])
    for p in data.cdr.get("plo", []):
        ws.append([p["ma"], p["noi_dung"], p.get("mien", "")])
    ws.append([])
    ws.append(["#", "Nội dung", "Target", "PLO"])
    for p in data.cdr.get("pi", []):
        ws.append([f"PI {p['ma']}", p["noi_dung"], p.get("target", 0.75), f"PLO{p['ma'].split('.')[0]}"])
    for row in ws.iter_rows():
        for c in row:
            c.font = bold if c.value in ("#", "Nội dung", "Năng lực", "Miền", "Target", "PLO") else Font(name=FONT, size=12)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 22
    m = wb.create_sheet("Map CLO-PI")
    pis = [p["ma"] for p in data.cdr.get("pi", [])]
    cmap = data.map_clo_pi()
    lead = ["", "", "", "", "", "", "", "", "", ""]
    m.append(["Ma trận CLO – PI (sinh từ bảng CLO của đề cương du-lieu/de-cuong/*.md)"])
    plo_row = lead[:]
    last = None
    for pi in pis:
        plo = f"PLO{pi.split('.')[0]}"
        plo_row.append(plo if plo != last else "")
        last = plo
    m.append(plo_row)
    m.append(lead + [f"PI {pi}" for pi in pis])
    counts = []
    for pi in pis:
        counts.append(sum(1 for code in cmap for k, s in cmap[code].items() if pi in s))
    m.append(["Mã HP", "Tên HP", "Học kỳ", "Số TC", "TCLT", "TCHT", "TCBT", "Tổng số tiết", "Khối", "Số CLO"] + counts)
    for c in data.courses:
        if not c["tc"]:
            continue
        mp = cmap.get(c["ma"], {})
        row = [int(c["ma"]) if c["ma"].isdigit() else c["ma"], c["ten"], c["hoc_ky"], c["tc"], c["lt"], c["th"], c["bt"], c["tiet"], c["khoi"], len(mp)]
        for pi in pis:
            row.append(", ".join(f"CLO{k}" for k in sorted(mp) if pi in mp[k]))
        m.append(row)
    for row in m.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=10, bold=c.row <= 4)
    m.column_dimensions["B"].width = 40
    wb.save(out / f"PLO-{data.fpart('khoa_tuyen')}-Chuan-dau-ra-va-ma-tran-CLO-PI.xlsx")
    zero = [pi for pi, n in zip(pis, counts) if n == 0]
    if zero:
        data.warnings.append(f"PI chưa có học phần nào đánh giá: {', '.join(zero)}")


def build_bm02(data, out):
    doc, kept = base_doc(data, "BM02", header_repl(data, "BM02"))
    admin_header(doc, data, "BM02", kept)
    para(doc, "CHUẨN ĐẦU RA", 14, bold=True, align="center", space_after=0)
    para(doc, f"NGÀNH {data.nganh().upper()}", 14, bold=True, align="center")
    para(doc, "Tên ngành đào tạo:", bold=True)
    para(doc, f"Tiếng Việt: {data.nganh().upper()}", indent=0.6)
    para(doc, f"Tiếng Anh: {data.i('nganh_en').upper()}", indent=0.6)
    para(doc, f"Mã ngành: {data.i('ma_nganh')}", indent=0.6)
    para(doc, f"Trình độ đào tạo: {data.i('trinh_do', 'Đại học')}", indent=0.6)
    heading(doc, "1. Mục tiêu đào tạo", 2)
    table(doc, [["Mục tiêu của chương trình (PEOs)", "Nội dung"]] + [[p["ma"], p["noi_dung"]] for p in data.cdr.get("peo", [])],
          widths=[3.5, 12.5], center_cols=(0,))
    heading(doc, "2. Chuẩn đầu ra", 2)
    rows = [["Chuẩn đầu ra CTĐT (PLOs/SOs)", "Nội dung"]]
    bold = []
    pis = defaultdict(list)
    for p in data.cdr.get("pi", []):
        pis[p["ma"].split(".")[0]].append(p)
    for mien in ("Kiến thức", "Kỹ năng", "Thái độ", "Năng lực tự chủ và trách nhiệm"):
        items = [p for p in data.cdr.get("plo", []) if p.get("mien") == mien]
        if not items:
            continue
        rows.append([mien, None])
        bold.append(len(rows) - 1)
        for p in items:
            text = p["noi_dung"]
            if data.info.get("bm02_kem_pi", True):
                text += "".join(f"<br>– PI{x['ma']}: {x['noi_dung']}" for x in pis[re.sub(r'\D', '', p['ma'])])
            rows.append([p["ma"], text])
    table(doc, rows, widths=[3.5, 12.5], bold_rows=bold, center_cols=(0,))
    heading(doc, "3. Vị trí làm việc của người học sau tốt nghiệp", 2)
    for i, grp in enumerate(data.cdr.get("vi_tri_viec_lam", []), 1):
        para(doc, f"3.{i}. {grp['nhom']}", bold=True)
        if grp.get("mo_ta"):
            para(doc, grp["mo_ta"], align="justify")
        for item in grp.get("muc", []):
            para(doc, "– " + item, indent=0.6, align="justify", space_after=2)
    heading(doc, "4. Khả năng học tập nâng cao trình độ sau khi ra trường", 2)
    para(doc, data.cdr.get("hoc_nang_cao") or PH % "khả năng học nâng cao", align="justify")
    heading(doc, "5. Hình thức và thời gian đào tạo", 2)
    para(doc, f"Hình thức đào tạo: {data.i('loai_hinh', 'Chính quy tập trung').lower()}.")
    para(doc, f"Thời gian đào tạo: {data.i('thoi_gian', '4 năm')}.")
    heading(doc, "6. Đối tượng tuyển sinh", 2)
    para(doc, data.i("doi_tuong_tuyen_sinh", "Theo Quy chế tuyển sinh đại học hệ chính quy hiện hành của Bộ Giáo dục và Đào tạo và Trường Đại học Lạc Hồng."), align="justify")
    signatures(doc, [("", ""), ("TRƯỞNG KHOA", data.i("truong_khoa"))])
    save_docx(doc, out / "BM02-Cong-bo-chuan-dau-ra.docx")


def build_bm04(data, out):
    if not data.doisanh:
        data.warnings.append("BM04: chưa có du-lieu/doi-sanh.csv, bỏ qua")
        return
    doc, kept = base_doc(data, "BM04", header_repl(data, "BM04"))
    admin_header(doc, data, "BM04", kept)
    lhu = [c for c in data.courses if c["tc"]]
    by_school = defaultdict(list)
    for r in data.doisanh:
        by_school[(nfc(r["truong"]), nfc(r.get("chuong_trinh", "")), nfc(r.get("nguon", "")), nfc(r.get("nam", "")))].append(r)
    for (truong, ct, nguon, nam), rows in by_school.items():
        para(doc, "BẢNG SO SÁNH CHƯƠNG TRÌNH ĐÀO TẠO CỦA TRƯỜNG ĐẠI HỌC LẠC HỒNG", 13, bold=True, align="center", space_after=0)
        para(doc, f"VỚI CHƯƠNG TRÌNH ĐÀO TẠO CỦA {truong.upper()}", 13, bold=True, align="center")
        para(doc, f"Chương trình đối sánh: {ct or PH % 'tên chương trình'}; năm áp dụng: {nam or PH % 'năm'}; nguồn: {nguon or PH % 'URL/văn bản'}.", 12, italic=True)
        matched = set()
        body = []
        for r in rows:
            codes = re.findall(r"\d{6}", r.get("ma_lhu", ""))
            matched.update(codes)
            names = "; ".join(data.by[c]["ten"] for c in codes if c in data.by)
            bad = [c for c in codes if c not in data.by]
            if bad:
                data.warnings.append(f"BM04: mã {', '.join(bad)} trong doi-sanh.csv không có trong ke-hoach.csv")
            body.append([nfc(r["mon_doi_sanh"]), names or "Không có", ""])
        for c in lhu:
            if c["ma"] not in matched:
                body.append(["", "", c["ten"]])
        n_eq = len([c for c in lhu if c["ma"] in matched])
        n = len(lhu)
        head = [[f"Tên môn học của {ct or truong} (tổng số {len(rows)} môn)", f"Tên môn học của ngành {data.nganh()} – Trường Đại học Lạc Hồng (tổng số {n} môn)", None],
                ["^", "Những môn học tương đương", "Những môn học không tương đương"]]
        foot = [["", f"Tổng số môn học tương đương là: {n_eq}/{n}<br>Tỷ lệ {n_eq / n * 100:.2f}%",
                 f"Tổng số môn học không tương đương là: {n - n_eq}/{n}<br>Tỷ lệ {(n - n_eq) / n * 100:.2f}%"]]
        table(doc, head + body + foot, header=2, size=11, widths=[5.5, 5.5, 5.5], bold_rows=(len(head + body),))
        para(doc, "Nhận xét:", bold=True)
        para(doc, PH % f"điểm giống, khác so với {truong} và điều chỉnh rút ra cho CTĐT", align="justify")
    signatures(doc, [("", ""), ("TRƯỞNG KHOA", data.i("truong_khoa"))])
    save_docx(doc, out / "BM04-Bang-doi-chieu-so-sanh-CTDT.docx")


def block_totals(courses):
    tot = defaultdict(float)
    for c in courses:
        if c["nhom"] and c is not first_option(courses, c):
            continue
        tot[c["khoi"]] += c["tc"] or 0
    return tot


def first_option(courses, c):
    key = c["nhom"].split(":")[0]
    k = int(c["nhom"].split(":")[1]) if ":" in c["nhom"] else 1
    same = [x for x in courses if x["nhom"].split(":")[0] == key]
    return c if c in same[:k] else None


def khoi_table_rows(courses):
    tot = block_totals(courses)
    rows = [["Khối kiến thức", None, "Tổng số tín chỉ", "Ghi chú"]]
    for nhom in ("Khối kiến thức giáo dục đại cương", "Khối kiến thức giáo dục chuyên nghiệp", "Tốt nghiệp"):
        first = True
        for k, grp in KHOI.items():
            if grp == nhom:
                label = {"Đại cương khác": "Môn học khác", "Chuyên nghiệp khác": "Môn học khác",
                         "Thực tập": "Thực tập doanh nghiệp, kiến tập, đồ án", "Tốt nghiệp": "Khóa luận hoặc chuyên đề tốt nghiệp"}.get(k, k)
                rows.append([nhom if first else "^", label, g(tot.get(k, 0.0)), ""])
                first = False
    rows.append(["Tổng số tín chỉ tích lũy tối thiểu toàn khóa", None, g(sum(v for k, v in tot.items() if KHOI.get(k))), ""])
    return rows


def build_bm05(data, out):
    folder = out / "BM05-De-cuong"
    folder.mkdir(exist_ok=True)
    for old in folder.glob("*.docx"):
        old.unlink()
    for c in data.courses:
        s = data.syl.get(c["ma"])
        if not s:
            if not c["ma"].startswith("102") and c["tc"]:
                data.warnings.append(f"BM05: {c['ma']} {c['ten']} chưa có du-lieu/de-cuong/{c['ma']}.md")
            continue
        meta = s["meta"]
        form = meta.get("mau", "BM05")
        if form not in ("BM05", "BM05a", "BM05b"):
            form = "BM05"
        doc, kept = base_doc(data, form)
        if not kept and form == "BM05":
            para(doc, "BỘ GIÁO DỤC VÀ ĐÀO TẠO", 13, align="center", space_after=0)
            para(doc, "TRƯỜNG ĐẠI HỌC LẠC HỒNG", 13, bold=True, align="center")
        if form != "BM05":
            for t in ("BỘ GIÁO DỤC VÀ ĐÀO TẠO", "TRƯỜNG ĐẠI HỌC LẠC HỒNG", "LAC HONG UNIVERSITY"):
                para(doc, t, 13, bold=t != "BỘ GIÁO DỤC VÀ ĐÀO TẠO", align="center", space_after=0)
        para(doc, "ĐỀ CƯƠNG HỌC PHẦN", 16, bold=True, align="center", space_after=0)
        if form != "BM05":
            para(doc, "COURSE SYLLABUS", 13, bold=True, align="center", space_after=0)
            para(doc, "HỌC PHẦN TÍCH HỢP TOÀN PHẦN HỌC LIỆU ASU (ASU FULL-ADOPTED COURSE)" if form == "BM05a"
                 else "HỌC PHẦN TĂNG CƯỜNG TỪ HỌC LIỆU ASU (ASU-ENHANCED COURSE)", 12, bold=True, align="center", space_after=0)
        para(doc, f"{c['ma']} – {c['ten'].upper()}", 14, bold=True, align="center")
        n = 0
        if form != "BM05":
            lines = section(s, "Bản quyền")
            heading(doc, "0. KHỐI BẢN QUYỀN VÀ NGUỒN HỌC LIỆU (Rights & Attribution)", 2)
            render_lines(doc, lines or [PH % "bảng khối bản quyền theo mẫu " + form])
        n += 1
        heading(doc, f"{n}. THÔNG TIN CHUNG (General information)", 2)
        grp = next((x for x in data.courses if x["nhom"] and x["nhom"] == c["nhom"]), None)
        tinh_chat = meta.get("tinh_chat") or (f"Tự chọn đối với sinh viên ngành {data.nganh()}" if c["nhom"]
                                               else f"Bắt buộc đối với sinh viên ngành {data.nganh()}")
        gv = meta.get("giang_vien") or PH % "giảng viên phụ trách"
        rows = [["Tên học phần (tiếng Việt):", c["ten"]],
                ["Tên học phần (tiếng Anh):", c.get("ten_en") or meta.get("ten_en") or PH % "tên tiếng Anh"],
                ["Mã môn học:", c["ma"]],
                ["Thuộc khối kiến thức:", KHOI_TEN_DE_CUONG.get(c["khoi"], c["khoi"]) or PH % "khối"],
                ["Khoa/Bộ môn phụ trách:", meta.get("bo_mon") or data.i("don_vi", "Khoa Công nghệ thông tin")],
                ["Giảng viên phụ trách:", f"{gv}, Email: {meta.get('email') or PH % 'email'}"],
                ["Giảng viên tham gia giảng dạy:", meta.get("giang_vien_tham_gia") or PH % "danh sách giảng viên"],
                ["Số tín chỉ:", g(c["tc"])], ["Lý thuyết:", g(c["lt"])], ["Thực hành:", g(c["th"])], ["Bài tập:", g(c["bt"])]]
        if form != "BM05":
            rows += [["Đối sánh khối lượng với ASU:", meta.get("doi_sanh_khoi_luong") or PH % "credit hours ASU ↔ tín chỉ LHU"],
                     ["Ngôn ngữ:", meta.get("ngon_ngu") or PH % "ngôn ngữ giảng dạy / học liệu / đánh giá"],
                     ["Hình thức tổ chức:", meta.get("hinh_thuc_to_chuc") or PH % "trực tiếp / trực tuyến / kết hợp"]]
        rows += [["Tính chất của học phần:", tinh_chat],
                 ["Môn học trước:", data.prereq_text(c, with_name=True) or "Không"]]
        table(doc, rows, header=0, size=12, widths=[5.5, 10.5])
        known = [("Mô tả", "MÔ TẢ HỌC PHẦN (Course description)"),
                 ("Chuẩn đầu ra", "CHUẨN ĐẦU RA HỌC PHẦN (Course learning outcomes)"),
                 ("Kế hoạch", "NỘI DUNG HỌC PHẦN, KẾ HOẠCH GIẢNG DẠY (Course content, Lesson plan)"),
                 ("Ma trận", "MA TRẬN CHUẨN ĐẦU RA HỌC PHẦN VÀ CHUẨN ĐẦU RA BÀI HỌC"),
                 ("Đánh giá", "ĐÁNH GIÁ HỌC PHẦN (Course assessment)"),
                 ("Quy định", "QUY ĐỊNH CỦA HỌC PHẦN (Course requirements and expectations)"),
                 ("Tài liệu", "TÀI LIỆU HỌC TẬP, THAM KHẢO (Learning resources)"),
                 ("Phần mềm", "PHẦN MỀM HAY CÔNG CỤ HỖ TRỢ THỰC HÀNH")]
        done = set()
        for key, title in known:
            n += 1
            heading(doc, f"{n}. {title}", 2)
            lines = section(s, key)
            has = lines and any(l.strip() and not l.strip().startswith("<!--") for l in lines)
            render_lines(doc, lines if has else [PH % f"nội dung mục {key}"])
            done.add(key)
            # mục phụ của mẫu ASU (Bảng 1b, Bảng 2b...) đặt ngay sau mục chính tương ứng
            for name, extra in s["sections"]:
                if name.casefold().startswith(f"{key.casefold()} –") or name.casefold().startswith(f"{key.casefold()} -"):
                    para(doc, name.split("–", 1)[-1].split("-", 1)[-1].strip(), bold=True)
                    render_lines(doc, extra)
        n += 1
        heading(doc, f"{n}. QUY ƯỚC KÝ HIỆU", 2)
        legend = [["Ký hiệu", "Mô tả"], ["PLO/SO", "Chuẩn đầu ra của chương trình đào tạo"], ["PI", "Chỉ số đánh giá (chuẩn đầu ra trung gian)"],
                  ["CLO", "Chuẩn đầu ra học phần"], ["LLO", "Chuẩn đầu ra của bài học"], ["GV", "Giảng viên"], ["SV", "Sinh viên"]]
        if form != "BM05":
            legend += [["[U]", "Nội dung/học liệu kế thừa nguyên bản từ Arizona State University"],
                       ["[U-adapt]", "Nội dung/học liệu có nguồn gốc ASU đã được dịch, biên tập hoặc hiệu chỉnh"],
                       ["[L]", "Nội dung/học liệu do giảng viên Trường Đại học Lạc Hồng biên soạn hoặc lựa chọn"]]
        table(doc, legend, size=11, widths=[3, 13])
        nam = (data.ngay("BM05") or "")[:4] or data.i("nam", "…")
        para(doc, f"{data.i('noi_ky', 'Đồng Nai')}, ngày … tháng … năm {nam}", 13, italic=True, align="right")
        cols = [("TRƯỞNG KHOA/TRƯỞNG NGÀNH", ""), ("GIẢNG VIÊN PHỤ TRÁCH", gv if "{{" not in gv else "")]
        if form != "BM05":
            cols.append(("XÁC NHẬN RÀ SOÁT BẢN QUYỀN HỌC LIỆU", ""))
        signatures(doc, cols, 12)
        save_docx(doc, folder / f"HK{c['hoc_ky']}-{c['ma']}-{safe_name(c['ten'])}.docx")
    for code in data.syl:
        if code not in data.by:
            data.warnings.append(f"BM05: de-cuong/{code}.md không có trong ke-hoach.csv (không sinh)")


def mo_ta(data, code):
    s = data.syl.get(code)
    if not s:
        return PH % f"mô tả học phần {code}"
    if s["meta"].get("tom_tat"):
        return s["meta"]["tom_tat"]
    lines = [l.strip() for l in (section(s, "Mô tả") or []) if l.strip() and not l.strip().startswith(("|", "<!--"))]
    return " ".join(l.lstrip("-* ") for l in lines) or PH % f"mô tả học phần {code}"


def tai_lieu(data, code, k=3):
    s = data.syl.get(code)
    lines = section(s, "Tài liệu") if s else None
    out = [l.strip().lstrip("-* ") for l in (lines or []) if re.match(r"^\s*[-*]?\s*\[\d+\]", l)]
    return out[:k]


def build_tt(data, out):
    doc, kept = base_doc(data, "TT")
    if not kept:
        para(doc, "TRƯỜNG ĐẠI HỌC LẠC HỒNG", 13, bold=True)
        para(doc, data.i("don_vi", "Khoa Công nghệ thông tin").upper(), 13, bold=True)
    para(doc, "TÓM TẮT HỌC PHẦN", 14, bold=True, align="center", space_after=0)
    para(doc, f"NGÀNH: {data.nganh().upper()}", 13, bold=True, align="center")
    for c in data.courses:
        if c["ma"].startswith("102") or not c["tc"]:
            continue
        heading(doc, f"{c['ma']} - {c['ten']}", 2)
        para(doc, mo_ta(data, c["ma"]), align="justify")
        refs = tai_lieu(data, c["ma"])
        if refs:
            para(doc, "Tài liệu tham khảo:", bold=True)
            for r in refs:
                para(doc, r, indent=0.6, space_after=2)
    save_docx(doc, out / "Tom-tat-hoc-phan.docx")


def build_bm01(data, out):
    people = data.info.get("nguoi_gop_y", [])
    folder = out / "BM01-Phieu-y-kien"
    folder.mkdir(exist_ok=True)
    blank = [{"ho_ten": "", "don_vi": "", "loai": "BM01.1"}, {"ho_ten": "", "don_vi": "", "loai": "BM01.2"}]
    for i, pp in enumerate(blank + people):
        form = "BM01.2" if pp.get("loai") == "BM01.2" else "BM01.1"
        doc, kept = base_doc(data, form)
        if not kept:
            para(doc, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", 13, bold=True, align="center", space_after=0)
            para(doc, "Độc lập - Tự do - Hạnh phúc", 13, bold=True, align="center")
        intro = (f"Sau khi đọc xong cuốn chương trình giáo dục đại học ngành {data.nganh()} của "
                 f"{data.i('don_vi', 'Khoa Công nghệ thông tin')} Trường Đại học Lạc Hồng (niên khóa {data.i('nien_khoa')}), tôi có một số nhận xét sau đây:")
        dots = "…" * 60
        if form == "BM01.1":
            para(doc, "Ý KIẾN CÁC BÊN LIÊN QUAN VỀ CHƯƠNG TRÌNH ĐÀO TẠO", 14, bold=True, align="center")
        else:
            para(doc, "PHIẾU GÓP Ý CHƯƠNG TRÌNH ĐÀO TẠO", 14, bold=True, align="center")
        para(doc, f"Họ và tên người tham gia góp ý: {pp.get('ho_ten') or '…' * 30}")
        para(doc, f"Đơn vị công tác: {pp.get('don_vi') or '…' * 34}")
        para(doc, intro, align="justify")
        if form == "BM01.1":
            items = [("Về nhu cầu kinh tế – xã hội đối với ngành đào tạo", []),
                     ("Về mục tiêu và chuẩn đầu ra của CTĐT", []),
                     ("Về khối lượng, cấu trúc, nội dung chương trình; phương pháp đào tạo", [
                         "Về khối kiến thức cơ sở", "Về khối kiến thức chuyên ngành",
                         "Về tỷ lệ giữa kiến thức lý thuyết và thực hành (học phần thực hành, tự nghiên cứu)",
                         "Về tính liên thông của CTĐT", "Về nguồn tài liệu tham khảo"]),
                     ("Về cơ sở vật chất phục vụ đào tạo", []), ("Về đội ngũ giảng viên", []),
                     ("Về đề cương chi tiết các học phần", []),
                     ("Về khả năng đáp ứng nhu cầu kinh tế – xã hội", [
                         "Khả năng đáp ứng yêu cầu về kiến thức, kỹ năng của trình độ đào tạo", "Khả năng đáp ứng yêu cầu của người học",
                         "Khả năng đáp ứng yêu cầu nguồn nhân lực của thị trường lao động",
                         "Tính hội nhập quốc tế về nội dung chương trình, phương pháp đào tạo"]),
                     ("Ý kiến khác", [])]
            for k, (title, subs) in enumerate(items, 1):
                para(doc, f"{k}. {title}", bold=True)
                for sub in subs:
                    para(doc, f"– {sub}", italic=True, indent=0.6)
                    para(doc, dots, size=11)
                if not subs:
                    para(doc, dots, size=11)
                    para(doc, dots, size=11)
            para(doc, "KẾT LUẬN CHUNG", bold=True)
            para(doc, dots, size=11)
        else:
            rows = [["TT", "Nội dung thẩm định", "Nhận xét của thành viên Hội đồng", "Kết luận (đáp ứng yêu cầu hay không đáp ứng yêu cầu)"],
                    ["1", "Nhu cầu KT – XH đối với ngành đào tạo", "", ""],
                    ["2", "Mục tiêu đào tạo và chuẩn đầu ra", "", ""],
                    ["3", "Cấu trúc chương trình đào tạo: sự hợp lý của các học phần; thời lượng của từng học phần", "", ""],
                    ["4", "Nội dung của CTĐT (đáp ứng mục tiêu, phù hợp trình độ đào tạo, bảo đảm tính hiện đại, tính hội nhập và phù hợp với yêu cầu phát triển KT – XH)", "", ""],
                    ["5", "Đề cương chi tiết học phần (mục tiêu, nội dung, sự trùng lặp với các học phần khác, phương pháp đánh giá, tài liệu tham khảo)", "", ""]]
            table(doc, rows, size=12, widths=[1, 5.5, 6.5, 3.5], center_cols=(0,))
            para(doc, "Kết luận chung:", bold=True)
            para(doc, dots, size=11)
        para(doc, f"{data.i('noi_ky', 'Đồng Nai')}, ngày … tháng … năm {data.i('nam', '…')}", italic=True, align="right")
        signatures(doc, [("", ""), ("NGƯỜI GÓP Ý KIẾN", pp.get("ho_ten", ""))])
        name = f"{form}-Mau-trong" if i < 2 else f"{form}-{i - 1:02d}-{safe_name(pp.get('ho_ten', ''))}"
        save_docx(doc, folder / f"{name}.docx")


def changes_vs_prev(data):
    if not data.prev:
        return None
    old = {c["ma"]: c for c in data.prev}
    moved, new, removed, changed = [], [], [], []
    for c in data.courses:
        o = old.get(c["ma"])
        if not o:
            new.append(c)
        else:
            if o["hoc_ky"] != c["hoc_ky"]:
                moved.append((c, o["hoc_ky"]))
            if (o["tc"], o["lt"], o["th"], o["bt"]) != (c["tc"], c["lt"], c["th"], c["bt"]) or nfc(o["ten"]) != nfc(c["ten"]):
                changed.append((c, o))
    for o in data.prev:
        if o["ma"] not in data.by:
            removed.append(o)
    return moved, new, removed, changed


def tc_str(c):
    return f"{g(c['tc'])}({g(c['lt'])}-{g(c['th'])}-{g(c['bt'])})"


def build_bm06(data, out):
    doc, kept = base_doc(data, "BM06", header_repl(data, "BM06"))
    admin_header(doc, data, "BM06", kept)
    hop = data.info.get("hop_hoi_dong", {})
    para(doc, "BIÊN BẢN HỌP HỘI ĐỒNG KHOA HỌC VÀ ĐÀO TẠO", 14, bold=True, align="center", space_after=0)
    para(doc, f"Về việc xây dựng chương trình đào tạo niên khóa {data.i('nien_khoa')}", 13, bold=True, align="center", space_after=0)
    para(doc, f"Ngành {data.nganh()}", 13, bold=True, align="center")
    d = parse_date(data.ngay("BM06"))
    when = f"ngày {d.day} tháng {d.month} năm {d.year}" if d else PH % "ngày họp"
    para(doc, f"Thời gian: {hop.get('gio') or PH % 'giờ'} {when}")
    para(doc, f"Địa điểm: {hop.get('dia_diem') or PH % 'địa điểm'}")
    para(doc, "Thành phần tham dự: (theo Quyết định thành lập Hội đồng khoa học và đào tạo Khoa)", bold=True)
    hd = data.info.get("hoi_dong", [])
    table(doc, [[m.get("ho_ten", ""), m.get("chuc_vu", ""), m.get("vai_tro", "")] for m in hd] or [[PH % "thành viên Hội đồng", "", ""]],
          header=0, size=12, widths=[5.5, 7.5, 3.5])
    for grp, members in data.info.get("khach_moi", {}).items():
        para(doc, grp, bold=True)
        table(doc, [[m.get("ho_ten", ""), m.get("chuc_vu", "")] for m in members] or [[PH % "họ tên", ""]], header=0, size=12, widths=[5.5, 11])
    heading(doc, "1. Căn cứ xây dựng chương trình đào tạo", 2)
    for cc in data.info.get("can_cu", []):
        txt = cc.get("van_ban", "")
        if cc.get("ngay") and not re.search(r"ngày\s+\d", txt):
            txt += f" ngày {ngay_ngan(cc['ngay'])}"
        para(doc, f"Căn cứ {txt};", align="justify", indent=0.6)
    heading(doc, "2. Nội dung cuộc họp", 2)
    para(doc, f"{data.i('nguoi_trinh_bay', data.i('truong_khoa'))} đại diện Khoa trình bày trước Hội đồng dự thảo xây dựng, điều chỉnh CTĐT "
              f"ngành {data.nganh()} niên khóa {data.i('nien_khoa')} (đã gửi các thành viên trước cuộc họp) và đề nghị góp ý các nội dung: "
              "mục tiêu đào tạo; chuẩn đầu ra; khối lượng và sự phù hợp giữa các khối kiến thức với thời gian đào tạo; kế hoạch đào tạo; đề cương chi tiết các học phần.",
         align="justify")
    heading(doc, "3. Ý kiến của các thành viên và khách mời", 2)
    para(doc, "(Thư ký ghi theo diễn biến cuộc họp. Không điền sẵn.)", italic=True)
    speakers = [m.get("ho_ten") for m in hd] + [m.get("ho_ten") for ms in data.info.get("khach_moi", {}).values() for m in ms]
    for sp in speakers or [PH % "người phát biểu"]:
        para(doc, sp, bold=True)
        para(doc, PH % f"ý kiến của {sp}", indent=0.6)
    heading(doc, "4. Kết luận", 2)
    para(doc, "Căn cứ ý kiến đóng góp của các bên liên quan và kết quả khảo sát CTĐT, Hội đồng thống nhất:", align="justify")
    for line in ("Về triết lý giáo dục, tầm nhìn và sứ mạng của Trường: " + PH % "kết luận",
                 "Về mục tiêu đào tạo và chuẩn đầu ra: " + PH % "kết luận",
                 "Tiếp thu, giải trình từng ý kiến ở mục 3: " + PH % "ý kiến – tiếp thu (thay đổi cụ thể) hoặc giải trình"):
        para(doc, "– " + line, indent=0.6, align="justify")
    ch = changes_vs_prev(data)
    if ch:
        moved, new, removed, changed = ch
        if moved:
            para(doc, "Thay đổi học kỳ của một số học phần:", bold=True)
            table(doc, [["Mã HP", "Tên học phần", "Số tín chỉ", "Học kỳ cũ", "Học kỳ mới"]] +
                  [[c["ma"], c["ten"], tc_str(c), str(o), str(c["hoc_ky"])] for c, o in moved], size=11, center_cols=(0, 2, 3, 4))
        if changed:
            para(doc, "Học phần thay đổi tên hoặc khối lượng:", bold=True)
            table(doc, [["Mã HP", "Tên cũ", "Tên mới", "Khối lượng cũ", "Khối lượng mới"]] +
                  [[c["ma"], o["ten"], c["ten"], tc_str(o), tc_str(c)] for c, o in changed], size=11, center_cols=(0, 3, 4))
        if new:
            para(doc, "Một số học phần mới:", bold=True)
            table(doc, [["Mã HP", "Tên học phần", "Số tín chỉ", "Số tiết", "Học kỳ"]] +
                  [[c["ma"], c["ten"], tc_str(c), g(c["tiet"]), str(c["hoc_ky"])] for c in new], size=11, center_cols=(0, 2, 3, 4))
        if removed:
            para(doc, "Học phần không còn trong chương trình:", bold=True)
            table(doc, [["Mã HP", "Tên học phần", "Số tín chỉ", "Học kỳ cũ"]] +
                  [[o["ma"], o["ten"], tc_str(o), str(o["hoc_ky"])] for o in removed], size=11, center_cols=(0, 2, 3))
    else:
        para(doc, PH % "bảng thay đổi so với CTĐT khóa trước (thêm du-lieu/ke-hoach-khoa-truoc.csv để sinh tự động)", italic=True)
    para(doc, f"Hội đồng thống nhất xây dựng CTĐT niên khóa {data.i('nien_khoa')} như sau:", bold=True)
    rows, bold = plan_table_rows(data)
    table(doc, rows, header=2, size=10, widths=[1.8, 6.5, 1.3, 1.5, 1.5, 1.3, 1.8], bold_rows=bold, center_cols=(0, 2, 3, 4, 5, 6))
    para(doc, f"Kết quả biểu quyết: {PH % 'số thành viên đồng ý / tổng số thành viên dự họp'}.")
    para(doc, f"Cuộc họp kết thúc lúc {hop.get('ket_thuc') or PH % 'giờ'} cùng ngày./.")
    signatures(doc, [("CHỦ TỊCH HỘI ĐỒNG", hop.get("chu_tich", "")), ("THƯ KÝ HỘI ĐỒNG", hop.get("thu_ky", ""))])
    save_docx(doc, out / "BM06-Bien-ban-hop-HDKH-va-DT.docx")


def build_bm07(data, out):
    doc, kept = base_doc(data, "BM07", header_repl(data, "BM07"))
    admin_header(doc, data, "BM07", kept)
    para(doc, "TỜ TRÌNH", 14, bold=True, align="center", space_after=0)
    para(doc, f"Về việc đề nghị ban hành chương trình đào tạo ngành {data.nganh()}", 13, bold=True, align="center", space_after=0)
    para(doc, f"trình độ {data.i('trinh_do', 'đại học').lower()}, niên khóa {data.i('nien_khoa')}", 13, bold=True, align="center")
    para(doc, "Kính gửi: Ban Giám hiệu Trường Đại học Lạc Hồng", align="center")
    d6 = parse_date(data.ngay("BM06"))
    para(doc, f"Căn cứ Biên bản số {data.so('BM06')} ngày {d6.day if d6 else '…'} tháng {d6.month if d6 else '…'} năm {d6.year if d6 else '…'} "
              f"của Hội đồng khoa học và đào tạo {data.i('don_vi', 'Khoa Công nghệ thông tin')} về việc xây dựng chương trình đào tạo niên khóa {data.i('nien_khoa')};",
         align="justify")
    para(doc, f"{data.i('don_vi', 'Khoa Công nghệ thông tin')} kính trình Ban Giám hiệu chương trình đào tạo ngành {data.nganh()} "
              f"(mã ngành {data.i('ma_nganh')}), tổng khối lượng {g(data.total()['tc'])} tín chỉ, như sau:", align="justify")
    rows, bold = plan_table_rows(data)
    table(doc, rows, header=2, size=10, widths=[1.8, 6.5, 1.3, 1.5, 1.5, 1.3, 1.8], bold_rows=bold, center_cols=(0, 2, 3, 4, 5, 6))
    para(doc, "Kính đề nghị Ban Giám hiệu nhà trường xem xét ra quyết định./.", align="justify")
    t = table(doc, [["Nơi nhận:\n- Ban Giám hiệu (báo cáo);\n- Phòng Đào tạo (để phối hợp thực hiện);\n" f"- Lưu: VT, K.{data.i('ky_hieu_don_vi', 'CNTT')}.",
                     "TRƯỞNG KHOA\n\n\n\n" + data.i("truong_khoa")]], header=0, size=12, borders=False)
    t.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    save_docx(doc, out / "BM07-To-trinh-de-nghi-ban-hanh-CTDT.docx")


def build_bm08(data, out):
    evals = data.info.get("danh_gia_ctdt", [])
    if not evals:
        data.warnings.append("BM08: chưa khai báo \"danh_gia_ctdt\" trong thong-tin.json, bỏ qua")
        return
    doc, kept = base_doc(data, "BM08", header_repl(data, "BM08"))
    admin_header(doc, data, "BM08", kept)
    for k, ev in enumerate(evals):
        if k:
            page_break(doc)
        para(doc, "BẢNG ĐÁNH GIÁ CHƯƠNG TRÌNH ĐÀO TẠO ĐẠI HỌC", 14, bold=True, align="center", space_after=0)
        para(doc, f"NIÊN KHÓA {ev.get('nien_khoa', PH % 'niên khóa')}", 13, bold=True, align="center", space_after=0)
        para(doc, f"NGÀNH {data.nganh().upper()} ({ev.get('trang_thai', PH % 'đã tốt nghiệp / đánh giá giữa chương trình')})", 13, bold=True, align="center")
        para(doc, f"Khoa, Ngành {data.nganh()} đánh giá chương trình đào tạo niên khóa {ev.get('nien_khoa', '')} "
                  f"({ev.get('trang_thai', '')}) làm căn cứ xây dựng chương trình đào tạo niên khóa {data.i('nien_khoa')}.", align="justify")
        plan = None
        if ev.get("ke_hoach"):
            p = data.dir / "du-lieu" / ev["ke_hoach"]
            plan = read_plan(p) if p.exists() else None
        heading(doc, "1. Mục tiêu đào tạo", 2)
        peo = "<br>".join(f"{p['ma']}: {p['noi_dung']}" for p in ev.get("peo", data.cdr.get("peo", [])))
        table(doc, [["Nội dung", "Bộ môn/Ngành đánh giá", "Khoa đánh giá"], [peo, PH % "đánh giá", PH % "đánh giá"]], size=11, widths=[8, 4, 4])
        heading(doc, "2. Chuẩn đầu ra", 2)
        plo = "<br>".join(f"{p['ma']}: {p['noi_dung']}" for p in ev.get("plo", data.cdr.get("plo", [])))
        table(doc, [["Nội dung", "Bộ môn/Ngành đánh giá", "Khoa đánh giá"], [plo, PH % "đánh giá", PH % "đánh giá"]], size=11, widths=[8, 4, 4])
        heading(doc, "3. Khối lượng và kiến thức toàn khóa", 2)
        if plan:
            rows = khoi_table_rows(plan)
            rows[0] = ["Khối kiến thức", None, f"CTĐT niên khóa {ev.get('nien_khoa', '')} (tín chỉ)", "Đánh giá"]
            table(doc, rows, size=11, widths=[5, 5, 3, 3], center_cols=(2,))
        else:
            para(doc, PH % "bảng khối kiến thức (khai báo \"ke_hoach\" của khóa được đánh giá)", italic=True)
        for i, title in enumerate(("Phương pháp dạy và học", "Đánh giá học phần", "Đội ngũ giảng viên", "Cơ sở vật chất"), 4):
            heading(doc, f"{i}. {title}", 2)
            table(doc, [["Nội dung", "Bộ môn/Ngành đánh giá", "Khoa đánh giá"],
                        [ev.get(title, PH % f"hiện trạng: {title.lower()} (số liệu có năm)"), PH % "đánh giá", PH % "đánh giá"]], size=11, widths=[8, 4, 4])
        heading(doc, "8. Kết luận", 2)
        para(doc, ev.get("ket_luan") or PH % "kết luận và hướng điều chỉnh cho CTĐT mới", align="justify")
    signatures(doc, [("BAN GIÁM HIỆU", ""), ("KHOA " + data.i("ky_hieu_don_vi", "CNTT"), data.i("truong_khoa")), ("BỘ MÔN/NGÀNH", data.i("truong_nganh", ""))], 12)
    save_docx(doc, out / "BM08-Danh-gia-CTDT.docx")


def build_bm09(data, out):
    noi = data.i("noi_ky", "Đồng Nai")
    d = parse_date(data.ngay("BM09"))
    reps = [(r"^(KỸ SƯ|CỬ NHÂN)\b.*CHÍNH QUY", f"{data.i('danh_hieu', PH % 'danh hiệu văn bằng').upper()} CHÍNH QUY"),
            (r"^NGÀNH\s", f"NGÀNH {data.nganh().upper()}"),
            (r"^Đồng Nai\s*[–-]\s*\d", f"{noi} – {d.month}/{d.year}" if d else f"{noi} – {PH % 'tháng/năm'}")]
    doc, kept = base_doc(data, "BM09", reps)
    if not kept:
        for t, sz in (("BỘ GIÁO DỤC VÀ ĐÀO TẠO", 14), ("TRƯỜNG ĐẠI HỌC LẠC HỒNG", 14), ("", 13), ("CHƯƠNG TRÌNH ĐÀO TẠO ĐẠI HỌC", 18),
                      (f"{data.i('danh_hieu', PH % 'danh hiệu văn bằng').upper()} CHÍNH QUY", 16), (f"NGÀNH {data.nganh().upper()}", 16)):
            para(doc, t, sz, bold=True, align="center")
        para(doc, f"{noi} – {d.month}/{d.year}" if d else f"{noi} – {PH % 'tháng/năm'}", 13, bold=True, align="center")
        page_break(doc)
    para(doc, "MỤC LỤC", 14, bold=True, align="center")
    add_toc(doc)
    page_break(doc)
    para(doc, "CHƯƠNG TRÌNH ĐÀO TẠO ĐẠI HỌC", 14, bold=True, align="center")
    tot = data.total()
    info_rows = [("Tên chương trình", data.i("ten_chuong_trinh")), ("Trình độ đào tạo", data.i("trinh_do", "Đại học")),
                 ("Ngành đào tạo", data.nganh()), ("Mã ngành đào tạo", data.i("ma_nganh")),
                 ("Loại hình đào tạo", data.i("loai_hinh", "Chính quy tập trung")), ("Thời gian đào tạo", data.i("thoi_gian", "4 năm")),
                 ("Khối lượng kiến thức toàn khóa", f"{g(tot['tc'])} tín chỉ"), ("Đối tượng áp dụng", f"Từ khóa tuyển năm {data.i('khoa_tuyen')}"),
                 ("Thời điểm xây dựng/cập nhật CTĐT", f"{d.month}/{d.year}" if d else PH % "tháng/năm")]
    table(doc, [[a + ":", b] for a, b in info_rows], header=0, size=12, widths=[6, 10], borders=False)
    cdr = data.cdr
    heading(doc, "1. GIỚI THIỆU CHUNG", 1)
    heading(doc, "1.1. Mục tiêu đào tạo", 2)
    para(doc, f"Người học tốt nghiệp ngành {data.nganh()} sẽ có khả năng:")
    table(doc, [["Mục tiêu của chương trình (PEOs)", "Nội dung"]] + [[p["ma"], p["noi_dung"]] for p in cdr.get("peo", [])], widths=[3.5, 12.5], center_cols=(0,))
    heading(doc, "1.2. Vị trí và khả năng làm việc sau tốt nghiệp", 2)
    for grp in cdr.get("vi_tri_viec_lam", []):
        para(doc, grp["nhom"], bold=True)
        for item in grp.get("muc", []):
            para(doc, "– " + item, indent=0.6, align="justify", space_after=2)
    para(doc, cdr.get("hoc_nang_cao", ""), align="justify")
    heading(doc, "1.3. Hình thức và thời gian đào tạo", 2)
    para(doc, f"Hình thức đào tạo: {data.i('loai_hinh', 'Chính quy tập trung').lower()}. Thời gian đào tạo: {data.i('thoi_gian', '4 năm')}.")
    heading(doc, "2. ĐỐI TƯỢNG TUYỂN SINH", 1)
    para(doc, data.i("doi_tuong_tuyen_sinh", "Theo Quy chế tuyển sinh đại học hệ chính quy hiện hành của Bộ Giáo dục và Đào tạo và Trường Đại học Lạc Hồng."), align="justify")
    heading(doc, "3. CHUẨN ĐẦU RA CHƯƠNG TRÌNH ĐÀO TẠO (PLO)", 1)
    para(doc, "Ngay tại thời điểm tốt nghiệp, người học đạt được:")
    rows = [["Chuẩn đầu ra CTĐT (PLOs/SOs)", "Chỉ số đánh giá (PIs)"]]
    pis = defaultdict(list)
    for p in cdr.get("pi", []):
        pis[p["ma"].split(".")[0]].append(p)
    for p in cdr.get("plo", []):
        items = pis[re.sub(r"\D", "", p["ma"])]
        for k, x in enumerate(items or [{"ma": "", "noi_dung": ""}]):
            rows.append([f"{p['ma']}: {p['noi_dung']}" if k == 0 else "^", f"PI{x['ma']}: {x['noi_dung']}" if x["ma"] else ""])
    table(doc, rows, size=11, widths=[7, 9])
    heading(doc, "4. MỐI LIÊN HỆ GIỮA MỤC TIÊU VÀ CHUẨN ĐẦU RA CHƯƠNG TRÌNH ĐÀO TẠO", 1)
    peos = [p["ma"] for p in cdr.get("peo", [])]
    pp = cdr.get("peo_plo", {})
    table(doc, [["PLOs"] + peos] + [[p["ma"]] + ["X" if q in pp.get(p["ma"], []) else "" for q in peos] for p in cdr.get("plo", [])],
          size=11, center_cols=tuple(range(len(peos) + 1)))
    heading(doc, "5. CÁC HỌC PHẦN VÀ MỐI QUAN HỆ VỚI CHUẨN ĐẦU RA", 1)
    cmap = data.map_clo_pi()
    pi_ids = [p["ma"] for p in cdr.get("pi", [])]
    plo_of = [f"PLO{x.split('.')[0]}" for x in pi_ids]
    groups = [("5.1. Các học phần đại cương", ("Lý luận chính trị", "Ngoại ngữ", "Đại cương khác")),
              ("5.2. Các học phần cơ sở ngành", ("Cơ sở ngành",)), ("5.3. Các học phần chuyên ngành", ("Chuyên ngành",)),
              ("5.4. Các học phần khác", ("Chuyên nghiệp khác",)), ("5.5. Thực tập và tốt nghiệp", ("Thực tập", "Tốt nghiệp"))]
    for title, khois in groups:
        cs = [c for c in data.courses if c["khoi"] in khois and c["tc"]]
        if not cs:
            continue
        heading(doc, title, 2)
        h1 = ["Học phần"]
        last = None
        for plo in plo_of:
            h1.append(plo if plo != last else None)
            last = plo
        rows = [h1, ["PIs"] + [f"PI {x}" for x in pi_ids]]
        for c in cs:
            mp = cmap.get(c["ma"], {})
            rows.append([c["ten"]] + ["x" if any(x in s for s in mp.values()) else "" for x in pi_ids])
        table(doc, rows, header=2, size=7, center_cols=tuple(range(1, len(pi_ids) + 1)))
    heading(doc, "6. CHUỖI CÁC MÔN HỌC", 1)
    chuoi = data.info.get("chuoi_mon") or {}
    sems = sorted({c["hoc_ky"] for c in data.courses})
    if chuoi:
        rows = [[""] + [f"HK{s}" for s in sems]]
        for name, codes in chuoi.items():
            row = [name]
            for s in sems:
                row.append("<br>".join(data.by[x]["ten"] for x in codes if x in data.by and data.by[x]["hoc_ky"] == s))
            rows.append(row)
        table(doc, rows, size=8)
    else:
        para(doc, PH % "chuỗi môn theo hướng (khai báo \"chuoi_mon\" trong thong-tin.json)", italic=True)
    heading(doc, "7. CHƯƠNG TRÌNH ĐÀO TẠO", 1)
    heading(doc, "7.1. Tỷ lệ các khối kiến thức", 2)
    para(doc, "Không kể giáo dục thể chất và giáo dục quốc phòng.", italic=True)
    table(doc, khoi_table_rows(data.courses), size=11, widths=[5.5, 6, 2.5, 2], center_cols=(2,))
    k = 1
    for nhom in ("Khối kiến thức giáo dục đại cương", "Khối kiến thức giáo dục chuyên nghiệp", "Tốt nghiệp"):
        k += 1
        heading(doc, f"7.{k}. {nhom if nhom != 'Tốt nghiệp' else 'Khối kiến thức tốt nghiệp'}", 2)
        cs = [c for c in data.courses if KHOI.get(c["khoi"]) == nhom or (nhom.endswith("đại cương") and KHOI.get(c["khoi"], "x") is None)]
        rows = [["STT", "Mã HP", "Tên học phần", "Loại", "TC", "LT", "TH", "BT"]]
        for i, c in enumerate(cs, 1):
            rows.append([str(i), c["ma"], c["ten"], c["khoi"] + (" (tự chọn)" if c["nhom"] else ""), g(c["tc"]), g(c["lt"]), g(c["th"]), g(c["bt"])])
        table(doc, rows, size=10, widths=[1, 1.8, 6.5, 3.2, 1, 1, 1, 1], center_cols=(0, 1, 4, 5, 6, 7))
    heading(doc, "8. KẾ HOẠCH GIẢNG DẠY MẪU", 1)
    for hk, items in data.semesters():
        heading(doc, f"Học kỳ {hk}", 3)
        rows = [["STT", "Mã HP", "Tên học phần", "TC", "LT", "TH", "BT", "Học phần trước"]]
        i = 0
        for it in items:
            if it[0] == "nhom":
                rows.append([f"Tự chọn {it[1]} trong {len(it[2])} học phần", None, None, None, None, None, None, None])
            for c in ([it[1]] if it[0] == "hp" else it[2]):
                i += 1
                rows.append([str(i), c["ma"], c["ten"], g(c["tc"]), g(c["lt"]), g(c["th"]), g(c["bt"]), data.prereq_text(c)])
        t = data.sem_total(items)
        rows.append(["Tổng", None, None, g(t["tc"]), g(t["lt"]), g(t["th"]), g(t["bt"]), ""])
        table(doc, rows, size=10, widths=[1, 1.8, 6.5, 1, 1, 1, 1, 2.7], bold_rows=(len(rows) - 1,), center_cols=(0, 1, 3, 4, 5, 6))
    heading(doc, "9. ĐIỀU KIỆN TỐT NGHIỆP", 1)
    for line in data.info.get("dieu_kien_tot_nghiep", [PH % "điều kiện tốt nghiệp"]):
        para(doc, "– " + line.replace("{tong}", g(tot["tc"])).replace("{khoa}", str(data.i("khoa_tuyen"))), indent=0.6, align="justify")
    heading(doc, "10. MÔ TẢ VẮN TẮT NỘI DUNG VÀ KHỐI LƯỢNG CÁC HỌC PHẦN", 1)
    for c in data.courses:
        if not c["tc"]:
            continue
        heading(doc, f"{c['ma']} – {c['ten']} – {tc_str(c)}", 3)
        para(doc, mo_ta(data, c["ma"]) if not c["ma"].startswith("102") or c["ma"] in data.syl
             else PH % f"mô tả học phần {c['ma']} (đơn vị phụ trách cung cấp)", align="justify")
    heading(doc, "PHỤ LỤC. ĐỀ CƯƠNG CHI TIẾT CÁC HỌC PHẦN", 1)
    para(doc, "Đề cương chi tiết từng học phần kèm theo trong thư mục BM05-De-cuong.", italic=True)
    save_docx(doc, out / f"BM09.2-Chuong-trinh-dao-tao-{data.fpart('ky_hieu_don_vi', 'CNTT')}-{data.fpart('khoa_tuyen')}.docx")


BUILDERS = [("BM01", build_bm01), ("BM02", build_bm02), ("BM03", build_bm03), ("PLO", build_plo), ("DH", build_dh),
            ("BM04", build_bm04), ("BM05", build_bm05), ("TT", build_tt), ("BM06", build_bm06), ("BM07", build_bm07),
            ("BM08", build_bm08), ("BM09", build_bm09)]


def cmd_build(folder, only):
    data = Data(folder)
    out = folder / "ho-so"
    out.mkdir(parents=True, exist_ok=True)
    data.check_dates()
    for name, fn in BUILDERS:
        if only and name not in only:
            continue
        try:
            fn(data, out)
            print(f"[+] {name}")
        except PermissionError as e:
            print(f"[!] {name}: không ghi được (file đang mở trong Word/Excel?): {e.filename}")
    for w in data.warnings:
        print(f"[CẢNH BÁO] {w}")
    rs = folder / "ra-soat"
    rs.mkdir(exist_ok=True)
    v = load_validator()
    sys.argv = ["validate-ctdt.py", str(out), "--report", str(rs / "ra-soat-tu-dong.md")]
    if re.fullmatch(r"\d{4}", str(data.info.get("nam", ""))):
        sys.argv += ["--nam", str(data.info["nam"])]
    import contextlib
    import io
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            v.main()
    except SystemExit:
        pass
    head = [l for l in buf.getvalue().splitlines() if l.startswith("Tổng:")]
    print(f"[i] Rà soát: {head[0] if head else ''} → {rs / 'ra-soat-tu-dong.md'}")


# ================================================================ init

TEMPLATE_INFO = {
    "nganh": "Công nghệ thông tin", "nganh_en": "Information Technology", "ma_nganh": "7480201",
    "trinh_do": "Đại học", "danh_hieu": "{{cử nhân / kỹ sư: đối chiếu quy định}}", "loai_hinh": "Chính quy tập trung", "thoi_gian": "4 năm",
    "ten_chuong_trinh": "Chương trình đào tạo trình độ đại học hệ chính quy",
    "nien_khoa": "{{2027 - 2031}}", "khoa_tuyen": "{{2027}}", "nam": "{{2027}}",
    "don_vi": "Khoa Công nghệ thông tin", "ky_hieu_don_vi": "CNTT", "noi_ky": "Đồng Nai",
    "truong_khoa": "{{họ tên Trưởng khoa}}", "truong_nganh": "{{họ tên Trưởng ngành}}", "nguoi_trinh_bay": "{{người trình bày dự thảo}}",
    "thu_muc_mau": "templates/xay-dung-ctdt-Dai hoc chinh quy",
    "quy_doi_tiet": {},
    "ngay": {f: None for f in THU_TU_NGAY},
    "so_van_ban": {"BM06": "{{số}}/BB-HĐKH", "BM07": "{{số}}/TTr-CNTT"},
    "hop_hoi_dong": {"gio": "{{8h30}}", "dia_diem": "{{Phòng họp Khoa}}", "ket_thuc": "{{11h30}}", "chu_tich": "", "thu_ky": ""},
    "hoi_dong": [{"ho_ten": "{{học vị, họ tên}}", "chuc_vu": "{{chức vụ, đơn vị}}", "vai_tro": "Chủ tịch"}],
    "khach_moi": {"Chuyên gia trong ngành": [], "Doanh nghiệp sử dụng lao động": [], "Cựu sinh viên": []},
    "can_cu": [
        {"van_ban": "Thông tư số 17/2021/TT-BGDĐT của Bộ Giáo dục và Đào tạo quy định về chuẩn chương trình đào tạo; xây dựng, thẩm định và ban hành chương trình đào tạo các trình độ của giáo dục đại học", "ngay": "2021-06-22"},
        {"van_ban": "Thông tư số 56/2026/TT-BGDĐT của Bộ Giáo dục và Đào tạo ban hành Quy chế đào tạo trình độ đại học", "ngay": "2026-07-07"},
        {"van_ban": "Quyết định số 1277/QĐ-ĐHLH của Hiệu trưởng Trường Đại học Lạc Hồng về khối lượng kiến thức tối thiểu, yêu cầu về năng lực và quy định xây dựng, thẩm định, ban hành chương trình đào tạo", "ngay": "2024-10-14"},
        {"van_ban": "{{Quy chế đào tạo của Trường áp dụng cho khóa tuyển, số, ngày}}", "ngay": None},
        {"van_ban": "{{Thông báo/kế hoạch của Hiệu trưởng về đợt xây dựng CTĐT, số, ngày}}", "ngay": None}],
    "nguoi_gop_y": [],
    "chuoi_mon": {},
    "dieu_kien_tot_nghiep": [
        "Tích lũy tối thiểu {tong} tín chỉ và hoàn thành các học phần bắt buộc của chương trình.",
        "Có chứng chỉ Giáo dục quốc phòng và an ninh, hoàn thành Giáo dục thể chất.",
        "Đạt chuẩn đầu ra tiếng Anh: TOEIC từ 400 trở lên hoặc chứng chỉ tương đương.",
        "{{Chuẩn đầu ra tin học, kỹ năng mềm theo quy định của Trường}}",
        "Đáp ứng các điều kiện khác theo quy chế đào tạo trình độ đại học của Trường áp dụng cho khóa tuyển {khoa}."],
    "danh_gia_ctdt": [{"nien_khoa": "{{niên khóa đã tốt nghiệp}}", "trang_thai": "đã tốt nghiệp", "ke_hoach": "ke-hoach-khoa-truoc.csv"}],
}

TEMPLATE_SYL = """---
ma: {ma}
mau: BM05
giang_vien: {{{{họ tên}}}}
email: {{{{email}}}}
giang_vien_tham_gia: {{{{danh sách}}}}
---
<!-- Tên, tín chỉ, LT/TH/BT, khối, môn học trước lấy từ ke-hoach.csv; không ghi lại ở đây. -->
## Mô tả
{{{{Vị trí trong CTĐT, mục đích, nội dung chính; 3–5 câu.}}}}

## Chuẩn đầu ra
| CLO | Nội dung chuẩn đầu ra học phần | Cấp độ (Bloom domain/level) | PLO/SO/PI |
|---|---|---|---|
| CLO1 | {{{{động từ đo được + đối tượng + bối cảnh}}}} | Kiến thức (3) | SO1 (PI1.1) |
| CLO2 | | Kỹ năng (3) | |
| CLO3 | | Thái độ (3) | SO5 (PI5.2) |

## Kế hoạch giảng dạy
| Tuần | Bài/Chương | Tên bài/chương | LLO | Hoạt động dạy và học | Phương pháp giảng dạy | Phương pháp đánh giá | Tài liệu |
|---|---|---|---|---|---|---|---|
| 1 | Bài 1 | | LLO1.1 | | | A1 | [1] |

## Ma trận LLO – CLO
| LLO | CLO1 | CLO2 | CLO3 | Thành phần đánh giá |
|---|---|---|---|---|
| LLO1.1 | X | | | A1 |

## Đánh giá
| Thành phần đánh giá | Hình thức đánh giá | CLOs | Tỷ lệ (%) |
|---|---|---|---|
| A1. Quá trình | | CLO1, CLO2 | 30% |
| A2. Giữa kỳ | | CLO1, CLO2 | 30% |
| A3. Cuối kỳ | | CLO2, CLO3 | 40% |

## Quy định
- Dự lớp: theo quy định của Nhà trường.
- Sinh viên đọc tài liệu học phần trên LMS (learn.lhu.edu.vn) trước mỗi buổi học.
- Liêm chính học thuật: tuân thủ quy định chống đạo văn; mức độ được phép dùng AI trong học phần: {{{{mức độ}}}}.

## Tài liệu
### Giáo trình
- [1] {{{{Tác giả (năm). Tên. Nơi xuất bản: Nhà xuất bản.}}}}
### Tài liệu tham khảo
- [2] {{{{...}}}}

## Phần mềm
- {{{{Hãng (phiên bản) – Tên phần mềm}}}}
"""


def cmd_init(folder, src):
    d = folder / "du-lieu"
    if (d / "thong-tin.json").exists():
        sys.exit(f"[!] {d} đã có dữ liệu; xóa hoặc chọn thư mục khác.")
    (d / "de-cuong").mkdir(parents=True, exist_ok=True)
    (folder / "ho-so").mkdir(exist_ok=True)
    (folder / "ra-soat").mkdir(exist_ok=True)
    if src:
        import_ho_so(folder, Path(src))
        return
    (d / "thong-tin.json").write_text(json.dumps(TEMPLATE_INFO, ensure_ascii=False, indent=2), encoding="utf-8")
    cdr = {"peo": [{"ma": "PEO1", "noi_dung": "{{...}}", "nang_luc": "Năng lực nghề nghiệp"}],
           "plo": [{"ma": "PLO1", "noi_dung": "{{...}}", "mien": "Kiến thức"}],
           "pi": [{"ma": "1.1", "noi_dung": "{{...}}", "target": 0.75}],
           "peo_plo": {"PLO1": ["PEO1"]},
           "vi_tri_viec_lam": [{"nhom": "{{Khối doanh nghiệp ...}}", "mo_ta": "", "muc": ["{{vị trí}}"]}],
           "hoc_nang_cao": "Có khả năng học sau đại học (thạc sĩ, tiến sĩ) tại các cơ sở đào tạo trong và ngoài nước."}
    (d / "chuan-dau-ra.json").write_text(json.dumps(cdr, ensure_ascii=False, indent=2), encoding="utf-8")
    write_plan(d / "ke-hoach.csv", [
        {"hoc_ky": 1, "ma": "111155", "ten": "Nhập môn lập trình", "ten_en": "Introduction to Programming", "tc": 3.0, "lt": 1.0, "th": 2.0, "bt": 0.0,
         "khoi": "Cơ sở ngành", "nhom": "", "hoc_truoc": [], "ghi_chu": ""},
        {"hoc_ky": 2, "ma": "111161", "ten": "Kỹ thuật lập trình", "ten_en": "Programming Techniques", "tc": 3.0, "lt": 1.0, "th": 2.0, "bt": 0.0,
         "khoi": "Cơ sở ngành", "nhom": "", "hoc_truoc": [("111155", "**")], "ghi_chu": ""}])
    with open(d / "doi-sanh.csv", "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows([["truong", "chuong_trinh", "nguon", "nam", "mon_doi_sanh", "ma_lhu"],
                                 ["{{Trường đối sánh}}", "{{Tên CTĐT}}", "{{URL}}", "{{năm}}", "{{Tên học phần}}", "111155"]])
    (d / "de-cuong" / "111155.md").write_text(TEMPLATE_SYL.format(ma="111155"), encoding="utf-8")
    print(f"[+] Đã tạo bộ dữ liệu mẫu tại {d}. Điền theo .claude/skills/ctdt-builder/references/thong-tin-dau-vao.md rồi chạy build.")


# ---------------------------------------------------------------- nhập từ hồ sơ Word/Excel

def docx_seq(path):
    """[('p', text) | ('t', rows)] theo thứ tự; ô bảng nối dòng bằng <br>."""
    doc = Document(path)
    out = []
    for el in doc.element.body.iterchildren():
        if el.tag.endswith('}p'):
            t = nfc(Paragraph(el, doc).text).strip()
            if t:
                out.append(("p", t))
        elif el.tag.endswith('}tbl'):
            rows = []
            for r in Table(el, doc).rows:
                cells, prev = [], None
                for c in r.cells:
                    if c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(nfc("<br>".join(p.text.strip() for p in c.paragraphs if p.text.strip())).replace("|", "/"))
                rows.append(cells)
            out.append(("t", rows))
    return out


def clean(s):
    s = re.sub(r"\s*\bDOCX\b", "", s)
    s = re.sub(r"\((Cần bổ sung[^)]*)\)", r"{{\1}}", s)
    s = re.sub(r"\s*\{\s*(Khóa\s*\d{4}|20\d\d[^}]*|CNTT)\s*\}", "", s)
    return s.strip()


SYL_KEYS = [("MÔ TẢ", "Mô tả"), ("CHUẨN ĐẦU RA", "Chuẩn đầu ra"), ("NỘI DUNG", "Kế hoạch giảng dạy"),
            ("MAPPING", "Ma trận LLO – CLO"), ("MA TRẬN", "Ma trận LLO – CLO"), ("ĐÁNH GIÁ", "Đánh giá"),
            ("QUY ĐỊNH", "Quy định"), ("TÀI LIỆU", "Tài liệu"), ("PHẦN MỀM", "Phần mềm"), ("QUY ƯỚC", None), ("PHÊ DUYỆT", None)]


def syl_heading(text):
    t = re.sub(r"^[\[\(]?\s*([IVX]+|\d+)\s*[.)]\s*", "", text).strip().upper()
    if len(t) > 100:
        return False, None
    for k, name in SYL_KEYS:
        if t.startswith(k):
            return True, name
    return False, None


def md_row(cells):
    return "| " + " | ".join(c.replace("\n", "<br>") for c in cells) + " |"


def import_syllabus(path, code):
    seq = docx_seq(path)
    meta = {"ma": code}
    secs = {}
    cur = None
    info = {}
    for kind, v in seq:
        if kind == "p":
            if re.match(r"^(Trưởng khoa|TRƯỞNG KHOA)", v):
                cur = None
                continue
            is_h, name = syl_heading(v)
            if is_h:
                cur = name
                if cur:
                    secs.setdefault(cur, [])
                continue
            if cur:
                secs[cur].append(("p", v))
        else:
            # tiêu đề mục đặt trong bảng một hàng: | 3. | CHUẨN ĐẦU RA ... |
            if len(v) == 1 and len(v[0]) <= 2 and syl_heading(v[0][-1])[0]:
                cur = syl_heading(v[0][-1])[1]
                if cur:
                    secs.setdefault(cur, [])
                continue
            if cur is None and not info:
                for r in v:
                    if len(r) >= 2:
                        info[r[0].rstrip(":").strip().casefold()] = r[1]
            elif cur:
                secs[cur].append(("t", v))
    meta["ten_en"] = clean(info.get("tên môn học (tiếng anh)", "") or info.get("tên học phần (tiếng anh)", ""))
    gv = info.get("giảng viên phụ trách", "")
    em = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", gv)
    meta["giang_vien"] = clean(re.split(r"[,/]?\s*Email", gv)[0].replace("<br>", " ").strip(" ,/")) or "{{họ tên}}"
    meta["email"] = em.group(0) if em else "{{email}}"
    meta["giang_vien_tham_gia"] = clean(info.get("giảng viên tham gia giảng dạy", "").replace("<br>", ", ")) or "{{danh sách}}"
    text_all = " ".join(v if k == "p" else "" for k, v in seq).upper()
    meta["mau"] = "BM05a" if "TÍCH HỢP TOÀN PHẦN" in text_all else "BM05b" if "TĂNG CƯỜNG TỪ HỌC LIỆU ASU" in text_all else "BM05"
    out = ["---"] + [f"{k}: {v}" for k, v in meta.items()] + ["---",
           f"<!-- Nhập tự động từ {path.name}. Kiểm tra lại từng mục; tên, tín chỉ, môn học trước lấy từ ke-hoach.csv. -->"]

    def paras(name):
        return [clean(v) for k, v in secs.get(name, []) if k == "p"]

    alltab = [v for k, v in seq if k == "t"]

    out += ["## Mô tả"] + [p for p in paras("Mô tả") if not p.startswith(("Sau khi", "Bảng"))][:12] + [""]
    # CLO
    out += ["## Chuẩn đầu ra", "| CLO | Nội dung chuẩn đầu ra học phần | Cấp độ (Bloom domain/level) | PLO/SO/PI |", "|---|---|---|---|"]
    seen = set()
    for t in alltab:
        for r in t:
            m = re.fullmatch(r"CLO\s*(\d+)", r[0].strip()) if r else None
            if not m or m[1] in seen or len(r) < 3 or any(c.startswith("LLO") for c in r[1:2]):
                continue
            seen.add(m[1])
            bloom = next((c for c in r[1:] if re.search(r"(Kiến thức|Kỹ năng|Thái độ|Năng lực)\s*\(?\d", c)), "")
            plo = next((c for c in r[1:] if re.search(r"\b(SO|PLO|PI)\s*\d", c) and len(c) < 60), "")
            rest = [c for c in r[1:] if c not in (bloom, plo)]
            text = max(rest, key=len) if rest else ""
            out.append(md_row([f"CLO{m[1]}", clean(text), bloom, plo]))
    out.append("")
    # Kế hoạch
    out += ["## Kế hoạch giảng dạy",
            "| Tuần | Bài/Chương | Tên bài/chương | LLO | Hoạt động dạy và học | Phương pháp giảng dạy | Phương pháp đánh giá | Tài liệu |",
            "|---|---|---|---|---|---|---|---|"]
    for t in alltab:
        if not t or not any("tuần" in c.casefold() for c in t[0]):
            continue
        for r in t[1:]:
            if not r or not r[0].strip():
                continue
            cells = (r + [""] * 8)[:8]
            out.append(md_row([clean(c) for c in cells]))
    out.append("")
    # Ma trận LLO – CLO
    nclo = len(seen) or 3
    llo = {}
    for _ in (0,):
        for t in alltab:
            hdr = next((r for r in t[:3] if sum(1 for c in r if re.fullmatch(r"CLO\s*\d+", c.strip())) >= 2), None)
            if hdr:
                idx = {i: int(re.sub(r"\D", "", c)) for i, c in enumerate(hdr) if re.fullmatch(r"CLO\s*\d+", c.strip())}
                for r in t:
                    # LLO ở cột đầu, hoặc cột thứ hai khi cột đầu là "Bài/Chương"
                    pos = next((i for i, c in enumerate(r[:2]) if c.startswith("LLO")), None)
                    if pos is not None:
                        cl = set()
                        for i, k in idx.items():
                            for j in (i, i - 1):
                                if pos < j < len(r) and r[j].strip().lower() == "x":
                                    cl.add(k)
                                    break
                        comp = ", ".join(re.findall(r"A\d", r[-1])) if r else ""
                        llo.setdefault(r[pos].split()[0], (cl, comp))
            for r in t:   # dạng "CLO1 | LLO1.1 | nội dung"
                if len(r) >= 2 and re.fullmatch(r"CLO\s*\d+", r[0].strip()) and r[1].startswith("LLO"):
                    llo.setdefault(r[1].split()[0], ({int(re.sub(r'\D', '', r[0]))}, ""))
    out += ["## Ma trận LLO – CLO", md_row(["LLO"] + [f"CLO{k}" for k in range(1, nclo + 1)] + ["Thành phần đánh giá"]),
            "|" + "---|" * (nclo + 2)]
    for name, (cl, comp) in llo.items():
        out.append(md_row([name] + ["X" if k in cl else "" for k in range(1, nclo + 1)] + [comp]))
    out.append("")
    # Đánh giá
    out += ["## Đánh giá", "| Thành phần đánh giá | Hình thức đánh giá | CLOs | Tỷ lệ (%) |", "|---|---|---|---|"]
    for t in alltab:
        for r in t:
            if r and re.match(r"^A\d+\s*[.:]", r[0].strip()):
                pct = re.fullmatch(r"\s*(\d+(?:[.,]\d+)?)\s*%?\s*", r[-1])
                pct = pct.group(1) if pct else (re.findall(r"(\d+)\s*%", " ".join(r)) or ["{{%}}"])[-1]
                clos = ", ".join(f"CLO{x}" for x in sorted(set(int(x) for x in re.findall(r"CLO\s*(\d+)", " ".join(r[1:])))))
                comp = re.sub(r"\s*\(\s*\d+\s*%\s*\)", "", r[0]).strip()
                out.append(md_row([comp, clean(r[1]) if len(r) > 1 else "", clos, f"{pct}%"]))
    out.append("")
    out += ["## Quy định"] + [f"- {p.lstrip('-–•+ ').strip()}" for p in paras("Quy định")] + [""]
    out += ["## Tài liệu"]
    for p in paras("Tài liệu"):
        if re.fullmatch(r"\$*\d+\$*", p):
            continue
        if re.match(r"^(Giáo trình|Tài liệu tham khảo|Tài liệu học tập)\s*:?$", p, re.I):
            out.append(f"### {p.rstrip(':')}")
        elif re.match(r"^(Phần mềm|PHẦN MỀM)", p):
            break
        else:
            out.append(f"- {p.lstrip('-–•+ ').strip()}")
    out += ["", "## Phần mềm"] + [f"- {p.lstrip('-–•+ ').strip()}" for p in paras("Phần mềm") if not p.startswith("Quy ước")] + [""]
    return "\n".join(out)


def import_ho_so(folder, src):
    v = load_validator()
    d = folder / "du-lieu"
    report = []
    bm03 = sorted(src.glob("BM03*.xlsx"))
    if not bm03:
        sys.exit("[!] Không thấy BM03*.xlsx trong thư mục hồ sơ")
    plan, _, _, _ = v.load_xlsx_plan(bm03[0])
    docs = [p for p in src.glob("*.docx") if not p.name.startswith("~$")]
    dh = next((p for p in docs if "điều hành" in nfc(p.name).casefold()), None)
    pre = {}
    if dh:
        for t in v.docx_blocks(dh)[1]:
            cs, _, _ = v.parse_plan_rows(t, dh.name)
            if cs:
                pre = {c.code: c.prereq for c in cs}
                break
    # khối kiến thức: theo các bảng khối trong BM09.2, dự phòng bằng cột ghi chú của file PLO
    khoi = {}
    bm09 = next((p for p in docs if nfc(p.name).startswith("BM09")), None)
    if bm09:
        state = None
        for kind, val in docx_seq(bm09):
            if kind == "p":
                t = val.casefold()
                if "kế hoạch giảng dạy mẫu" in t:
                    break
                if t.startswith(("khối kiến thức giáo dục đại cương",)):
                    state = "DC"
                elif "cơ sở ngành" in t and t.startswith("nhóm"):
                    state = "Cơ sở ngành"
                elif "chuyên ngành" in t and t.startswith("nhóm"):
                    state = "Chuyên ngành"
                elif t.startswith("nhóm các học phần khác"):
                    state = "Chuyên nghiệp khác"
                elif t.startswith(("thực tập", "nhóm các học phần đồ án")):
                    state = "Thực tập"
                elif t.startswith(("khóa luận", "nhóm các học phần chuyên đề")):
                    state = "Tốt nghiệp"
            elif state:
                for r in val:
                    for cell in r[:3]:
                        if re.fullmatch(r"\d{6}", cell.strip()):
                            khoi.setdefault(cell.strip(), state)
    plo_x = sorted(src.glob("PLO*.xlsx"))
    plo_note = {}
    if plo_x:
        wb = openpyxl.load_workbook(plo_x[0], data_only=True)
        ws = next((wb[s] for s in wb.sheetnames if "clo" in s.casefold()), None)
        if ws:
            for r in ws.iter_rows(values_only=True):
                m = v.CODE_RE.match(str(r[0] or ""))
                if m and len(r) > 8:
                    plo_note[m.group(1)] = str(r[8] or "")
    syl_files = [p for dd in src.glob("BM05*") if dd.is_dir() for p in sorted(dd.glob("*.docx"))
                 if re.search(r"(?<!\d)1\d{5}(?!\d)", p.stem) and not p.name.startswith("~$")]
    syl_info = {}
    for p in syl_files:
        try:
            s = v.parse_syllabus(p)
        except Exception:
            continue
        fm = re.search(r"(?<!\d)(1\d{5})(?!\d)", p.stem)
        # tên file và nội dung mang hai mã khác nhau: chọn mã có trong BM03, ưu tiên mã theo tên file
        in_plan = {c.code for c in plan}
        cands = [x for x in (fm.group(1) if fm else None, s["code"]) if x]
        code = next((x for x in cands if x in in_plan), cands[0] if cands else None)
        if fm and s["code"] and s["code"] != fm.group(1):
            report.append(f"{p.name}: tên file mang mã {fm.group(1)}, nội dung mang mã {s['code']}; nhập vào {code}, CẦN KIỂM TRA nội dung")
        syl_info.setdefault(code, []).append((p, s))

    def khoi_of(c):
        name = c.name.casefold()
        if "thể chất" in name:
            return "Giáo dục thể chất"
        if "quốc phòng" in name:
            return "Giáo dục quốc phòng"
        k = khoi.get(c.code)
        if k == "DC" or (k is None and c.code.startswith("102")):
            if re.search(r"mác|hồ chí minh|lịch sử đảng|xã hội khoa học", name):
                return "Lý luận chính trị"
            if "ngoại ngữ" in name or "anh văn" in name:
                return "Ngoại ngữ"
            return "Đại cương khác"
        if k:
            return k
        note = plo_note.get(c.code, "").casefold()
        if "cơ sở" in note:
            return "Cơ sở ngành"
        if "thực tập" in note:
            return "Tốt nghiệp" if "tốt nghiệp" == name else "Thực tập"
        report.append(f"{c.code} {c.name}: không xác định được khối kiến thức, tạm ghi Chuyên ngành")
        return "Chuyên ngành"

    rows = []
    groups = defaultdict(list)
    for c in plan:
        if c.group:
            groups[c.group].append(c)
    for c in plan:
        note = []
        name = re.sub(r"\s*\{[^}]*\}", "", c.name).strip()
        if name != c.name.strip():
            note.append(f"tên gốc trong BM03: {c.name.strip()}")
        tc, lt, th, bt = c.tc, c.lt, c.th, c.bt
        if c.group and tc is None:
            sib = next((x for x in groups[c.group] if x.tc), None)
            if sib:
                tc, lt, th, bt = sib.tc, sib.lt, sib.th, sib.bt
                note.append("{{CẦN KIỂM TRA: BM03 bỏ trống tín chỉ, chép từ môn cùng nhóm tự chọn}}")
        ten_en = ""
        for p, _ in syl_info.get(c.code, [])[:1]:
            for kind, val in docx_seq(p)[:12]:
                if kind == "t":
                    for r in val:
                        if len(r) > 1 and r[0].casefold().startswith(("tên môn học (tiếng anh)", "tên học phần (tiếng anh)")):
                            ten_en = clean(r[1])
        rows.append({"hoc_ky": c.sem, "ma": c.code, "ten": name, "ten_en": ten_en, "tc": tc, "lt": lt or 0.0, "th": th or 0.0, "bt": bt or 0.0,
                     "khoi": khoi_of(c), "nhom": f"HK{c.sem}-{c.group[0]}" + (f":{c.group[1]}" if c.group[1] != 1 else "") if c.group else "",
                     "hoc_truoc": [(a, {"tiên quyết": "*", "học trước": "**", "song hành": "***"}[k]) for a, k in pre.get(c.code, [])],
                     "ghi_chu": "; ".join(note)})
    write_plan(d / "ke-hoach.csv", rows)
    # đề cương
    codes = {r["ma"] for r in rows}
    for code, items in syl_info.items():
        if code not in codes:
            report.append(f"Đề cương {items[0][0].name}: mã {code} không có trong BM03, không nhập")
            continue
        if len(items) > 1:
            report.append(f"{code}: {len(items)} file đề cương ({'; '.join(p.name for p, _ in items)}), nhập file đầu tiên")
        (d / "de-cuong" / f"{code}.md").write_text(import_syllabus(items[0][0], code), encoding="utf-8")
    for r in rows:
        if r["ma"] not in syl_info and not r["ma"].startswith("102") and r["tc"]:
            (d / "de-cuong" / f"{r['ma']}.md").write_text(TEMPLATE_SYL.format(ma=r["ma"]), encoding="utf-8")
            report.append(f"{r['ma']} {r['ten']}: chưa có đề cương, đã tạo khung trống")
    # chuẩn đầu ra
    cdr = {"peo": [], "plo": [], "pi": [], "peo_plo": {}, "vi_tri_viec_lam": [], "hoc_nang_cao": ""}
    if plo_x:
        ws = openpyxl.load_workbook(plo_x[0], data_only=True)["PLO"]
        for r in ws.iter_rows(values_only=True):
            a = str(r[0] or "").strip()
            if re.fullmatch(r"PEO\d+", a):
                cdr["peo"].append({"ma": a, "noi_dung": nfc(r[1]).strip(), "nang_luc": nfc(r[2] or "")})
            elif re.fullmatch(r"PLO\d+", a):
                cdr["plo"].append({"ma": a, "noi_dung": nfc(r[1]).strip(), "mien": nfc(r[2] or "")})
            elif re.match(r"PI\s*\d", a):
                cdr["pi"].append({"ma": re.sub(r"PI\s*", "", a), "noi_dung": nfc(r[1]).strip(), "target": r[2]})
        cdr["plo"].sort(key=lambda p: p["ma"])
    if bm09:
        for kind, val in docx_seq(bm09):
            if kind == "t" and val and any(re.fullmatch(r"PEO\d", c) for r in val[:2] for c in r):
                hdr = next(r for r in val[:2] if any(re.fullmatch(r"PEO\d", c) for c in r))
                for r in val:
                    m = re.fullmatch(r"(SO|PLO)(\d)", r[0].strip()) if r else None
                    if m:
                        cdr["peo_plo"][f"PLO{m[2]}"] = [hdr[i] for i in range(1, min(len(r), len(hdr))) if r[i].strip().lower() == "x"]
                break
    bm02 = next((p for p in docs if nfc(p.name).startswith("BM02")), None)
    info = json.loads(json.dumps(TEMPLATE_INFO))
    if bm02:
        seq = docx_seq(bm02)
        on = False
        for kind, val in seq:
            if kind != "p":
                continue
            if val.startswith("Tiếng Anh:"):
                info["nganh_en"] = val.split(":", 1)[1].strip().title()
            if val.startswith("Vị trí làm việc"):
                on = True
                continue
            if val.startswith("Khả năng học tập"):
                on = False
                continue
            if on:
                if re.match(r"^(\d+(\.\d+)?\.?\s+)?Khối\b", val):
                    cdr["vi_tri_viec_lam"].append({"nhom": re.sub(r"^\d+(\.\d+)?\.?\s*", "", val), "mo_ta": "", "muc": []})
                elif cdr["vi_tri_viec_lam"]:
                    grp = cdr["vi_tri_viec_lam"][-1]
                    if val.startswith("Bao gồm") and not grp["muc"]:
                        grp["mo_ta"] = val
                    else:
                        grp["muc"].append(val)
            if val.startswith("Có khả năng học sau đại học"):
                cdr["hoc_nang_cao"] = val
    (d / "chuan-dau-ra.json").write_text(json.dumps(cdr, ensure_ascii=False, indent=2), encoding="utf-8")
    # thông tin chung, Hội đồng, căn cứ từ BM06
    info["nien_khoa"], info["khoa_tuyen"], info["nam"] = "{{niên khóa}}", "{{khóa tuyển}}", "{{năm}}"
    bm06 = next((p for p in docs if nfc(p.name).startswith("BM06")), None)
    if bm06:
        seq = docx_seq(bm06)
        info["hoi_dong"], info["khach_moi"], info["can_cu"] = [], {}, []
        label = None
        for kind, val in seq:
            if kind == "p":
                m = re.search(r"niên khóa\s*(\d{4})\s*[-–]\s*(\d{4})", val)
                if m and info["nien_khoa"].startswith("{{"):
                    info["nien_khoa"], info["khoa_tuyen"], info["nam"] = f"{m[1]} - {m[2]}", m[1], int(m[1])
                if val.startswith("Căn cứ") and len(val) > 20 and not val.startswith("Căn cứ vào các ý kiến") and not val.startswith("Căn cứ vào báo cáo"):
                    dm = re.search(r"ngày\s*(\d{1,2})\s*/\s*(\d{1,2})\s*/\s*(\d{4})", val)
                    txt = re.sub(r"^Căn cứ (vào )?", "", val).rstrip(";.")
                    if dm:
                        txt = txt[:dm.start()].strip() + " " + txt[dm.end():].strip()
                    info["can_cu"].append({"van_ban": re.sub(r"\s+", " ", txt).strip(),
                                           "ngay": f"{dm[3]}-{int(dm[2]):02d}-{int(dm[1]):02d}" if dm else None})
                if val.startswith("Trưởng khoa:"):
                    info["truong_khoa"] = val.split(":", 1)[1].strip()
                if val.casefold().startswith("trưởng ngành:"):
                    info["truong_nganh"] = val.split(":", 1)[1].strip()
                if len(val) < 80 and re.match(r"^(Chuyên gia|Doanh nghiệp|Cựu sinh viên)", val):
                    label = val.split(":")[0].strip()
            elif kind == "t":
                if val and len(val[0]) == 3 and any(r[2] in ("Chủ tịch", "Phó chủ tịch") for r in val if len(r) > 2):
                    info["hoi_dong"] = [{"ho_ten": r[0], "chuc_vu": r[1].replace("<br>", " "), "vai_tro": r[2]} for r in val if len(r) >= 3]
                elif label and val and len(val[0]) >= 2 and not any("TỔNG CỘNG" in " ".join(r) for r in val):
                    people = []
                    for r in val:
                        names = r[0].split("<br>")
                        roles = (r[1].split("<br>") + [""] * len(names))[:len(names)]
                        people += [{"ho_ten": n.strip(), "chuc_vu": ro.strip()} for n, ro in zip(names, roles) if n.strip()]
                    info["khach_moi"][label] = people
                    label = None
        info["hop_hoi_dong"]["chu_tich"] = next((m["ho_ten"] for m in info["hoi_dong"] if m["vai_tro"] == "Chủ tịch"), "")
        info["hop_hoi_dong"]["thu_ky"] = next((m["ho_ten"] for m in info["hoi_dong"] if "Thư ký" in m["vai_tro"]), "")
    people = []
    for sub, loai in (("BM1.1*", "BM01.1"), ("BM1.2*", "BM01.2")):
        for dd in src.glob(sub):
            for p in sorted(dd.glob("*.docx")):
                name = unit = ""
                for kind, val in docx_seq(p)[:12]:
                    if kind == "p" and val.startswith("Họ và tên"):
                        name = val.split(":", 1)[1].strip()
                    if kind == "p" and val.startswith("Đơn vị công tác"):
                        unit = val.split(":", 1)[1].strip()
                if name:
                    people.append({"ho_ten": name, "don_vi": unit, "loai": loai})
    info["nguoi_gop_y"] = people
    info["danh_gia_ctdt"] = []
    info["ngay"] = {f: None for f in THU_TU_NGAY}
    (d / "thong-tin.json").write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding="utf-8")
    # đối sánh BM04
    bm04 = next((p for p in docs if nfc(p.name).startswith("BM04")), None)
    if bm04:
        names = {v.norm(r["ten"]): r["ma"] for r in rows}
        out_rows = [["truong", "chuong_trinh", "nguon", "nam", "mon_doi_sanh", "ma_lhu", "ghi_chu"]]
        school = ""
        for kind, val in docx_seq(bm04):
            if kind == "p" and val.upper().startswith("VỚI CHƯƠNG TRÌNH ĐÀO TẠO CỦA"):
                school = re.sub(r"^VỚI CHƯƠNG TRÌNH ĐÀO TẠO CỦA\s*", "", val, flags=re.I).strip().title()
            if kind == "t" and school:
                for r in val[2:]:
                    if len(r) < 2 or not r[0].strip() or r[0].startswith("Tổng"):
                        continue
                    eq = [x.strip() for x in re.split(r"<br>|\+", r[1]) if x.strip() and x.strip() != "Không có"]
                    codes = [names.get(v.norm(x), "") for x in eq]
                    miss = [x for x, c in zip(eq, codes) if not c]
                    out_rows.append([school, "{{tên CTĐT}}", "{{URL}}", "{{năm}}", clean(r[0]), "; ".join(c for c in codes if c),
                                     ("không khớp tên: " + "; ".join(miss)) if miss else ""])
        with open(d / "doi-sanh.csv", "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(out_rows)
    (folder / "ra-soat").mkdir(exist_ok=True)
    note = folder / "ra-soat" / "nhap-du-lieu.md"
    note.write_text("# Ghi chú nhập dữ liệu từ hồ sơ\n\nNguồn: `" + str(src) + "`\n\n" + "\n".join(f"- {r}" for r in report) + "\n", encoding="utf-8")
    print(f"[+] Đã nhập dữ liệu vào {d} ({len(rows)} học phần, {len(list((d / 'de-cuong').glob('*.md')))} đề cương). Ghi chú: {note}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a1 = sub.add_parser("init", help="tạo bộ dữ liệu cho một CTĐT mới")
    a1.add_argument("folder", type=Path)
    a1.add_argument("--tu-ho-so", help="nhập dữ liệu từ thư mục hồ sơ Word/Excel đã có")
    a2 = sub.add_parser("build", help="sinh hồ sơ vào ho-so/ và rà soát")
    a2.add_argument("folder", type=Path)
    a2.add_argument("--chi", help="chỉ sinh một số biểu mẫu, ví dụ BM03,BM05,PLO")
    a = ap.parse_args()
    folder = a.folder if a.folder.is_absolute() else Path.cwd() / a.folder
    if a.cmd == "init":
        cmd_init(folder, a.tu_ho_so)
    else:
        cmd_build(folder, set(a.chi.upper().split(",")) if a.chi else None)


if __name__ == "__main__":
    main()
