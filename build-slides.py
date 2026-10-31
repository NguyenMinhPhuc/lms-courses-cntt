"""
Sinh slide PowerPoint (Slides/bai-XX.pptx) từ ghi chú slide Markdown (Slides/bai-XX.markdown),
dựng trên mẫu templates/Template-slide.pptx (logo LHU, khung tiêu đề, slide cảm ơn).

Cách dùng:
    uv run --with python-pptx python build-slides.py nhap-mon-lap-trinh          # mọi bài của học phần
    uv run --with python-pptx python build-slides.py nhap-mon-lap-trinh 03       # chỉ bài 03

Định dạng Markdown đầu vào: xem .claude/skills/lms-lecture-builder/references/slides.md
"""
import re
import sys
import copy
import json
import math
import argparse
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.parts.image import Image
from pptx.util import Inches, Pt

# Fix Windows console encoding for Vietnamese unicode characters
if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if sys.stderr and hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

ROOT_DIR = Path(__file__).resolve().parent
TEMPLATE_FILE = ROOT_DIR / "templates" / "Template-slide.pptx"
SLIDE_FILE_RE = re.compile(r"^bai-(\d{2})\.markdown$")
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(\s*<?([^)>]+?)>?\s*\)$")

# Tên layout trong Template-slide.pptx
LAYOUT_TITLE = "TITLE"
LAYOUT_CONTENT = "OBJECT"

BLUE = RGBColor(0x00, 0x33, 0xCC)
TEXT = RGBColor(0x21, 0x25, 0x29)
MUTED = RGBColor(0x6C, 0x75, 0x7D)
CODE_BG = RGBColor(0xF1, 0xF3, 0xF5)
PLACEHOLDER_BG = RGBColor(0xF8, 0xF9, 0xFA)
CODE_FONT = "Consolas"

# Vùng nội dung của layout OBJECT (inch): dưới thanh tiêu đề, trên đường cong cuối slide
BODY_LEFT, BODY_TOP, BODY_WIDTH, BODY_BOTTOM = 0.35, 1.2, 12.6, 6.05
IMAGE_SIDE_LEFT, IMAGE_SIDE_WIDTH = 6.85, 6.1   # cột hình khi slide có cả chữ và hình
IMAGE_BOTTOM = 5.45                             # hình không đè lên đường cong trang trí góc phải
CAPTION_H = 0.4
NOTES_LABEL = "ghi chú giảng viên"
HIDDEN_LABELS = {"nội dung chính"}              # nhãn mặc định, không cần hiện trên slide


# ---------------------------------------------------------------- parse
def parse_markdown(text: str) -> tuple[str, str, list[dict]]:
    """Trả về (tiêu đề bài, phụ đề, danh sách slide). Mỗi slide: {kind, title, blocks, notes}.
    kind = "content" (## ...) hoặc "section" (# ... sau tiêu đề bài).
    block = ("para", level, text, is_label), ("code", text) hoặc ("image", alt, path)."""
    deck_title, subtitle, slides = "", "", []
    current = None
    in_notes = in_code = False
    code_lines: list[str] = []

    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if in_code:
            if line.strip().startswith("```"):
                current["blocks"].append(("code", "\n".join(code_lines)))
                in_code, code_lines = False, []
            else:
                code_lines.append(raw.rstrip("\n"))
            continue

        if line.startswith("# "):
            if not deck_title:
                deck_title = line[2:].strip()
            else:
                current = {"kind": "section", "title": line[2:].strip(), "blocks": [], "notes": []}
                slides.append(current)
                in_notes = False
            continue
        if line.startswith("## "):
            title = re.sub(r"^Slide\s*\d+\s*[:.\-–]\s*", "", line[3:].strip(), flags=re.IGNORECASE)
            current = {"kind": "content", "title": title, "blocks": [], "notes": []}
            slides.append(current)
            in_notes = False
            continue
        if current is None:
            if line.strip() and not subtitle and line.strip() != "---":
                subtitle = line.strip()
            continue
        if line.strip() == "---":
            in_notes = False
            continue
        if line.strip().startswith("```") and not in_notes:
            in_code = True
            continue

        label = re.match(r"^\*\*(.+?)\*\*\s*:\s*(.*)$", line.strip())
        if label and label.group(1).strip().lower() == NOTES_LABEL:
            in_notes = True
            if label.group(2).strip():
                current["notes"].append(label.group(2).strip())
            continue
        if in_notes:
            if line.strip():
                current["notes"].append(re.sub(r"^\s*[-*]\s+", "", line).strip())
            continue
        if not line.strip():
            continue

        image = IMAGE_RE.match(line.strip())
        if image:
            current["blocks"].append(("image", image.group(1).strip(), image.group(2).strip()))
            continue
        if label:
            name, rest = label.group(1).strip(), label.group(2).strip()
            if name.lower() in HIDDEN_LABELS and not rest:
                continue
            content = f"**{name}**" + (f": {rest}" if rest else "")
            current["blocks"].append(("para", 0, content, True))
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if bullet:
            level = len(bullet.group(1).replace("\t", "  ")) // 2
            current["blocks"].append(("para", level, bullet.group(2).strip(), False))
        else:
            current["blocks"].append(("para", 0, line.strip(), False))
    return deck_title, subtitle, slides


# ---------------------------------------------------------------- helpers
def add_runs(paragraph, text: str, size: int, color: RGBColor | None = None, bold_all: bool = False):
    """Ghi text vào paragraph, hỗ trợ **in đậm** và `code` trong dòng. Font chữ thường lấy theo template."""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith("**") and part.endswith("**"):
            run.text, run.font.bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            run.text, run.font.name = part[1:-1], CODE_FONT
        else:
            run.text = part
            if bold_all:
                run.font.bold = True
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def set_text(shape, text: str, size: int | None = None, color: RGBColor | None = None):
    """Thay toàn bộ chữ của placeholder, giữ định dạng kế thừa từ layout."""
    frame = shape.text_frame
    for p in list(frame.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = frame.paragraphs[0]
    for r in list(first._p):
        if r.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            first._p.remove(r)
    if size is None:
        first.add_run().text = text
    else:
        add_runs(first, text, size, color)


def set_bullet(paragraph, char: str | None):
    """char=None: bỏ ký hiệu đầu dòng (dòng nhãn)."""
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if char is None:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))
    else:
        etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def plain_len(text: str) -> int:
    return len(re.sub(r"\*\*|`", "", text))


def fit_size(paras, width_in: float, height_in: float, sizes=(26, 24, 22, 20, 18, 16, 14)) -> int:
    """Chọn cỡ chữ lớn nhất mà phần chữ (ước lượng số dòng sau khi xuống dòng) vừa khung."""
    for size in sizes:
        chars_per_line = width_in * 72 / (size * 0.5)
        lines = sum(max(1, math.ceil(plain_len(t) / max(chars_per_line - 4 * lv - 3, 10))) for _, lv, t, _ in paras)
        if lines * size * 1.3 / 72 + len(paras) * 6 / 72 <= height_in:
            return size
    return sizes[-1]


def title_size(text: str, steps) -> int:
    """steps: [(số ký tự tối đa, cỡ chữ), ...] theo thứ tự tăng dần, phần tử cuối là mặc định."""
    for limit, size in steps:
        if len(text) <= limit:
            return size
    return steps[-1][1]


def placeholder(slide, idx: int):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def add_slide_number(slide):
    """add_slide() không chép placeholder số trang của layout, nên chép thủ công."""
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            sp = copy.deepcopy(ph._element)
            sp.nvSpPr.cNvPr.id = slide.shapes._next_shape_id
            slide.shapes._spTree.append(sp)


def get_layout(prs, name: str):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    raise ValueError(f"Template thiếu layout '{name}'")


# ---------------------------------------------------------------- render
def title_slide(prs, title: str, subtitle: str, tagline: str = ""):
    slide = prs.slides.add_slide(get_layout(prs, LAYOUT_TITLE))
    set_text(placeholder(slide, 0), title, title_size(title, [(28, 54), (50, 44), (80, 36), (999, 30)]))
    sub = placeholder(slide, 1)
    if subtitle:
        set_text(sub, subtitle, title_size(subtitle, [(40, 24), (60, 20), (999, 16)]))
    else:
        remove_shape(sub)
    if tagline:
        box = slide.shapes.add_textbox(Inches(2.72), Inches(5.35), Inches(7.9), Inches(0.45))
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, tagline, 18, MUTED)
        p.runs[0].font.italic = True
    return slide


def add_image(slide, md_dir: Path, alt: str, path: str, left, top, width, height) -> bool:
    """Đặt hình vừa khung (giữ tỉ lệ) kèm chú thích. Thiếu file thì vẽ khung chờ để giảng viên chèn sau."""
    file = (md_dir / path).resolve()
    ok = not re.match(r"^https?://", path, re.IGNORECASE) and file.is_file()
    img_h = height - Inches(CAPTION_H) if alt else height
    if ok:
        try:
            px_w, px_h = Image.from_file(str(file)).size
        except Exception:
            ok = False
    if ok:
        scale = min(width / px_w, img_h / px_h)
        w, h = int(px_w * scale), int(px_h * scale)
        img_top = top + (img_h - h) // 2
        slide.shapes.add_picture(str(file), left + (width - w) // 2, img_top, w, h)
        img_h = img_top + h - top
    else:
        rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, img_h)
        rect.fill.solid()
        rect.fill.fore_color.rgb = PLACEHOLDER_BG
        rect.line.color.rgb = MUTED
        rect.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        rect.shadow.inherit = False
        frame = rect.text_frame
        frame.word_wrap = True
        p = frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, f"[Chèn hình] {alt or path}", 14, MUTED)
        p2 = frame.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        add_runs(p2, path, 11, MUTED)
    if alt:
        box = slide.shapes.add_textbox(left, top + img_h, width, Inches(CAPTION_H))
        box.text_frame.word_wrap = True
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, alt, 12, MUTED)
        for r in p.runs:
            r.font.italic = True
    return ok


def content_slide(prs, data, header: str, md_dir: Path) -> list[str]:
    """Trả về danh sách đường dẫn hình không tìm thấy."""
    slide = prs.slides.add_slide(get_layout(prs, LAYOUT_CONTENT))
    set_text(placeholder(slide, 0), header, title_size(header, [(40, 30), (60, 24), (999, 20)]))
    set_text(placeholder(slide, 2), data["title"], title_size(data["title"], [(55, 22), (75, 18), (999, 16)]))
    add_slide_number(slide)

    paras = [b for b in data["blocks"] if b[0] == "para"]
    codes = [b for b in data["blocks"] if b[0] == "code"]
    images = [b for b in data["blocks"] if b[0] == "image"][:3]
    missing = []

    text_w = (IMAGE_SIDE_LEFT - BODY_LEFT - 0.25) if images and (paras or codes) else BODY_WIDTH
    body_h = BODY_BOTTOM - BODY_TOP
    code_lines = sum(c[1].count("\n") + 1 for c in codes)
    code_size = 18 if code_lines <= 8 else 16 if code_lines <= 12 else 14 if code_lines <= 16 else 12
    longest = max((len(l) for c in codes for l in c[1].split("\n")), default=0)
    while code_size > 12 and longest * code_size * 0.55 / 72 > text_w - 0.5:  # tránh dòng code bị ngắt
        code_size -= 2
    code_h = (code_size * 1.25 / 72) * code_lines + 0.3 * len(codes) + 0.15 * len(codes)
    text_h = max(1.0, body_h - code_h)

    body = placeholder(slide, 1)
    if paras:
        size = fit_size(paras, text_w - 0.6, text_h)
        if codes:  # khối code nằm ngay dưới phần chữ thay vì dồn xuống cuối slide
            text_h = min(text_h, len(paras) * size * 1.6 / 72 + 0.3)
        body.left, body.top, body.width, body.height = Inches(BODY_LEFT), Inches(BODY_TOP), Inches(text_w), Inches(text_h)
        frame = body.text_frame
        for i, (_, level, text, is_label) in enumerate(paras):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.level = min(level, 4)
            p.space_before = Pt(8 if is_label and i > 0 else 4)
            set_bullet(p, None if is_label else ("❖" if level == 0 else "–"))
            add_runs(p, text, max(size - 2 * level, 12), BLUE if is_label else None)
    else:
        remove_shape(body)

    top = Inches(BODY_TOP + (text_h + 0.1 if paras else 0))
    for _, code in codes:
        height = Inches((code_size * 1.25 / 72) * (code.count("\n") + 1) + 0.3)
        rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(BODY_LEFT + 0.1), top, Inches(text_w - 0.2), height)
        rect.fill.solid()
        rect.fill.fore_color.rgb = CODE_BG
        rect.line.fill.background()
        rect.shadow.inherit = False
        frame = rect.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = frame.margin_top = Inches(0.15)
        for i, code_line in enumerate(code.split("\n")):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = code_line or " "
            run.font.name, run.font.size, run.font.color.rgb = CODE_FONT, Pt(code_size), TEXT
        top += height + Inches(0.15)

    if images:
        gap, img_area_h = 0.2, IMAGE_BOTTOM - BODY_TOP
        if paras or codes:   # cột hình bên phải, các hình xếp dọc
            left, width = IMAGE_SIDE_LEFT, IMAGE_SIDE_WIDTH
            cell_h = (img_area_h - gap * (len(images) - 1)) / len(images)
            cells = [(left, BODY_TOP + i * (cell_h + gap), width, cell_h) for i in range(len(images))]
        else:                # chỉ có hình: xếp ngang, chiếm cả vùng nội dung
            cell_w = (BODY_WIDTH - gap * (len(images) - 1)) / len(images)
            cells = [(BODY_LEFT + i * (cell_w + gap), BODY_TOP, cell_w, img_area_h) for i in range(len(images))]
        for (_, alt, path), (x, y, w, h) in zip(images, cells):
            if not add_image(slide, md_dir, alt, path, Inches(x), Inches(y), Inches(w), Inches(h)):
                missing.append(path)

    if data["notes"]:
        slide.notes_slide.notes_text_frame.text = "\n".join(data["notes"])
    return missing


def fill_closing_slide(slide, config: dict, course_line: str):
    """Slide cảm ơn của template: tên, email, website giảng viên; ngày và chân trang."""
    email = config.get("instructorEmail", "")
    website = config.get("website") or (config.get("faculty") or {}).get("website") or "lhu.edu.vn"
    semester = " • ".join(str(x) for x in (config.get("semester"), config.get("year")) if x)
    values = {1: config.get("instructor", ""), 3: email, 4: website, 10: semester, 11: course_line}
    for idx, value in values.items():
        ph = placeholder(slide, idx)
        if ph is not None:
            set_text(ph, value)
    if email:
        placeholder(slide, 3).text_frame.paragraphs[0].runs[0].hyperlink.address = f"mailto:{email}"
    else:  # không có email thì bỏ cả dòng email và biểu tượng của nó
        for idx in (3, 6):
            if placeholder(slide, idx) is not None:
                remove_shape(placeholder(slide, idx))


def build_deck(md_file: Path, out_file: Path, config: dict) -> tuple[int, list[str]]:
    deck_title, subtitle, slides = parse_markdown(md_file.read_text(encoding="utf-8"))
    if not slides:
        raise ValueError(f"{md_file.name}: không tìm thấy slide nào (mỗi slide bắt đầu bằng '## ')")
    template = ROOT_DIR / config["slideTemplate"] if config.get("slideTemplate") else TEMPLATE_FILE
    prs = Presentation(str(template))

    # Bỏ các slide mẫu, chỉ giữ slide cuối (cảm ơn) để chuyển xuống cuối bộ slide
    sld_ids = prs.slides._sldIdLst
    closing = prs.slides[len(prs.slides) - 1]
    closing.part.partname = PackURI("/ppt/slides/slide-closing.xml")  # tránh trùng tên với slide mới
    for sld_id in list(sld_ids)[:-1]:
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)

    name, code = config.get("name", ""), config.get("code", "")
    course_line = f"{name} ({code})" if name and code else name or code
    deck_title = deck_title or md_file.stem
    title_slide(prs, deck_title, course_line, subtitle)
    missing, header = [], deck_title
    for data in slides:
        if data["kind"] == "section":
            header = data["title"]
            section = title_slide(prs, data["title"], deck_title)
            if data["notes"]:
                section.notes_slide.notes_text_frame.text = "\n".join(data["notes"])
        else:
            missing += content_slide(prs, data, header, md_file.parent)

    closing_id = list(sld_ids)[0]
    sld_ids.remove(closing_id)
    sld_ids.append(closing_id)
    fill_closing_slide(closing, config, course_line)
    prs.save(out_file)
    return len(prs.slides), missing


def main():
    parser = argparse.ArgumentParser(description="Sinh Slides/bai-XX.pptx từ Slides/bai-XX.markdown")
    parser.add_argument("course", help="Thư mục học phần")
    parser.add_argument("lessons", nargs="*", help="Số bài (01, 02...). Mặc định: mọi bài")
    args = parser.parse_args()

    course_dir = ROOT_DIR / args.course.strip("/\\")
    config_file = course_dir / "course.json"
    config = json.loads(config_file.read_text(encoding="utf-8")) if config_file.is_file() else {}
    slides_dir = course_dir / "Slides"
    files = sorted(f for f in slides_dir.glob("bai-*.markdown") if SLIDE_FILE_RE.match(f.name)) if slides_dir.is_dir() else []
    if args.lessons:
        files = [f for f in files if SLIDE_FILE_RE.match(f.name).group(1) in args.lessons]
    if not files:
        print(f"[LỖI] Không có file Slides/bai-XX.markdown trong {course_dir.name}")
        return 1

    exit_code = 0
    for md_file in files:
        try:
            count, missing = build_deck(md_file, md_file.with_suffix(".pptx"), config)
            print(f"[+] {course_dir.name}/Slides/{md_file.with_suffix('.pptx').name}: {count} slide")
            for path in missing:
                print(f"    [CẢNH BÁO] Chưa có hình {path} (đã chèn khung chờ; tải hình về Slides/images/ rồi build lại)")
        except (ValueError, PermissionError) as e:
            print(f"[LỖI] {md_file.name}: {e}")
            exit_code = 1
    return exit_code


if __name__ == "__main__":
    sys.exit(main())
