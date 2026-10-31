"""
Rà soát bộ hồ sơ xây dựng chương trình đào tạo (CTĐT) đại học chính quy.

Đọc thư mục hồ sơ gồm các biểu mẫu BM01–BM09 (docx, xlsx) và thư mục đề cương BM05,
đối chiếu chéo giữa các văn bản rồi in báo cáo LỖI / CẢNH BÁO dạng Markdown.

Cách dùng:
    uv run --with python-docx --with openpyxl python validate-ctdt.py "<thư mục hồ sơ>"
    uv run --with python-docx --with openpyxl python validate-ctdt.py "<thư mục hồ sơ>" --nam 2026 --report ra-soat.md

Tự nhận diện file theo tên:
    BM03*.xlsx              kế hoạch đào tạo (nguồn chuẩn về mã, tên, tín chỉ, học kỳ)
    *điều hành* / *dieu-hanh*.docx  chương trình đào tạo điều hành (học phần trước, tiên quyết)
    PLO*.xlsx               PEO/PLO/PI và ma trận CLO–PI (sheet "Map CLO-PI")
    BM05*/ (thư mục)        đề cương học phần
    mọi .docx khác          dò mã học phần, bảng kế hoạch, tỷ lệ % để đối chiếu

Exit code 1 nếu có LỖI; CẢNH BÁO không làm thất bại.
"""
import re
import sys
import argparse
import unicodedata
from collections import defaultdict
from pathlib import Path

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
import openpyxl

CODE_RE = re.compile(r"^\s*(\d{6})(?:\.0)?\s*$")
GIO_MOI_TIN_CHI = 50
# Dấu vết mẫu chưa điền hoặc nhãn còn sót từ khóa cũ
MARKERS = [
    (r"DOCX", "chữ 'DOCX' còn sót"),
    (r"……|\.{4,}", "dấu '……' chưa điền"),
    (r"[Cc]ần bổ sung", "'Cần bổ sung'"),
    (r"Tên tác giả \(năm", "tài liệu tham khảo còn dòng mẫu"),
    (r"\{\s*[Kk]hóa\s*\d{4}\s*\}|\{\s*20\d\d[^}]*\}", "nhãn '{Khóa …}' / '{20xx …}' còn sót"),
    (r"<\s*MÃ HỌC PHẦN", "tiêu đề mẫu <MÃ HỌC PHẦN …>"),
    (r"xoá toàn bộ khung|xóa toàn bộ khung", "khung hướng dẫn sử dụng mẫu chưa xóa"),
]


# ---------------------------------------------------------------- đọc file

def norm(s):
    """Chuẩn hóa tên học phần để so sánh: bỏ dấu câu, {…}, khoảng trắng, chữ hoa."""
    s = unicodedata.normalize("NFC", str(s or ""))
    s = re.sub(r"\{[^}]*\}", "", s)
    s = re.sub(r"[\s\-–—.:,()/&]+", " ", s)
    return s.strip().casefold()


def num(v):
    if v is None or str(v).strip() == "":
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def cell_text(c):
    return " / ".join(p.text.strip() for p in c.paragraphs if p.text.strip())


def docx_blocks(path):
    """Trả về (danh sách dòng văn bản, danh sách bảng[hàng[ô]]), giữ thứ tự xuất hiện."""
    doc = Document(path)
    lines, tables = [], []
    for el in doc.element.body.iterchildren():
        if el.tag.endswith('}p'):
            t = Paragraph(el, doc).text.strip()
            if t:
                lines.append(t)
        elif el.tag.endswith('}tbl'):
            rows = []
            for r in Table(el, doc).rows:
                cells, prev = [], None
                for c in r.cells:
                    if c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(cell_text(c))
                rows.append(cells)
                lines.append(" | ".join(cells))
            tables.append(rows)
    return lines, tables


class Report:
    def __init__(self):
        self.items = defaultdict(list)   # nhóm -> [(mức, nội dung)]

    def err(self, group, msg):
        self.items[group].append(("LỖI", msg))

    def warn(self, group, msg):
        self.items[group].append(("CẢNH BÁO", msg))

    def count(self, level):
        return sum(1 for v in self.items.values() for lv, _ in v if lv == level)

    def markdown(self):
        out = ["# Báo cáo rà soát hồ sơ CTĐT", "",
               f"Tổng: **{self.count('LỖI')} lỗi**, **{self.count('CẢNH BÁO')} cảnh báo**.", ""]
        for group, items in self.items.items():
            out.append(f"## {group}")
            out.append("")
            for lv, msg in items:
                out.append(f"- **{lv}:** {msg}")
            out.append("")
        return "\n".join(out)


# ---------------------------------------------------------------- kế hoạch đào tạo

class Course:
    def __init__(self, code, name, sem, tc, lt, th, bt, tiet, group=None, src=""):
        self.code, self.name, self.sem = code, name, sem
        self.tc, self.lt, self.th, self.bt, self.tiet = tc, lt, th, bt, tiet
        self.group, self.src = group, src
        self.prereq = []          # [(mã, loại)] loại: tiên quyết / học trước / song hành


def find_cols(header_rows):
    """Xác định vị trí cột từ 1–2 hàng tiêu đề của bảng kế hoạch."""
    keys = {"code": ("mã mh", "mã hp", "mã môn"), "name": ("môn học", "tên học phần", "tên môn"),
            "tc": ("tổng tc", "số tín chỉ"), "lt": ("lý thuyết",), "th": ("thực hành",),
            "bt": ("bài tập",), "tiet": ("số tiết tổng", "tổng số tiết"),
            "pre": ("học phần trước", "môn học trước")}
    cols = {}
    for row in header_rows:
        for i, v in enumerate(row):
            t = str(v or "").strip().casefold()
            for k, words in keys.items():
                if any(t.startswith(w) for w in words):
                    # hàng tiêu đề thứ hai chi tiết hơn ("Tổng TC", "Lý thuyết"…) nên ghi đè hàng đầu
                    cols[k] = i
    return cols


def parse_plan_rows(rows, src):
    """Phân tích bảng kế hoạch (từ xlsx hoặc bảng docx). Trả về (courses, tổng HK khai báo, tổng toàn khóa)."""
    hdr_idx = next((i for i, r in enumerate(rows[:15])
                    if any(str(v or "").strip().casefold().startswith(("mã mh", "mã hp")) for v in r)), None)
    if hdr_idx is None:
        return [], {}, None
    cols = find_cols(rows[hdr_idx:hdr_idx + 2])
    if "code" not in cols or "tc" not in cols:
        return [], {}, None
    courses, declared, total = [], {}, None
    sem, group, gid = 1, None, 0

    def get(r, k):
        i = cols.get(k)
        return r[i] if i is not None and i < len(r) else None

    for r in rows[hdr_idx + 1:]:
        first = " ".join(str(v) for v in r[:2] if v not in (None, "")).strip()
        up = first.upper()
        m = re.search(r"TỔNG CỘNG HỌC KỲ\s*(\d+)", up)
        # dòng tổng thường gộp ô nhãn nên lệch cột: lấy số đầu tiên sau nhãn làm tổng TC
        nums = [num(v) for v in r if num(v) is not None]
        if m:
            declared[int(m.group(1))] = {"tc": nums[0] if nums else None}
            sem, group = int(m.group(1)) + 1, None
            continue
        if "TỔNG CỘNG TOÀN KHÓA" in up:
            total = nums[0] if nums else None
            break
        if "TỰ CHỌN" in up:
            gid += 1
            k = re.search(r"CHỌN\s*(\d+)\s*TRONG\s*(\d+)", up)
            group = (gid, int(k.group(1)) if k else 1)
            continue
        cm = CODE_RE.match(str(get(r, "code") or ""))
        if not cm:
            continue
        c = Course(cm.group(1), str(get(r, "name") or "").strip(), sem,
                   num(get(r, "tc")), num(get(r, "lt")), num(get(r, "th")), num(get(r, "bt")),
                   num(get(r, "tiet")), group, src)
        pre = str(get(r, "pre") or "")
        for pm in re.finditer(r"(\d{6})\s*(\*{1,3})?", pre):
            kind = {"*": "tiên quyết", "**": "học trước", "***": "song hành"}.get(pm.group(2) or "**", "học trước")
            c.prereq.append((pm.group(1), kind))
        courses.append(c)
    return courses, declared, total


def load_xlsx_plan(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return parse_plan_rows(rows, path.name) + (ws.title,)


def check_plan(rep, label, courses, declared, total):
    g = f"Kế hoạch đào tạo – {label}"
    if not courses:
        rep.warn(g, "không đọc được bảng kế hoạch")
        return
    seen = {}
    for c in courses:
        if c.code in seen:
            rep.err(g, f"mã {c.code} xuất hiện hai lần ({seen[c.code].name} / {c.name})")
        seen[c.code] = c
        if c.tc is None:
            if c.group:
                rep.warn(g, f"{c.code} {c.name}: học phần tự chọn chưa ghi số tín chỉ, LT/TH/BT")
            continue
        parts = [v or 0 for v in (c.lt, c.th, c.bt)]
        if any(v is not None for v in (c.lt, c.th, c.bt)) and abs(sum(parts) - c.tc) > 1e-6:
            rep.err(g, f"{c.code} {c.name}: LT+TH+BT = {sum(parts):g} ≠ {c.tc:g} tín chỉ")
        if c.tiet is not None and c.tc and abs(c.tiet - GIO_MOI_TIN_CHI * c.tc) > 1e-6:
            rep.warn(g, f"{c.code} {c.name}: số tiết tổng {c.tiet:g} ≠ {GIO_MOI_TIN_CHI}×{c.tc:g} = {GIO_MOI_TIN_CHI * c.tc:g}")
    # tổng học kỳ: học phần bắt buộc + k học phần đầu tiên có TC của mỗi nhóm tự chọn
    by_sem = defaultdict(list)
    for c in courses:
        by_sem[c.sem].append(c)
    grand = 0
    for sem in sorted(by_sem):
        tc = sum(c.tc or 0 for c in by_sem[sem] if not c.group)
        groups = defaultdict(list)
        for c in by_sem[sem]:
            if c.group:
                groups[c.group].append(c)
        for (gid, k), items in groups.items():
            with_tc = [c for c in items if c.tc]
            tc += sum(c.tc for c in with_tc[:k])
            tcs = {c.tc for c in with_tc}
            if len(tcs) > 1:
                rep.err(g, f"HK{sem}: nhóm tự chọn có số tín chỉ khác nhau {sorted(tcs)} ({', '.join(c.code for c in items)})")
        grand += tc
        d = declared.get(sem, {}).get("tc")
        if d is not None and abs(d - tc) > 1e-6:
            rep.err(g, f"HK{sem}: dòng tổng ghi {d:g} tín chỉ, cộng các học phần được {tc:g}")
    if total is not None and abs(total - grand) > 1e-6:
        rep.err(g, f"Tổng toàn khóa ghi {total:g} tín chỉ, cộng các học kỳ được {grand:g}")
    dsum = sum(v["tc"] or 0 for v in declared.values())
    if total is not None and declared and abs(dsum - total) > 1e-6:
        rep.err(g, f"Tổng toàn khóa {total:g} ≠ tổng các dòng 'TỔNG CỘNG HỌC KỲ' ({dsum:g})")


def check_prereq(rep, label, courses):
    g = f"Học phần trước / tiên quyết – {label}"
    by = {c.code: c for c in courses}
    for c in courses:
        for code, kind in c.prereq:
            p = by.get(code)
            if not p:
                rep.err(g, f"{c.code} {c.name}: học phần {kind} {code} không có trong chương trình")
                continue
            if kind == "song hành":
                if p.sem > c.sem:
                    rep.err(g, f"{c.code} (HK{c.sem}) song hành với {code} ở HK{p.sem} sau đó")
            elif p.sem >= c.sem:
                rep.err(g, f"{c.code} {c.name} (HK{c.sem}): học phần {kind} {code} {p.name} xếp ở HK{p.sem}, không trước")
            if p.group and not c.group:
                rep.warn(g, f"{c.code} {c.name} bắt buộc nhưng yêu cầu {code} {p.name} là học phần TỰ CHỌN: "
                            "sinh viên chọn môn còn lại sẽ không đủ điều kiện")
            elif p.group and c.group:
                rep.warn(g, f"{c.code} {c.name} (tự chọn) yêu cầu {code} {p.name} (tự chọn): "
                            "chuỗi tự chọn khóa lựa chọn của sinh viên, cần ghi rõ lộ trình")


def compare_plans(rep, ref, other, label):
    g = f"Đối chiếu kế hoạch BM03 ↔ {label}"
    a = {c.code: c for c in ref}
    b = {c.code: c for c in other}
    for code in sorted(set(a) - set(b)):
        rep.err(g, f"{code} {a[code].name} có trong BM03, không có trong {label}")
    for code in sorted(set(b) - set(a)):
        rep.err(g, f"{code} {b[code].name} có trong {label}, không có trong BM03")
    for code in sorted(set(a) & set(b)):
        x, y = a[code], b[code]
        if norm(x.name) != norm(y.name):
            rep.warn(g, f"{code}: tên khác nhau – BM03 '{x.name}' / {label} '{y.name}'")
        if x.sem != y.sem:
            rep.err(g, f"{code} {x.name}: BM03 xếp HK{x.sem}, {label} xếp HK{y.sem}")
        for f, nm in (("tc", "tín chỉ"), ("lt", "LT"), ("th", "TH"), ("bt", "BT"), ("tiet", "số tiết")):
            u, v = getattr(x, f), getattr(y, f)
            if u is not None and v is not None and abs(u - v) > 1e-6:
                rep.err(g, f"{code} {x.name}: {nm} BM03 = {u:g}, {label} = {v:g}")


# ---------------------------------------------------------------- PLO / ma trận CLO–PI

def load_plo(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    pis, cmap = set(), {}
    if "PLO" in wb.sheetnames:
        for r in wb["PLO"].iter_rows(values_only=True):
            m = re.match(r"PI\s*(\d+\.\d+)", str(r[0] or ""))
            if m:
                pis.add(m.group(1))
    sheet = next((s for s in wb.sheetnames if "clo" in s.casefold()), None)
    if sheet:
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        hdr = next((r for r in rows[:6] if sum(1 for v in r if re.match(r"PI\s*\d", str(v or ""))) > 3), None)
        if hdr:
            picol = {i: re.sub(r"PI\s*", "", str(v)) for i, v in enumerate(hdr) if re.match(r"PI\s*\d", str(v or ""))}
            for r in rows:
                m = CODE_RE.match(str(r[0] or ""))
                if not m:
                    continue
                d = defaultdict(set)
                for i, pi in picol.items():
                    v = str(r[i] or "") if i < len(r) else ""
                    for clo in re.findall(r"CLO\s*(\d+)", v):
                        d[int(clo)].add(pi)
                    if v.strip().lower() == "x":
                        d[0].add(pi)
                # số CLO = số CLO lớn nhất được ghi trong hàng (cột "đếm" của sheet là số ô PI đã đánh dấu)
                cmap[m.group(1)] = {"name": r[1], "nclo": max((k for k in d if k), default=None), "map": d}
    return pis, cmap


# ---------------------------------------------------------------- đề cương

def parse_syllabus(path):
    lines, tables = docx_blocks(path)
    text = "\n".join(lines)
    info = {"path": path, "text": text, "lines": lines}

    def field(*labels):
        for row in (r for t in tables for r in t):
            if len(row) >= 2 and any(row[0].strip().casefold().startswith(l) for l in labels):
                return row[1].strip()
        return None

    code = field("mã môn học", "mã học phần")
    cm = re.search(r"\d{6}", code or "") or re.search(r"(?<!\d)(1\d{5})(?!\d)", path.stem)
    info["code"] = cm.group(0) if cm else None
    info["name"] = field("tên môn học (tiếng việt)", "tên học phần (tiếng việt)")
    tcm = re.match(r"\s*(\d+)", field("số tín chỉ") or "")
    info["tc"] = float(tcm.group(1)) if tcm else None
    def lead(v):
        m = re.match(r"\s*(\d+(?:[.,]\d+)?)", v or "")
        return float(m.group(1).replace(",", ".")) if m else 0.0

    ltb = field("lý thuyết / thực hành", "lý thuyết/thực hành")
    paren = re.search(r"\(\s*(\d+)\s*[-,]\s*(\d+)\s*[-,]\s*(\d+)\s*\)", field("số tín chỉ") or "")
    if ltb:
        n = [float(x) for x in re.findall(r"(\d+(?:[.,]\d+)?)\s*(?:lý|thực|bài)", ltb.casefold())]
        info["ltb"] = tuple(n[:3]) if len(n) >= 3 else None
    elif paren:                                   # "2(2-0-0)"
        info["ltb"] = tuple(float(x) for x in paren.groups())
    else:
        lt, th, bt = field("lý thuyết"), field("thực hành"), field("bài tập")
        info["ltb"] = tuple(lead(x) for x in (lt, th, bt)) if lt is not None else None
    pre = field("môn học trước", "học phần học trước", "học phần trước") or ""
    info["prereq"] = re.findall(r"\d{6}", pre)
    info["prereq_text"] = pre
    # CLO khai báo trong Bảng 1
    clos = {}
    for t in tables:
        for r in t:
            m = re.match(r"^CLO\s*(\d+)$", r[0].strip()) if r else None
            if m and len(r) >= 3 and int(m.group(1)) not in clos:
                clos[int(m.group(1))] = r
    info["clos"] = clos
    # Bảng đánh giá: hàng bắt đầu bằng A1., A2.…
    assess = []
    for t in tables:
        for r in t:
            if r and re.match(r"^A\d+\s*[.:]", r[0].strip()):
                last = re.fullmatch(r"\s*(\d+(?:[.,]\d+)?)\s*%?\s*", r[-1])
                pct = [last.group(1)] if last else re.findall(r"(\d+(?:[.,]\d+)?)\s*%", " ".join(r))
                assess.append((r[0].strip()[:40], set(int(x) for x in re.findall(r"CLO\s*(\d+)", " ".join(r))),
                               float(pct[-1].replace(",", ".")) if pct else None))
    info["assess"] = assess
    # Ma trận LLO–CLO: cột tiêu đề CLO1..CLOn
    mapped = set()
    for t in tables:
        hdr = next((r for r in t[:3] if sum(1 for v in r if re.fullmatch(r"CLO\s*\d+", v.strip())) >= 2), None)
        # LLO ở cột đầu, hoặc cột thứ hai khi cột đầu là "Bài/Chương"
        def llo_pos(r):
            return next((i for i, c in enumerate(r[:2]) if c.startswith("LLO")), None)
        if not hdr or not any(llo_pos(r) is not None for r in t if r):
            continue
        idx = {i: int(re.sub(r"\D", "", v)) for i, v in enumerate(hdr) if re.fullmatch(r"CLO\s*\d+", v.strip())}
        for r in t:
            pos = llo_pos(r) if r else None
            if pos is not None:
                for i, k in idx.items():
                    # bảng có thể lệch một cột do ô gộp; dò cả hai vị trí
                    for j in (i, i - 1):
                        if pos < j < len(r) and r[j].strip().lower() == "x":
                            mapped.add(k)
                            break
    info["mapped"] = mapped
    info["pi"] = re.findall(r"PI\s*(\d+\.\d+)", " ".join(" ".join(r) for r in clos.values()))
    years = re.findall(r"[Nn]gày\s*[.…\d ]*\s*tháng\s*[.…\d ]*\s*năm\s*(20\d\d)", text)
    info["years"] = years
    info["kind"] = ("BM05a" if "TÍCH HỢP TOÀN PHẦN" in text.upper()
                    else "BM05b" if "TĂNG CƯỜNG TỪ HỌC LIỆU ASU" in text.upper() else "BM05")
    return info


def check_syllabi(rep, files, plan, plo, year):
    by = {c.code: c for c in plan}
    seen = defaultdict(list)
    for f in files:
        try:
            s = parse_syllabus(f)
        except Exception as e:      # file hỏng hoặc không phải docx
            rep.err("Đề cương – đọc file", f"{f.name}: không đọc được ({e})")
            continue
        g = f"Đề cương {s['code'] or '?'} – {f.name}"
        if not s["code"]:
            rep.err(g, "không xác định được mã học phần")
            continue
        seen[s["code"]].append(f.name)
        fm = re.search(r"(?<!\d)(1\d{5})(?!\d)", f.stem)
        if fm and fm.group(1) != s["code"]:
            rep.err(g, f"mã trong tên file ({fm.group(1)}) khác mã trong đề cương ({s['code']})")
        c = by.get(s["code"])
        if plan and not c:
            rep.err(g, f"mã {s['code']} không có trong kế hoạch đào tạo BM03 (học phần đã bỏ hoặc sai mã?)")
        if c:
            if s["name"] and norm(s["name"]) != norm(c.name):
                rep.warn(g, f"tên '{s['name']}' khác BM03 '{c.name}'")
            if s["tc"] and c.tc and abs(s["tc"] - c.tc) > 1e-6:
                rep.err(g, f"số tín chỉ {s['tc']:g} khác BM03 {c.tc:g}")
            ltb = s.get("ltb")
            if ltb and any(ltb) and c.tc and None not in (c.lt, c.th):
                plan_ltb = (c.lt or 0, c.th or 0, c.bt or 0)
                shown = '-'.join(f'{x:g}' for x in ltb)
                if sum(ltb) >= 10 * c.tc:
                    rep.warn(g, f"LT/TH/BT ghi {shown} (số tiết, không phải tín chỉ); BM03 là {'-'.join(f'{x:g}' for x in plan_ltb)} tín chỉ")
                elif tuple(round(x, 2) for x in ltb) != plan_ltb:
                    rep.err(g, f"LT/TH/BT {shown} khác BM03 {'-'.join(f'{x:g}' for x in plan_ltb)}")
            want = {p for p, _ in c.prereq}
            if want and set(s["prereq"]) != want:
                rep.err(g, f"học phần trước ghi '{s['prereq_text'][:60]}', chương trình điều hành yêu cầu {', '.join(sorted(want))}")
            if not want and s["prereq"]:
                rep.warn(g, f"đề cương ghi học phần trước {', '.join(s['prereq'])} nhưng chương trình điều hành không có")
        # CLO và đánh giá
        n = len(s["clos"])
        if n == 0:
            rep.err(g, "không tìm thấy bảng CLO")
        else:
            if sorted(s["clos"]) != list(range(1, n + 1)):
                rep.warn(g, f"CLO đánh số không liên tục: {sorted(s['clos'])}")
            for k, r in s["clos"].items():
                # cột PLO/SO/PI không cố định (có mẫu đặt ngay sau CLO), nên dò cả hàng
                if not re.search(r"(SO|PLO|PI)\s*\d", " ".join(r[1:])) and s["code"] and not s["code"].startswith("102"):
                    rep.warn(g, f"CLO{k} chưa gắn PLO/SO/PI")
            if s["mapped"]:
                for k in s["clos"]:
                    if k not in s["mapped"]:
                        rep.warn(g, f"CLO{k} không có bài học (LLO) nào đánh dấu trong ma trận LLO–CLO")
        if s["assess"]:
            tot = sum(p for _, _, p in s["assess"] if p is not None)
            if all(p is None for _, _, p in s["assess"]):
                rep.warn(g, "Bảng đánh giá không ghi tỷ lệ % cho các thành phần")
            elif abs(tot - 100) > 1e-6:
                rep.err(g, f"tổng tỷ lệ đánh giá = {tot:g}% (phải 100%)")
            assessed = set().union(*(a for _, a, _ in s["assess"]))
            for k in s["clos"]:
                if k not in assessed:
                    rep.err(g, f"CLO{k} không được đánh giá ở thành phần nào (Bảng đánh giá)")
            for k in sorted(assessed - set(s["clos"])):
                rep.err(g, f"Bảng đánh giá nhắc CLO{k} nhưng Bảng CLO không có")
        else:
            rep.warn(g, "không tìm thấy Bảng đánh giá (A1., A2., …)")
        # PI và ma trận CLO–PI
        if plo:
            pis, cmap = plo
            for pi in s["pi"]:
                if pis and pi not in pis:
                    rep.err(g, f"PI{pi} không có trong danh mục PI của chương trình")
            m = cmap.get(s["code"])
            if m and m["nclo"] and n and int(m["nclo"]) != n:
                rep.warn(g, f"đề cương có {n} CLO, ma trận CLO–PI (file PLO) ghi {int(m['nclo'])}")
            if m:
                for k, r in s["clos"].items():
                    doc_pi = set(re.findall(r"PI\s*(\d+\.\d+)", " ".join(r)))
                    xl_pi = m["map"].get(k, set())
                    if doc_pi and xl_pi and doc_pi != xl_pi:
                        rep.warn(g, f"CLO{k}: đề cương gắn PI {', '.join(sorted(doc_pi))}, ma trận CLO–PI gắn {', '.join(sorted(xl_pi))}")
        # dấu vết mẫu, năm
        for pat, label in MARKERS:
            hits = re.findall(pat, s["text"])
            if hits:
                rep.warn(g, f"{label} ({len(hits)} chỗ)")
        if re.search(r"Email\s*:\s*(\||$|\n)", s["text"]) or "(Cần bổ sung Email)" in s["text"]:
            rep.warn(g, "email giảng viên phụ trách để trống")
        if year:
            bad = sorted({y for y in s["years"] if y != str(year)})
            if bad:
                rep.warn(g, f"ngày ký ghi năm {', '.join(bad)} (đợt xây dựng {year})")
    for code, names in seen.items():
        if len(names) > 1:
            rep.err("Đề cương – trùng mã", f"{code} có {len(names)} file: {'; '.join(names)}")
    for c in plan:
        if c.code not in seen and not c.code.startswith("102"):
            rep.warn("Đề cương – còn thiếu", f"{c.code} {c.name} (HK{c.sem}) chưa có đề cương trong thư mục BM05")


# ---------------------------------------------------------------- văn bản khác

def check_documents(rep, files, plan, year):
    by = {c.code: c for c in plan}
    for f in files:
        try:
            lines, tables = docx_blocks(f)
        except Exception as e:
            rep.err("Văn bản – đọc file", f"{f.name}: không đọc được ({e})")
            continue
        g = f"Văn bản – {f.name}"
        text = "\n".join(lines)
        # bảng kế hoạch nhúng trong văn bản (BM06, BM07, chương trình điều hành)
        for t in tables:
            if any("TỔNG CỘNG HỌC KỲ" in " ".join(r).upper() for r in t):
                courses, declared, total = parse_plan_rows(t, f.name)
                if courses:
                    check_plan(rep, f.name, courses, declared, total)
                    if plan:
                        compare_plans(rep, plan, courses, f.name)
        # mã học phần nhắc trong văn bản nhưng không thuộc kế hoạch
        if plan:
            for code in sorted(set(re.findall(r"(?<!\d)(1[019]\d{4})(?!\d)", text))):
                if code not in by:
                    ctx = re.search(rf"{code}[^\n|]{{0,50}}", text).group(0)
                    rep.warn(g, f"nhắc mã {code} không có trong BM03: '{ctx.strip()}'")
        # tỷ lệ phần trăm dạng a/b … x%
        for m in re.finditer(r"(\d+)\s*/\s*(\d+)[\s/]*Tỷ lệ\s*([\d.,]+)\s*%", text):
            a, b, p = int(m.group(1)), int(m.group(2)), float(m.group(3).replace(",", "."))
            if b and abs(a / b * 100 - p) > 0.05:
                rep.err(g, f"{a}/{b} = {a / b * 100:.2f}% nhưng văn bản ghi {p:g}%")
        for pat, label in MARKERS[2:5]:
            if re.search(pat, text):
                rep.warn(g, label)
        if year:
            dates = re.findall(r"ngày\s+(\d{1,2})\s*/\s*(\d{1,2})\s*/\s*(20\d\d)", text)
            heads = re.findall(r"ngày\s+(\d{1,2})\s+tháng\s+(\d{1,2})\s+năm\s+(20\d\d)", text[:600])
            if heads and dates:
                d0 = tuple(int(x) for x in heads[0][::-1])
                later = sorted({"/".join(d) for d in dates if tuple(int(x) for x in d[::-1]) > d0})
                if later:
                    rep.err(g, f"văn bản đề ngày {'/'.join(heads[0])} nhưng căn cứ/nhắc văn bản ngày sau đó: {', '.join(later)}")


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder", type=Path)
    ap.add_argument("--nam", type=int, help="năm của đợt xây dựng CTĐT (kiểm tra ngày ký)")
    ap.add_argument("--report", type=Path, help="ghi báo cáo Markdown ra file")
    a = ap.parse_args()
    root = a.folder
    rep = Report()

    plan_x = sorted(root.glob("BM03*.xlsx"))
    plan, prereq_src = [], None
    if plan_x:
        plan, declared, total, sheet = load_xlsx_plan(plan_x[0])
        check_plan(rep, f"{plan_x[0].name} (sheet {sheet})", plan, declared, total)
    else:
        rep.warn("Hồ sơ", "không tìm thấy BM03*.xlsx (kế hoạch đào tạo)")

    docs = [p for p in root.glob("*.docx") if not p.name.startswith("~$")]
    dh = next((p for p in docs if re.search(r"điều hành|dieu-hanh", unicodedata.normalize("NFC", p.name).casefold())), None)
    if dh:
        _, tables = docx_blocks(dh)
        for t in tables:
            courses, declared, total = parse_plan_rows(t, dh.name)
            if courses:
                prereq_src = courses
                break
        if prereq_src:
            check_prereq(rep, dh.name, prereq_src)
            pre = {c.code: c.prereq for c in prereq_src}
            for c in plan:
                c.prereq = pre.get(c.code, [])
            if not plan:
                plan = prereq_src

    plo_x = sorted(root.glob("PLO*.xlsx"))
    plo = load_plo(plo_x[0]) if plo_x else None

    bm05 = [p for d in root.glob("BM05*") if d.is_dir() for p in sorted(d.glob("*.docx"))
            if not p.name.startswith("~$")]
    syllabi = [p for p in bm05 if re.search(r"(?<!\d)1\d{5}(?!\d)", p.stem)]
    check_syllabi(rep, syllabi, plan, plo, a.nam)
    check_documents(rep, docs, plan, a.nam)

    md = rep.markdown()
    if a.report:
        a.report.write_text(md, encoding="utf-8")
        print(f"[+] Đã ghi {a.report}")
    print(md)
    sys.exit(1 if rep.count("LỖI") else 0)


if __name__ == "__main__":
    main()
